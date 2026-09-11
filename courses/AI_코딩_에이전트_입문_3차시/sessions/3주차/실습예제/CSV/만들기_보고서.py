#!/usr/bin/env python3
"""만들기_보고서.py — 지표기준값.json만 읽어 분석보고서.docx(강사용 완성본)를 만든다.

    python 만들기_보고서.py              → 분석보고서.docx        (제목 + 다섯 칸 + 표 하나 + 차트 하나)
    python 만들기_보고서.py --recovery   → 복구/분석보고서.docx   (한 쪽 · 다섯 칸)
    python 만들기_보고서.py --out 경로   → 원하는 곳에 저장

- 필요 패키지: python-docx · Pillow(차트 이미지). 한글 글꼴(Malgun Gothic 등)이 있어야 차트 글자가 보인다.
  matplotlib은 쓰지 않는다 — 없는 환경이 있어 Pillow 한 경로로만 그려 결과를 같게 한다.
- 지표 숫자는 전부 지표기준값.json에서 가져온다. 이 파일에 손으로 적은 지표 숫자는 없다.
- 차트 PNG에는 그린 글자 목록을 'labels' 텍스트 청크로 넣는다. 결과물_대조.py가 그 목록의 숫자를 대조한다
  (이미지 픽셀을 읽는 OCR은 하지 않는다).
"""
import argparse
import io
import json
import os
import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(errors="replace")
except Exception:
    pass

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Mm, Pt, RGBColor
except ImportError:
    sys.exit("FAIL python-docx가 없다 — 설치가 필요하다(pip install python-docx). 파일을 만들지 않았다.")

HERE = Path(__file__).resolve().parent
FONT = "Malgun Gothic"
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x4B, 0x55, 0x63)
SECTIONS = ["핵심 요약", "주요 지표", "발견한 특징", "개선이 필요한 부분", "데이터의 한계"]


# ── 데이터 ──────────────────────────────────────────────────────────
def won(v):
    return f"{int(v):,}원"


def cnt(v):
    return f"{int(v):,}"


def pct(v):
    return f"{v:.1f}%"


def month_name(ym):
    return f"{int(ym.split('-')[1])}월"


def load(path):
    b = json.loads(Path(path).read_text(encoding="utf-8"))
    for k in ("ROW", "ORD", "REV", "MGN", "MONTH", "CAT", "DELIV", "RATE", "RET", "HOLD"):
        if k not in b:
            sys.exit(f"FAIL 지표기준값.json에 {k} 절이 없다")
    cats = b["CAT"]["CAT-02_목록"]
    top_rev = max(cats, key=lambda c: c["순매출_krw"])
    low_rev = min(cats, key=lambda c: c["순매출_krw"])
    top_ord = max(cats, key=lambda c: c["주문수"])
    months, mrev = b["MONTH"]["MONTH-01_월목록"], b["MONTH"]["MONTH-02_월별순매출_krw"]
    imax, imin = mrev.index(max(mrev)), mrev.index(min(mrev))
    hold_del = b["HOLD"]["HOLD-03_배송이상"]
    del_days = {h["delivery_days"] for h in hold_del}
    reasons = b["RET"]["RET-03_사유분포"]
    checks = [
        (top_rev is min(cats, key=lambda c: c["마진율_pct"]), "매출 1위 상품군의 마진율이 가장 낮다"),
        (low_rev is max(cats, key=lambda c: c["마진율_pct"]), "매출 최하위 상품군의 마진율이 가장 높다"),
        (top_ord is not top_rev, "주문 수 1위와 매출 1위가 다르다"),
        (months[imax] == b["MONTH"]["MONTH-05_최대월"] and months[imin] == b["MONTH"]["MONTH-06_최소월"], "최대·최소월"),
        (len(b["HOLD"]["HOLD-02_중복충돌"]) == b["ROW"]["ROW-05_보류행수"], "중복 충돌 행 수 = ROW-05"),
        (len(hold_del) == b["DELIV"]["DELIV-07_보류제외행수"] and len(del_days) == 1, "배송 이상 행"),
    ]
    bad = [m for ok, m in checks if not ok]
    if bad:
        sys.exit("FAIL 본문의 결론이 데이터와 맞지 않아 만들지 않았다: " + " / ".join(bad))
    f = dict(cats=cats, top_rev=top_rev, low_rev=low_rev, top_ord=top_ord, months=months, mrev=mrev,
             imax=imax, imin=imin, del_days=del_days.pop(), reasons=reasons, top_reason=max(reasons, key=reasons.get))
    return b, f


# ── 문서 도구 ────────────────────────────────────────────────────────
def set_style_font(style, size=None, bold=None, color=None):
    style.font.name = FONT
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)


def new_document(title):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(25)
    sec.top_margin = sec.bottom_margin = Mm(22)
    set_style_font(doc.styles["Normal"], 10.5, None, INK)
    set_style_font(doc.styles["Title"], 22, True, INK)
    set_style_font(doc.styles["Heading 1"], 14, True, INK)
    set_style_font(doc.styles["List Bullet"], 10.5, None, INK)
    doc.styles["Normal"].paragraph_format.space_after = Pt(4)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(4)
    doc.core_properties.title = title
    doc.core_properties.author = "AI 코딩 에이전트 입문 3차시 강사용 예시"
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = ""
    return doc


def para(doc, text, style=None, size=None, color=None, italic=False, align=None):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    if size:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    run.font.italic = italic
    if align is not None:
        p.alignment = align
    return p


def bullets(doc, lines):
    for t in lines:
        doc.add_paragraph(t, style="List Bullet")


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def category_table(doc, b, f):
    head = ["상품군", "총매출(원)", "매출 비중", "마진율"]
    rows = [[c["category"], cnt(c["순매출_krw"]), pct(c["매출비중_pct"]), pct(c["마진율_pct"])] for c in f["cats"]]
    rows.append(["전체", cnt(b["REV"]["REV-01_순매출_krw"]), "—", pct(b["MGN"]["MGN-03_마진율_pct"])])
    table = doc.add_table(rows=1 + len(rows), cols=len(head))
    table.style = doc.styles["Table Grid"]
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Mm(45), Mm(45), Mm(35), Mm(35)]
    for r, values in enumerate([head] + rows):
        for c, text in enumerate(values):
            cell = table.cell(r, c)
            cell.width = widths[c]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(text)
            run.font.bold = r == 0 or r == len(rows)
            if r == 0:
                shade(cell, "E5E7EB")
    return table


# ── 차트 이미지(Pillow) ──────────────────────────────────────────────
def find_font(bold):
    names = ["malgunbd.ttf", "malgun.ttf"] if bold else ["malgun.ttf"]
    names += ["NanumGothicBold.ttf", "NanumGothic.ttf"] if bold else ["NanumGothic.ttf"]
    names += ["AppleSDGothicNeo.ttc"]
    dirs = [Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts",
            Path("/usr/share/fonts/truetype/nanum"), Path("/System/Library/Fonts"), Path("/Library/Fonts")]
    for n in names:
        for d in dirs:
            if (d / n).exists():
                return str(d / n)
    return None


def month_chart_png(f):
    try:
        from PIL import Image, ImageDraw, ImageFont
        from PIL.PngImagePlugin import PngInfo
    except ImportError:
        sys.exit("FAIL Pillow가 없다 — 차트 이미지를 그릴 수 없어 멈췄다(pip install Pillow).")
    regular, bold = find_font(False), find_font(True)
    if regular is None:
        sys.exit("FAIL 한글 글꼴(Malgun Gothic·NanumGothic 등)을 찾지 못해 차트 글자를 그릴 수 없어 멈췄다.")
    f_lab = ImageFont.truetype(regular, 32)
    f_val = ImageFont.truetype(bold or regular, 32)
    f_leg = ImageFont.truetype(regular, 30)
    W, H = 1800, 900
    left, right, top, bottom = 60, 40, 130, 90
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    labels = []

    def text(xy, s, font, fill, anchor):
        d.text(xy, s, font=font, fill=fill, anchor=anchor)
        labels.append(s)

    vals, n = f["mrev"], len(f["mrev"])
    plot_w, plot_h = W - left - right, H - top - bottom
    base = top + plot_h
    scale = plot_h / (max(vals) * 1.12)
    for k in (1, 2, 3, 4):
        y = base - plot_h * k / 4
        d.line([(left, y), (W - right, y)], fill=(229, 231, 235), width=2)
    slot = plot_w / n
    bw = slot * 0.62
    colors = {f["imax"]: (29, 78, 216), f["imin"]: (180, 83, 9)}
    for i, v in enumerate(vals):
        x0 = left + slot * i + (slot - bw) / 2
        y0 = base - v * scale
        d.rectangle([x0, y0, x0 + bw, base], fill=colors.get(i, (148, 163, 184)))
        text((x0 + bw / 2, base + 14), month_name(f["months"][i]), f_lab, (75, 85, 99), "mt")
        if i in colors:
            text((x0 + bw / 2, y0 - 10), won(v), f_val, colors[i], "mb")
    d.line([(left, base), (W - right, base)], fill=(107, 114, 128), width=3)
    text((left, 40), "막대 높이 = 월별 총매출(원)", f_leg, (31, 41, 55), "lm")
    lx = left + 520
    for label, color in (("가장 높은 달", colors[f["imax"]]), ("가장 낮은 달", colors[f["imin"]])):
        d.rectangle([lx, 26, lx + 28, 54], fill=color)
        text((lx + 40, 40), label, f_leg, (31, 41, 55), "lm")
        lx += 300
    info = PngInfo()
    info.add_text("labels", json.dumps(labels, ensure_ascii=False))
    info.add_text("source", "지표기준값.json MONTH")
    buf = io.BytesIO()
    img.save(buf, "PNG", pnginfo=info, dpi=(250, 250))
    buf.seek(0)
    return buf


def add_chart(doc, f):
    doc.add_picture(month_chart_png(f), width=Mm(160))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.inline_shapes[-1]._inline.docPr.set("descr", "월별 총매출 막대그래프 — 가장 높은 달과 가장 낮은 달을 색으로 표시")
    para(doc, "그림. 월별 총매출(원) — 가장 높은 달과 가장 낮은 달에만 금액을 적었다", size=9, color=MUTED,
         italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)


# ── 본문 ─────────────────────────────────────────────────────────────
def body(doc, b, f, recovery):
    R, O, REV, MGN, RATE, D, RET, H = (b[k] for k in ("ROW", "ORD", "REV", "MGN", "RATE", "DELIV", "RET", "HOLD"))
    top, low, tord = f["top_rev"], f["low_rev"], f["top_ord"]
    hi, lo = f["months"][f["imax"]], f["months"][f["imin"]]
    rev01, ord01, rate01 = won(REV["REV-01_순매출_krw"]), cnt(O["ORD-01_분석대상주문수"]), f"{RATE['RATE-01_평균']:.2f}점"

    doc.add_heading(SECTIONS[0], level=1)
    para(doc, f"원본 {cnt(R['ROW-01_원본행수'])}행을 정리해 주문 {cnt(R['ROW-03_분석대상행수'])}건(원본의 {pct(R['ROW-06_채택비율_pct'])})을 "
              f"분석했습니다. 총매출은 {rev01}, 주문 수는 {ord01}건, 평균 평점은 {rate01}입니다. "
              + ("" if recovery else f"월별 매출은 한 방향 추세 없이 오르내렸고, 매출 1위 상품군 {top['category']}의 마진율이 가장 낮았습니다. ")
              + f"결론으로 쓰기 전에 사람이 정해야 할 기록 {H['HOLD-01_승인필요건수']}건이 남아 있습니다.")

    doc.add_heading(SECTIONS[1], level=1)
    lines = [
        f"총매출 {rev01} — 취소·대기 제외, 반품 미차감",
        f"주문 수 {ord01}건 — 중복을 정리해 한 행이 한 주문",
        f"평균 평점 {rate01} — 응답 {cnt(RATE['RATE-02_분모행수'])}건 기준(무응답 {RATE['RATE-03_무응답행수']}건은 0점으로 세지 않음)",
        f"마진율 {pct(MGN['MGN-03_마진율_pct'])} — 상품 원가만 뺀 값이라 이익률이 아님",
    ]
    if not recovery:
        lines += [
            f"마진 {won(MGN['MGN-02_마진_krw'])} · 고유 고객 {cnt(O['ORD-03_고유고객수'])}명",
            f"평균 주문금액 {won(REV['REV-06_평균주문금액_krw'])} · 평균 배송일 {D['DELIV-01_평균일']:.2f}일(배송 기록 {cnt(D['DELIV-05_분모행수'])}건 기준)",
        ]
    bullets(doc, lines)
    if not recovery:
        para(doc, "표. 상품군별 총매출과 마진율", size=9, color=MUTED, italic=True)
        category_table(doc, b, f)

    doc.add_heading(SECTIONS[2], level=1)
    if recovery:
        bullets(doc, [
            f"매출 1위 상품군 {top['category']}(비중 {pct(top['매출비중_pct'])})의 마진율은 {pct(top['마진율_pct'])}로 가장 낮습니다.",
            f"평점 4점 이상이 {pct(RATE['RATE-06_4점이상_pct'])}이고, 반품은 {cnt(RET['RET-01_반품건수'])}건({pct(RET['RET-02_반품률_pct'])})입니다.",
        ])
    else:
        para(doc, f"가장 높은 달은 {month_name(hi)}({won(f['mrev'][f['imax']])}), 가장 낮은 달은 {month_name(lo)}"
                  f"({won(f['mrev'][f['imin']])})입니다. 한 해 자료라 계절 효과인지는 판단하지 않았습니다.")
        add_chart(doc, f)
        tr = f["top_reason"]
        bullets(doc, [
            f"{top['category']}는 매출 비중 {pct(top['매출비중_pct'])}로 1위지만 마진율은 {pct(top['마진율_pct'])}로 가장 낮고, "
            f"{low['category']}는 매출 비중 {pct(low['매출비중_pct'])}로 가장 작지만 마진율은 {pct(low['마진율_pct'])}로 가장 높습니다.",
            f"주문 수 1위는 {tord['category']}({cnt(tord['주문수'])}건)로, 매출 1위 {top['category']}({cnt(top['주문수'])}건)와 다릅니다.",
            f"평점 4점 이상이 {pct(RATE['RATE-06_4점이상_pct'])}이고 1점은 {RATE['RATE-05_분포']['1']}건입니다. "
            f"반품은 {cnt(RET['RET-01_반품건수'])}건({pct(RET['RET-02_반품률_pct'])})이며 사유 중 {tr}({f['reasons'][tr]}건)가 가장 많습니다.",
        ])

    doc.add_heading(SECTIONS[3], level=1)
    bullets(doc, [
        f"사람이 정해야 할 기록 {H['HOLD-01_승인필요건수']}건: 같은 주문번호에 수량이 다른 기록 {R['ROW-05_보류행수']}행, "
        f"배송 {f['del_days']}일로 적힌 기록 {D['DELIV-07_보류제외행수']}행.",
        f"배송비 {won(REV['REV-05_배송비합계_krw'])}을 매출에 넣을지, 반품을 뺄지(빼면 총매출 {won(REV['REV-07_반품제외순매출_krw'])}) 정해야 합니다.",
    ] + ([] if recovery else ["배송일을 주문일부터 셌는지 발송일부터 셌는지, 평점을 어떻게 모았는지 확인해야 합니다."]))

    doc.add_heading(SECTIONS[4], level=1)
    bullets(doc, [
        "연습용 모의 데이터입니다. 실제 시장 판단의 근거로 쓰지 않습니다.",
        "마진은 상품 원가만 뺀 값입니다. 인건비·마케팅비·수수료 정보가 없어 이익을 계산할 수 없습니다.",
        "한 해 자료뿐이라 전년 대비나 계절 효과를 말할 수 없습니다.",
        "두 숫자가 함께 움직여 보여도 이 데이터만으로 원인을 말할 수 없습니다.",
    ])


def build(ref_path, out_path, recovery):
    b, f = load(ref_path)
    title = "2025년 주문 데이터 분석 보고서" + (" (복구본)" if recovery else "")
    doc = new_document(title)
    doc.add_paragraph(title, style="Title")
    O = b["ORD"]
    para(doc, f"연습용 모의 데이터 · 기간 {O['ORD-06_기간시작']} ~ {O['ORD-07_기간종료']} · 숫자 출처 지표기준값.json",
         size=9, color=MUTED)
    body(doc, b, f, recovery)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return len(doc.tables), len(doc.inline_shapes)


def main():
    ap = argparse.ArgumentParser(description="지표기준값.json → 분석보고서.docx")
    ap.add_argument("--recovery", action="store_true", help="복구본(한 쪽)을 만든다")
    ap.add_argument("--ref", default=str(HERE / "지표기준값.json"))
    ap.add_argument("--out", default=None)
    a = ap.parse_args()
    out = Path(a.out) if a.out else (HERE / "복구" / "분석보고서.docx" if a.recovery else HERE / "분석보고서.docx")
    tables, images = build(a.ref, out, a.recovery)
    print(f"OK {out.name} · 칸 {len(SECTIONS)} · 표 {tables} · 그림 {images} · 출처 지표기준값.json → {out}")


if __name__ == "__main__":
    main()
