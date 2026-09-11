#!/usr/bin/env python3
"""만들기_발표자료.py — 지표기준값.json만 읽어 발표자료.pptx(강사용 완성본)를 만든다.

    python 만들기_발표자료.py              → 발표자료.pptx        (장마다 발표자 노트)
    python 만들기_발표자료.py --recovery   → 복구/발표자료.pptx   (문제와 데이터 · 핵심 지표 · 한계)
    python 만들기_발표자료.py --out 경로   → 원하는 곳에 저장

- 필요 패키지: python-pptx. 원본 CSV와 이 폴더의 기존 파일은 고치지 않는다.
- 지표 숫자는 전부 지표기준값.json에서 가져온다. 이 파일에 손으로 적은 지표 숫자는 없다.
- 제목에 쓴 결론(예: «매출 1위 상품군의 마진율이 가장 낮다»)은 만들기 전에 데이터로 다시 확인하고,
  맞지 않으면 파일을 만들지 않고 멈춘다.
- 넘침·겹침·18pt 하한은 결과물_대조.py가 추정 검사한다(PowerPoint 실제 렌더는 별도 확인).
"""
import argparse
import json
import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(errors="replace")
except Exception:
    pass

try:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Emu, Inches, Pt
except ImportError:
    sys.exit("FAIL python-pptx가 없다 — 설치가 필요하다(pip install python-pptx). 파일을 만들지 않았다.")

HERE = Path(__file__).resolve().parent
FONT = "Malgun Gothic"
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x4B, 0x55, 0x63)
ACCENT = RGBColor(0x1D, 0x4E, 0xD8)
ACCENT_SOFT = RGBColor(0x93, 0xC5, 0xFD)
WARN = RGBColor(0xB4, 0x53, 0x09)
BAR = RGBColor(0x94, 0xA3, 0xB8)
CARD = RGBColor(0xF3, 0xF4, 0xF6)
LINE = RGBColor(0xD1, 0xD5, 0xDB)

SLIDE_W, SLIDE_H = Emu(12192000), Emu(6858000)   # 16:9 · 13.333 × 7.5 in
MX = 0.6                                        # 좌우 여백(in)
CW = 13.333 - 2 * MX                            # 본문 폭(in)
BODY_TOP = 2.0


# ── 데이터 ──────────────────────────────────────────────────────────
def load(path):
    b = json.loads(Path(path).read_text(encoding="utf-8"))
    for k in ("ROW", "ORD", "REV", "MGN", "MONTH", "CAT", "DELIV", "RATE", "RET", "HOLD"):
        if k not in b:
            sys.exit(f"FAIL 지표기준값.json에 {k} 절이 없다")
    return b


def won(v):
    return f"{int(v):,}원"


def cnt(v):
    return f"{int(v):,}"


def pct(v):
    return f"{v:.1f}%"


def month_name(ym):
    return f"{int(ym.split('-')[1])}월"


def facts(b):
    """슬라이드 제목의 결론을 데이터로 확인하고, 필요한 값을 모은다."""
    cats = b["CAT"]["CAT-02_목록"]
    top_rev = max(cats, key=lambda c: c["순매출_krw"])
    low_rev = min(cats, key=lambda c: c["순매출_krw"])
    low_mgn = min(cats, key=lambda c: c["마진율_pct"])
    high_mgn = max(cats, key=lambda c: c["마진율_pct"])
    top_ord = max(cats, key=lambda c: c["주문수"])
    months, mrev = b["MONTH"]["MONTH-01_월목록"], b["MONTH"]["MONTH-02_월별순매출_krw"]
    imax, imin = mrev.index(max(mrev)), mrev.index(min(mrev))
    hold_dup, hold_del = b["HOLD"]["HOLD-02_중복충돌"], b["HOLD"]["HOLD-03_배송이상"]
    del_days = {h["delivery_days"] for h in hold_del}
    reasons = b["RET"]["RET-03_사유분포"]
    top_reason = max(reasons, key=reasons.get)
    checks = [
        (top_rev is low_mgn, "매출 1위 상품군의 마진율이 가장 낮다"),
        (low_rev is high_mgn, "매출 최하위 상품군의 마진율이 가장 높다"),
        (top_ord is not top_rev, "주문 수 1위와 매출 1위가 다르다"),
        (months[imax] == b["MONTH"]["MONTH-05_최대월"], "최대월 = MONTH-05"),
        (months[imin] == b["MONTH"]["MONTH-06_최소월"], "최소월 = MONTH-06"),
        (len(hold_dup) == b["ROW"]["ROW-05_보류행수"], "중복 충돌 행 수 = ROW-05"),
        (len(hold_del) == b["DELIV"]["DELIV-07_보류제외행수"], "배송 이상 행 수 = DELIV-07"),
        (len(hold_dup) + len(hold_del) == b["HOLD"]["HOLD-01_승인필요건수"], "승인 필요 = 중복 충돌 + 배송 이상"),
        (len(del_days) == 1, "배송 이상 행의 배송일이 하나의 값"),
    ]
    bad = [msg for ok, msg in checks if not ok]
    if bad:
        sys.exit("FAIL 제목·본문의 결론이 데이터와 맞지 않아 만들지 않았다: " + " / ".join(bad))
    return dict(cats=cats, top_rev=top_rev, low_rev=low_rev, top_ord=top_ord, months=months, mrev=mrev,
                imax=imax, imin=imin, del_days=del_days.pop(), reasons=reasons, top_reason=top_reason)


# ── 그리기 도구 ──────────────────────────────────────────────────────
def apply_font(font, size, bold=False, color=INK):
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = color
    font.name = FONT
    rpr = font._rPr
    latin = rpr.find(qn("a:latin"))
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = OxmlElement("a:ea")
        latin.addnext(ea)
    ea.set("typeface", FONT)
    cs = rpr.find(qn("a:cs"))
    if cs is None:
        cs = OxmlElement("a:cs")
        ea.addnext(cs)
    cs.set("typeface", FONT)


def P(text, size=20, bold=False, color=INK, before=0):
    return dict(text=text, size=size, bold=bold, color=color, before=before)


def add_text(slide, name, x, y, w, h, paras, fill=None, pad=0.1):
    if fill is None:
        shp = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        style = shp._element.find(qn("p:style"))
        if style is not None:
            shp._element.remove(style)
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
        shp.line.fill.background()
    shp.name = name
    tf = shp.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Inches(pad)
    tf.margin_top = tf.margin_bottom = Inches(0.08)
    for i, spec in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        if spec["before"]:
            p.space_before = Pt(spec["before"])
        run = p.add_run()
        run.text = spec["text"]
        apply_font(run.font, spec["size"], spec["bold"], spec["color"])
    return shp


def header(slide, kicker, title):
    add_text(slide, "머리말", MX, 0.35, CW, 1.45, [P(kicker, 18, True, ACCENT), P(title, 28, True, INK, 2)])
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(MX), Inches(1.86), Inches(CW), Inches(0.03))
    style = rule._element.find(qn("p:style"))
    if style is not None:
        rule._element.remove(style)
    rule.fill.solid()
    rule.fill.fore_color.rgb = LINE
    rule.line.fill.background()
    rule.name = "장식_구분선"


def column_chart(slide, name, x, y, w, h, categories, series, colors, unit_title):
    data = CategoryChartData(number_format="#,##0")
    data.categories = categories
    for label, values in series:
        data.add_series(label, values)
    frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data)
    frame.name = name
    chart = frame.chart
    apply_font(chart.font, 18, False, MUTED)
    chart.has_title = False          # 계열 하나일 때 PowerPoint가 계열 이름을 제목으로 자동 표시하는 것을 끈다
    chart.has_legend = len(series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    va = chart.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = LINE
    va.tick_labels.number_format = '#,##0,,'
    va.tick_labels.number_format_is_linked = False
    va.format.line.fill.background()
    va.has_title = True
    va.axis_title.text_frame.text = unit_title
    for p in va.axis_title.text_frame.paragraphs:
        for r in p.runs:
            apply_font(r.font, 18, False, MUTED)
    chart.category_axis.format.line.color.rgb = LINE
    plot = chart.plots[0]
    plot.gap_width = 60
    if len(series) > 1:
        plot.overlap = -10
    for s, color in zip(plot.series, colors):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = color
    return chart


def bullets(lines, size=20, color=INK, before=8):
    return [P("• " + t, size, False, color, before) for t in lines]


# ── 슬라이드 ─────────────────────────────────────────────────────────
def s_problem(slide, b, f, recovery):
    R, O = b["ROW"], b["ORD"]
    header(slide, "문제와 데이터",
           f"분석 대상은 원본 {cnt(R['ROW-01_원본행수'])}행을 정리한 주문 {cnt(O['ORD-01_분석대상주문수'])}건입니다")
    add_text(slide, "질문", MX, BODY_TOP, 5.9, 4.9,
             [P("우리가 답하려는 질문", 22, True, ACCENT)] + bullets([
                 "월별 매출은 어떻게 움직였나?",
                 "어느 상품군이 매출과 마진을 만들었나?",
                 "고객 평점과 반품은 어떤 모습인가?",
                 "결론 전에 무엇을 확인해야 하나?",
             ], before=12), fill=CARD, pad=0.25)
    add_text(slide, "데이터", MX + 5.9 + 0.333, BODY_TOP, 5.9, 4.9,
             [P("사용한 데이터", 22, True, ACCENT)] + bullets([
                 f"원본 {cnt(R['ROW-01_원본행수'])}행 · 고유 주문번호 {cnt(R['ROW-02_고유주문ID수'])}건",
                 f"분석 대상 {cnt(R['ROW-03_분석대상행수'])}건 (원본의 {pct(R['ROW-06_채택비율_pct'])})",
                 f"제외 {cnt(R['ROW-04_제외행수'])}행 · 확인이 필요해 보류 {cnt(R['ROW-05_보류행수'])}행",
                 f"기간 {O['ORD-06_기간시작']} ~ {O['ORD-07_기간종료']}",
                 "연습용 모의 데이터 (실제 기록 아님)",
             ], before=12), fill=CARD, pad=0.25)
    tail = ("이어서 핵심 지표와 확인할 부분을 보겠습니다."
            if recovery else "오늘은 핵심 지표, 월별 매출과 상품군별 매출, 눈에 띈 특징, 확인할 부분을 차례로 보겠습니다.")
    return (f"오늘 발표는 한 해 동안의 주문 기록을 정리해서 본 결과입니다. 먼저 이 데이터는 연습용으로 만든 모의 데이터이고, "
            f"실제 거래 기록이 아니라는 점을 말씀드립니다. 원본은 {cnt(R['ROW-01_원본행수'])}행이었지만 같은 주문이 두 번 적힌 행이 있어서 "
            f"고유 주문번호는 {cnt(R['ROW-02_고유주문ID수'])}건이었습니다. 취소나 대기 상태의 주문과 수량·배송일이 불가능한 기록을 빼고 "
            f"{cnt(R['ROW-03_분석대상행수'])}건을 분석했습니다. 원본의 {pct(R['ROW-06_채택비율_pct'])}입니다. "
            f"어느 값이 맞는지 정할 수 없는 {cnt(R['ROW-05_보류행수'])}행은 지우거나 채우지 않고 보류로 남겼습니다. {tail}")


def s_kpi(slide, b, f, recovery):
    REV, O, RATE, MGN, D = b["REV"], b["ORD"], b["RATE"], b["MGN"], b["DELIV"]
    header(slide, "핵심 지표", f"한 해 총매출은 {won(REV['REV-01_순매출_krw'])}입니다")
    cards = [
        ("총매출", won(REV["REV-01_순매출_krw"]), "취소·대기 제외, 반품 미차감"),
        ("주문 수", f"{cnt(O['ORD-01_분석대상주문수'])}건", "중복을 정리한 뒤의 건수"),
        ("평균 평점", f"{RATE['RATE-01_평균']:.2f}점", f"응답 {cnt(RATE['RATE-02_분모행수'])}건 기준"),
        ("마진율", pct(MGN["MGN-03_마진율_pct"]), "원가만 뺀 값, 이익률 아님"),
        ("평균 주문금액", won(REV["REV-06_평균주문금액_krw"]), "주문 한 건당 금액"),
        ("평균 배송일", f"{D['DELIV-01_평균일']:.2f}일", f"배송 기록 {cnt(D['DELIV-05_분모행수'])}건 기준"),
    ]
    cw, gap, ch = (CW - 2 * 0.25) / 3, 0.25, 2.3
    for i, (label, value, sub) in enumerate(cards):
        x = MX + (i % 3) * (cw + gap)
        y = BODY_TOP + (i // 3) * (ch + gap)
        add_text(slide, f"지표카드_{label}", x, y, cw, ch,
                 [P(label, 20, True, MUTED), P(value, 32, True, ACCENT if i < 3 else INK, 6), P(sub, 18, False, MUTED, 8)],
                 fill=CARD, pad=0.22)
    return (f"핵심 지표 여섯 개입니다. 총매출은 {won(REV['REV-01_순매출_krw'])}입니다. 여기서 총매출은 취소와 대기 주문을 빼고, "
            f"반품은 빼지 않은 금액입니다. 주문 수는 {cnt(O['ORD-01_분석대상주문수'])}건으로, 중복을 정리해 한 행이 한 주문이 되게 한 뒤의 숫자입니다. "
            f"평균 평점은 {RATE['RATE-01_평균']:.2f}점인데, 평점을 남기지 않은 {cnt(RATE['RATE-03_무응답행수'])}건은 0점으로 세지 않고 빼서 "
            f"{cnt(RATE['RATE-02_분모행수'])}건을 기준으로 계산했습니다. 마진율 {pct(MGN['MGN-03_마진율_pct'])}는 상품 원가만 뺀 값이라 "
            f"이익률로 읽으면 안 됩니다. 평균 배송일 {D['DELIV-01_평균일']:.2f}일도 배송 기록 {cnt(D['DELIV-05_분모행수'])}건을 기준으로 한 값입니다. "
            f"평균을 말할 때는 몇 건으로 나눈 값인지 함께 말해 주세요.")


def s_month(slide, b, f, recovery):
    M = b["MONTH"]
    header(slide, "핵심 차트 · 월별 매출", "월별 총매출은 한 방향 추세 없이 오르내렸습니다")
    chart = column_chart(slide, "차트_월별총매출", MX, BODY_TOP - 0.05, 8.35, 5.15,
                         [month_name(m) for m in f["months"]], [("총매출(원)", f["mrev"])], [BAR], "단위: 백만원")
    pts = chart.plots[0].series[0].points
    for idx, color in ((f["imax"], ACCENT), (f["imin"], WARN)):
        pts[idx].format.fill.solid()
        pts[idx].format.fill.fore_color.rgb = color
    orders = M["MONTH-03_월별주문수"]
    hi, lo = f["months"][f["imax"]], f["months"][f["imin"]]
    add_text(slide, "월별_요점", MX + 8.35 + 0.25, BODY_TOP, CW - 8.35 - 0.25, 5.1, [
        P("가장 높은 달", 18, True, ACCENT),
        P(f"{month_name(hi)} · {won(f['mrev'][f['imax']])}", 20, True, INK, 2),
        P("가장 낮은 달", 18, True, WARN, 14),
        P(f"{month_name(lo)} · {won(f['mrev'][f['imin']])}", 20, True, INK, 2),
        P("월별 주문 수", 18, True, MUTED, 14),
        P(f"{min(orders)}~{max(orders)}건 사이", 20, False, INK, 2),
        P("한 해 자료뿐이라", 18, False, MUTED, 14),
        P("계절 효과인지 알 수 없음", 18, False, MUTED, 2),
    ], pad=0.1)
    return (f"월별 총매출입니다. 막대가 한 방향으로 오르거나 내리지 않고 달마다 오르내리는 모습입니다. "
            f"가장 높은 달은 {month_name(hi)}로 {won(f['mrev'][f['imax']])}, 가장 낮은 달은 {month_name(lo)}로 "
            f"{won(f['mrev'][f['imin']])}이었습니다. 월별 주문 수도 {min(orders)}건에서 {max(orders)}건 사이에서 오르내렸습니다. "
            f"여기서 조심할 점이 있습니다. 이 데이터는 한 해 자료뿐이어서, {month_name(hi)}이 높았던 것이 해마다 반복되는 계절 효과인지 "
            f"그해에만 생긴 일인지 구분할 수 없습니다. 그래서 {month_name(hi)}을 성수기라고 부르지 않고, 높은 달과 낮은 달이 있었다는 "
            f"사실까지만 전달하겠습니다.")


def s_category(slide, b, f, recovery):
    cats, top, low = f["cats"], f["top_rev"], f["low_rev"]
    header(slide, "핵심 차트 · 상품군별 매출과 마진",
           f"매출 1위 {top['category']}는 마진율이 가장 낮았습니다")
    column_chart(slide, "차트_상품군별매출마진", MX, BODY_TOP - 0.05, 7.85, 5.15,
                 [c["category"] for c in cats],
                 [("총매출(원)", [c["순매출_krw"] for c in cats]), ("마진(원)", [c["마진_krw"] for c in cats])],
                 [ACCENT, ACCENT_SOFT], "단위: 백만원")
    rows = [P("상품군별 마진율", 18, True, MUTED)]
    for c in cats:
        hot = c is top or c is low
        rows.append(P(f"{c['category']}  {pct(c['마진율_pct'])}", 20, hot, ACCENT if c is top else (WARN if c is low else INK), 6))
    rows += [P("마진은 상품 원가만 뺀 값", 18, False, MUTED, 14), P("실제 이익이 아님", 18, False, MUTED, 2)]
    add_text(slide, "상품군_마진율", MX + 7.85 + 0.25, BODY_TOP, CW - 7.85 - 0.25, 5.1, rows, pad=0.1)
    return (f"상품군별 총매출과 마진입니다. 매출이 가장 큰 상품군은 {top['category']}로 전체 매출의 {pct(top['매출비중_pct'])}를 "
            f"차지했습니다. 그런데 마진율은 {pct(top['마진율_pct'])}로 여섯 상품군 가운데 가장 낮았습니다. 반대로 {low['category']}는 "
            f"매출 비중이 {pct(low['매출비중_pct'])}로 가장 작지만 마진율은 {pct(low['마진율_pct'])}로 가장 높았습니다. "
            f"다만 이 데이터에서는 상품마다 단가와 원가가 정해져 있어서, 마진율 차이를 누가 관리를 잘한 결과로 읽을 수는 없습니다. "
            f"그리고 마진은 상품 원가만 뺀 값이라 실제 이익과는 다르다는 점도 함께 말씀드립니다.")


def s_features(slide, b, f, recovery):
    RATE, RET = b["RATE"], b["RET"]
    top, tord = f["top_rev"], f["top_ord"]
    reasons, tr = f["reasons"], f["top_reason"]
    header(slide, "발견한 특징", f"평점은 높은 편이고, 주문 수 1위와 매출 1위는 다른 상품군입니다")
    one = RATE["RATE-05_분포"]["1"]
    cards = [
        ("고객 평점", f"4점 이상 {pct(RATE['RATE-06_4점이상_pct'])}",
         [("평균 평점", f"{RATE['RATE-01_평균']:.2f}점"), ("1점 평가", f"{one}건")], "평점 수집 방식은 확인 전"),
        ("주문 수와 매출", f"{tord['category']} {cnt(tord['주문수'])}건",
         [("매출 1위", top["category"]), (f"{top['category']} 주문 수", f"{cnt(top['주문수'])}건")], "주문 1위와 매출 1위가 다름"),
        ("반품", f"{cnt(RET['RET-01_반품건수'])}건 ({pct(RET['RET-02_반품률_pct'])})",
         [("가장 많은 사유", f"{tr} · {reasons[tr]}건"), ("사유별 건수", f"{min(reasons.values())}~{max(reasons.values())}건 사이")],
         "반품 금액은 매출에 포함"),
    ]
    cw, gap = (CW - 2 * 0.25) / 3, 0.25
    for i, (label, value, pairs, note) in enumerate(cards):
        paras = [P(label, 22, True, MUTED), P(value, 28, True, ACCENT, 8)]
        for k, v in pairs:
            paras += [P(k, 18, True, MUTED, 16), P(v, 20, False, INK, 2)]
        paras.append(P(note, 18, False, MUTED, 16))
        add_text(slide, f"특징카드_{label}", MX + i * (cw + gap), BODY_TOP, cw, 4.9, paras, fill=CARD, pad=0.22)
    return (f"눈에 띈 특징 세 가지입니다. 첫째, 평점을 남긴 주문 가운데 4점 이상이 {pct(RATE['RATE-06_4점이상_pct'])}였고 "
            f"1점은 한 건도 없었습니다. 다만 1점이 없다는 것이 실제로 불만이 없었다는 뜻인지, 평점을 모으는 방식의 영향인지는 "
            f"이 데이터로 알 수 없습니다. 둘째, 주문 수가 가장 많은 상품군은 {tord['category']}로 {cnt(tord['주문수'])}건이었고, "
            f"매출 1위인 {top['category']}는 {cnt(top['주문수'])}건이었습니다. 많이 팔린 것과 매출이 큰 것은 다를 수 있습니다. "
            f"셋째, 반품은 {cnt(RET['RET-01_반품건수'])}건으로 {pct(RET['RET-02_반품률_pct'])}였고, 사유는 {tr}가 {reasons[tr]}건으로 "
            f"가장 많았습니다. 사유별 건수는 {min(reasons.values())}건에서 {max(reasons.values())}건 사이였습니다.")


def s_limits(slide, b, f, recovery):
    R, REV, RATE, H, D = b["ROW"], b["REV"], b["RATE"], b["HOLD"], b["DELIV"]
    header(slide, "확인이 필요한 부분과 한계", "결론으로 쓰기 전에 사람이 정할 것과 데이터의 한계가 있습니다")
    add_text(slide, "확인필요", MX, BODY_TOP, 5.9, 4.9,
             [P(f"사람이 정할 것 (승인 필요 {H['HOLD-01_승인필요건수']}건)", 22, True, WARN)] + bullets([
                 f"같은 주문번호에 수량이 다른 기록 {R['ROW-05_보류행수']}행",
                 f"배송 {f['del_days']}일로 적힌 기록 {D['DELIV-07_보류제외행수']}행",
                 f"배송비 {won(REV['REV-05_배송비합계_krw'])}을 매출에 넣을지",
                 "반품을 매출에서 뺄지",
             ], before=12) + [P(f"   빼면 총매출 {won(REV['REV-07_반품제외순매출_krw'])}", 20, False, MUTED, 2)],
             fill=CARD, pad=0.25)
    add_text(slide, "한계", MX + 5.9 + 0.333, BODY_TOP, 5.9, 4.9,
             [P("데이터의 한계", 22, True, ACCENT)] + bullets([
                 "연습용 모의 데이터 — 판단 근거 아님",
                 "마진은 원가만 뺀 값 — 이익이 아님",
                 "한 해 자료뿐 — 계절 효과 판단 불가",
                 "함께 움직여도 원인은 알 수 없음",
                 f"평점 무응답 {RATE['RATE-03_무응답행수']}건은 0점으로 세지 않음",
             ], before=12), fill=CARD, pad=0.25)
    return (f"마지막으로 이 결과를 결론으로 쓰기 전에 확인할 것들입니다. 먼저 사람이 정해야 할 기록이 {H['HOLD-01_승인필요건수']}건 있습니다. "
            f"같은 주문번호인데 수량이 다르게 적힌 {R['ROW-05_보류행수']}행은 어느 쪽이 맞는지 알 수 없어 보류했고, "
            f"배송이 {f['del_days']}일로 적힌 {D['DELIV-07_보류제외행수']}행은 매출에는 넣고 배송 평균에서만 뺐습니다. "
            f"배송비를 매출에 넣을지, 반품을 뺄지도 정해야 합니다. 반품을 빼면 총매출은 {won(REV['REV-07_반품제외순매출_krw'])}이 됩니다. "
            f"한계도 분명합니다. 연습용 모의 데이터이고, 한 해 자료뿐이며, 마진은 이익이 아닙니다. "
            f"두 숫자가 함께 움직이는 것을 보더라도, 그것만으로 원인을 말할 수는 없습니다.")


FULL = [s_problem, s_kpi, s_month, s_category, s_features, s_limits]
RECOVERY = [s_problem, s_kpi, s_limits]


def build(ref_path, out_path, recovery):
    b = load(ref_path)
    f = facts(b)
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    prs.core_properties.title = "2025년 주문 데이터 분석 발표" + (" (복구본)" if recovery else "")
    prs.core_properties.author = "AI 코딩 에이전트 입문 3차시 강사용 예시"
    blank = prs.slide_layouts[6]
    for maker in (RECOVERY if recovery else FULL):
        slide = prs.slides.add_slide(blank)
        notes = maker(slide, b, f, recovery)
        slide.notes_slide.notes_text_frame.text = notes
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    return len(prs.slides)


def main():
    ap = argparse.ArgumentParser(description="지표기준값.json → 발표자료.pptx")
    ap.add_argument("--recovery", action="store_true", help="복구본(3장)을 만든다")
    ap.add_argument("--ref", default=str(HERE / "지표기준값.json"))
    ap.add_argument("--out", default=None)
    a = ap.parse_args()
    out = Path(a.out) if a.out else (HERE / "복구" / "발표자료.pptx" if a.recovery else HERE / "발표자료.pptx")
    n = build(a.ref, out, a.recovery)
    print(f"OK {out.name} · 슬라이드 {n}장 · 장마다 발표자 노트 · 출처 지표기준값.json → {out}")


if __name__ == "__main__":
    main()
