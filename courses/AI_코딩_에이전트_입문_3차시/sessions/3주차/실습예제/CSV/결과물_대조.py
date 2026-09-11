#!/usr/bin/env python3
"""결과물_대조.py — 결과물의 숫자가 지표기준값.json과 같은지, 발표자료 레이아웃이 넘치지 않는지 기계로 판정한다.

    python 결과물_대조.py               # 이 폴더의 결과물을 판정
    python 결과물_대조.py --detail      # 텍스트 상자마다 추정 높이/가용 높이도 출력
    python 결과물_대조.py --root 폴더   # 다른 폴더(사본)를 판정 — 음성 시험용

대상과 읽는 곳
  발표자료.pptx · 복구/발표자료.pptx      슬라이드 텍스트 · 표 · 차트(제목·축 제목·항목·계열 이름·값) · 발표자 노트
  분석보고서.docx · 복구/분석보고서.docx  본문·표·머리글·바닥글 문단 · 그림 대체텍스트 · 그림 라벨 목록(PNG 'labels')
  분석결과_요약.md                        모든 줄
  dashboard.html                          화면 정적 텍스트 · baseline-data(JSON 전체 비교) · rows-data(재집계)

판정
  ① 핵심값 — 총매출(REV-01) · 주문 수(ORD-01) · 평균 평점(RATE-01)이 파일마다 있는가
  ② 숫자 — 모든 숫자 토큰이 기준값에 있는가. 표기 변형(쉼표·%·원·일·점·건·5.0=5)은 같은 수로 본다.
     지표 ID(REV-01)·코드(E1·ORD00852)는 숫자로 세지 않는다(식별자). 비지표 숫자는 ALLOW 목록으로만 뺀다.
  ③ 레이아웃(PPTX) — 텍스트 넘침 추정(글자 폭×pt로 줄 수 계산 → 줄 수×행 높이 ≤ 상자 높이) · 슬라이드 밖 ·
     도형 겹침('장식'으로 시작하는 이름만 제외) · 글자 18pt 하한 · 글꼴 Malgun Gothic · 16:9
  ④ 구성·보조 — 칸/장 순서 · 장 수 · 표·그림·차트 개수 · 장마다 발표자 노트(1분 분량 글자 수) ·
     신규 결과물의 원인 단정 표현(때문·덕분·로 인해·탓에)

한계(판정하지 않는 것 — 실행할 때마다 «범위 밖» 줄로 개수를 밝힌다)
  - 다른 지표와 우연히 같은 값은 통과한다. 숫자와 라벨의 대응(«마진율 53.3%»가 틀렸는지)은 보지 않는다.
  - 넘침은 추정이다. PowerPoint·Word의 실제 렌더는 사람이 열어 확인해야 한다.
  - dashboard.html의 화면 값은 브라우저에서 JS가 그린다 — 정적으로는 데이터 블록만 판정하고, 화면은 자기검증 배너가 맡는다.

종료코드: 0 = 위반 0 · 미판정 0 · 입력 전부 있음 / 1 = 그 밖 전부(«대상 계수 0»은 통과가 아니다)
"""
import argparse
import html
import io
import json
import re
import sys
from collections import Counter
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

try:
    sys.stdout.reconfigure(errors="replace")
except Exception:
    pass

HERE = Path(__file__).resolve().parent
REF_NAME = "지표기준값.json"
FIVE = ["핵심 요약", "주요 지표", "발견한 특징", "개선이 필요한 부분", "데이터의 한계"]
TARGETS = [
    dict(path="발표자료.pptx", kind="pptx", new=True, slides_max=6, charts_min=2,
         sections=["문제와 데이터", "핵심 지표", "핵심 차트", "발견한 특징", "한계"]),
    dict(path="복구/발표자료.pptx", kind="pptx", new=True, slides_exact=3, charts_min=0,
         sections=["문제와 데이터", "핵심 지표", "한계"]),
    dict(path="분석보고서.docx", kind="docx", new=True, sections=FIVE, tables_min=1, images_min=1),
    dict(path="복구/분석보고서.docx", kind="docx", new=True, sections=FIVE, tables_min=0, images_min=0),
    dict(path="분석결과_요약.md", kind="md", new=True, sections=FIVE, max_lines=40),
    dict(path="dashboard.html", kind="html", new=False),
]
CORE = [("REV", "REV-01_순매출_krw", "총매출"), ("ORD", "ORD-01_분석대상주문수", "주문 수"), ("RATE", "RATE-01_평균", "평균 평점")]
FONT = "Malgun Gothic"
MIN_PT = 18
NOTE_MIN, NOTE_MAX = 200, 450      # 1분 분량 추정(말하기 분당 약 250~350자) — 공백 제외 글자 수
LINE_FACTOR = 1.30                 # 줄 간격 1.0일 때 행 높이 = pt × 1.30 (Malgun Gothic · 보수적)
CAUSAL_RE = re.compile(r"때문|덕분|로 인해|탓에")

# 비지표 숫자 허용 목록 — (이름, 조건(토큰, 앞 글자, 뒤 글자), 근거). 여기에 없는 숫자는 기준값에 있어야 한다.
ALLOW = [
    ("연도", lambda t, pre, post: t == "2025" and post.startswith("년"), "데이터 기간의 연도(ORD-06·ORD-07 날짜의 일부)"),
    ("과정 차시", lambda t, pre, post: post.startswith("차시"), "과정 회차 표기"),
    ("대시보드 요소 번호", lambda t, pre, post: t in set("12345678") and pre.rstrip().endswith("요소"), "커리큘럼 대시보드 필수 요소 1~8"),
    ("원본 열 수", lambda t, pre, post: t == "20" and post.startswith("열"), "원본 CSV 20열 — 지표가 아닌 구조 사실"),
    ("반올림 자리", lambda t, pre, post: t in ("1", "2") and pre.rstrip().endswith("소수") and post.startswith("자리"), "반올림 규칙 문구"),
    ("반올림 횟수", lambda t, pre, post: t == "1" and post.startswith("회"), "«1회만 반올림» 규칙 문구"),
    ("순위", lambda t, pre, post: t == "1" and post.startswith("위"), "순위 표기 «1위»"),
]

HEX_RE = re.compile(r"(?<![0-9A-Za-z])[0-9a-f]{16,}(?![0-9A-Za-z])")
DATE_RE = re.compile(r"(?<![0-9])\d{4}-\d{2}(?:-\d{2})?(?![0-9])")
ID_RE = re.compile(r"(?<![A-Za-z0-9])[A-Z]{2,6}-\d{2}(?:\s*~\s*\d{2})?(?![0-9])")
CODE_RE = re.compile(r"(?<![0-9A-Za-z])[A-Za-z]+\d+[0-9A-Za-z]*")
NUM_RE = re.compile(r"(?<![0-9A-Za-z.,])(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?(?![0-9])")


def dec(s):
    return Decimal(str(s).replace(",", ""))


def scan(text):
    """→ (날짜 토큰, 숫자 토큰, 식별자 수, 읽지 못한 숫자 덩어리). 위치는 원문 기준."""
    buf = list(text)

    def mask(m):
        for i in range(m.start(), m.end()):
            buf[i] = " "

    ids = 0
    for m in HEX_RE.finditer(text):
        mask(m)
        ids += 1
    dates = list(DATE_RE.finditer("".join(buf)))
    for m in dates:
        mask(m)
    for rx in (ID_RE, CODE_RE):
        for m in list(rx.finditer("".join(buf))):
            mask(m)
            ids += 1
    masked = "".join(buf)
    nums = list(NUM_RE.finditer(masked))
    covered = [False] * len(masked)
    for m in nums:
        for i in range(m.start(), m.end()):
            covered[i] = True
    residual = [m for m in re.finditer(r"\d+", masked) if not all(covered[m.start():m.end()])]
    return dates, nums, ids, residual


class Ref:
    def __init__(self, data):
        self.data = data
        self.nums, self.dates = set(), set()
        for k, v in data.items():
            if k != "_meta":
                self._walk(v)
        self.months = {int(d[5:7]) for d in self.dates if len(d) == 7}

    def _walk(self, o):
        if isinstance(o, bool) or o is None:
            return
        if isinstance(o, (int, float)):
            self.nums.add(dec(o))
        elif isinstance(o, str):
            dates, nums, _, _ = scan(o)
            self.dates.update(m.group(0) for m in dates)
            self.nums.update(dec(m.group(0)) for m in nums)
        elif isinstance(o, dict):
            for k, v in o.items():
                if isinstance(k, str) and k.isdigit():
                    self.nums.add(Decimal(k))
                self._walk(v)
        elif isinstance(o, list):
            for x in o:
                self._walk(x)


class Tally:
    def __init__(self, name, new):
        self.name, self.new = name, new
        self.c = Counter()
        self.fails, self.unjudged, self.detail, self.texts = [], [], [], []
        self.values = set()
        self.allow = Counter()

    def ok(self, cat, n=1):
        self.c[cat + "_판정"] += n

    def bad(self, cat, msg):
        self.c[cat + "_판정"] += 1
        self.c[cat + "_위반"] += 1
        self.fails.append(f"{cat}: {msg}")

    def unk(self, cat, msg):
        self.c["미판정"] += 1
        self.unjudged.append(f"{cat}: {msg}")


def snippet(text, s, e):
    return re.sub(r"\s+", " ", text[max(0, s - 14):e + 10]).strip()


def judge_text(t, text, where, ref):
    if not text or not text.strip():
        return
    t.texts.append(text)
    dates, nums, ids, residual = scan(text)
    t.c["식별자"] += ids
    for m in dates:
        if m.group(0) in ref.dates:
            t.ok("숫자")
            t.c["숫자_일치"] += 1
        else:
            t.bad("숫자", f"{where} 기준에 없는 날짜 {m.group(0)!r} «{snippet(text, m.start(), m.end())}»")
    for m in nums:
        tok = m.group(0)
        pre, post = text[max(0, m.start() - 6):m.start()], text[m.end():m.end() + 3]
        rule = next((name for name, fn, _ in ALLOW if fn(tok, pre, post)), None)
        if rule:
            t.ok("숫자")
            t.c["숫자_허용"] += 1
            t.allow[rule] += 1
            continue
        v = dec(tok)
        if post.startswith("월"):
            good = v == v.to_integral_value() and int(v) in ref.months
        else:
            good = v in ref.nums
        if good:
            t.ok("숫자")
            t.c["숫자_일치"] += 1
            t.values.add(v)
        else:
            t.bad("숫자", f"{where} 기준값에 없는 숫자 {tok!r} «{snippet(text, m.start(), m.end())}»")
    for m in residual:
        t.unk("숫자", f"{where} 숫자 형식을 읽지 못함 «{snippet(text, m.start(), m.end())}»")


def judge_value(t, value, where, ref):
    if value is None:
        t.unk("숫자", f"{where} 값이 비어 있음")
        return
    try:
        v = dec(repr(float(value)) if isinstance(value, float) else value)
    except (InvalidOperation, ValueError):
        t.unk("숫자", f"{where} 값을 수로 읽지 못함 {value!r}")
        return
    if v in ref.nums:
        t.ok("숫자")
        t.c["숫자_일치"] += 1
        t.values.add(v)
    else:
        t.bad("숫자", f"{where} 기준값에 없는 차트 값 {value!r}")


def judge_order(t, found, labels, what):
    """found: 라벨 → 첫 위치(없으면 None)."""
    pos = []
    for label in labels:
        if found.get(label) is None:
            t.bad("구성", f"{what} «{label}» 없음")
        else:
            t.ok("구성")
            pos.append(found[label])
    if len(pos) == len(labels):
        if pos == sorted(pos):
            t.ok("구성")
        else:
            t.bad("구성", f"{what} 순서가 {' → '.join(labels)}가 아님")


# ── PPTX ─────────────────────────────────────────────────────────────
def char_em(ch):
    if ch == " ":
        return 0.30
    if ch in ",.;:'\"|!":
        return 0.35
    if ch in "()[]":
        return 0.42
    if ch.isdigit():
        return 0.60
    if "A" <= ch <= "Z":
        return 0.70
    if "a" <= ch <= "z":
        return 0.56
    if ch == "%":
        return 0.90
    if ch in "~-+=/_*#&@<>":
        return 0.60
    return 1.0   # 한글·CJK·전각 기호(·—→÷ 등)는 보수적으로 1em


def text_width(s, size):
    return sum(char_em(c) for c in s) * size


def wrap_lines(text, size, avail):
    total = 0
    for seg in re.split(r"[\v\n]", text):
        lines, cur = 1, 0.0
        for word in seg.split(" "):
            ww = text_width(word, size)
            sp = char_em(" ") * size if cur > 0 else 0.0
            if cur + sp + ww <= avail:
                cur += sp + ww
            elif ww <= avail:
                lines, cur = lines + 1, ww
            else:
                if cur > 0:
                    lines, cur = lines + 1, 0.0
                for ch in word:
                    cw = char_em(ch) * size
                    if cur + cw > avail:
                        lines, cur = lines + 1, cw
                    else:
                        cur += cw
        total += lines
    return total


def frame_height(tf, width_emu, height_emu):
    """→ (추정 높이 pt, 가용 높이 pt, 가로 넘침 메시지 또는 None) · 크기를 모르면 None."""
    from pptx.util import Length
    avail_w = (width_emu - tf.margin_left - tf.margin_right) / 12700
    avail_h = (height_emu - tf.margin_top - tf.margin_bottom) / 12700
    wrap = tf.word_wrap is not False
    total, prev, wide = 0.0, None, None
    for p in tf.paragraphs:
        sizes = [r.font.size.pt for r in p.runs if r.font.size is not None]
        if p.font.size is not None:
            sizes.append(p.font.size.pt)
        if any(r.text.strip() and r.font.size is None for r in p.runs) and p.font.size is None:
            return None
        size = max(sizes) if sizes else prev
        if size is None:
            return None
        prev = size
        if wrap:
            lines = wrap_lines(p.text, size, avail_w)
        else:
            segs = re.split(r"[\v\n]", p.text)
            lines = len(segs)
            widest = max(text_width(s, size) for s in segs)
            if widest > avail_w:
                wide = f"줄바꿈 꺼짐 · 가로 {widest:.0f}pt > {avail_w:.0f}pt"
        ls = p.line_spacing
        if ls is None:
            lh = size * LINE_FACTOR
        elif isinstance(ls, Length):
            lh = ls.pt
        else:
            lh = size * LINE_FACTOR * float(ls)
        total += lines * lh
        total += p.space_before.pt if p.space_before is not None else 0
        total += p.space_after.pt if p.space_after is not None else 0
    return total, avail_h, wide


def judge_frame(t, tf, width, height, where, ref, detail):
    from pptx.oxml.ns import qn
    for p in tf.paragraphs:
        judge_text(t, p.text, where, ref)
        for r in p.runs:
            if not r.text.strip():
                continue
            size = r.font.size or p.font.size
            if size is None:
                t.unk("레이아웃", f"{where} 글자 크기 상속 — 18pt 하한 미판정 «{r.text[:16]}»")
            elif size.pt < MIN_PT:
                t.bad("레이아웃", f"{where} {size.pt:g}pt < {MIN_PT}pt «{r.text[:16]}»")
            else:
                t.ok("레이아웃")
            ea = r.font._rPr.find(qn("a:ea"))
            ea_face = ea.get("typeface") if ea is not None else None
            if r.font.name == FONT and ea_face == FONT:
                t.ok("레이아웃")
            else:
                t.bad("레이아웃", f"{where} 글꼴 latin={r.font.name} ea={ea_face} «{r.text[:16]}»")
    if not tf.text.strip():
        return
    est = frame_height(tf, width, height)
    if est is None:
        t.unk("레이아웃", f"{where} 글자 크기를 알 수 없어 넘침 미판정")
        return
    need, avail, wide = est
    detail.append(f"    {where}: 추정 {need:.0f}pt / 가용 {avail:.0f}pt (여유 {avail - need:.0f})")
    if need > avail:
        t.bad("레이아웃", f"{where} 넘침 추정 {need:.0f}pt > 가용 {avail:.0f}pt")
    elif wide:
        t.bad("레이아웃", f"{where} {wide}")
    else:
        t.ok("레이아웃")


def judge_chart(t, chart, where, ref):
    from pptx.oxml.ns import qn
    if chart.has_title and chart.chart_title.has_text_frame:
        judge_text(t, chart.chart_title.text_frame.text, f"{where} 차트 제목", ref)
    for axis_name in ("category_axis", "value_axis"):
        try:
            axis = getattr(chart, axis_name)
        except (ValueError, NotImplementedError, AttributeError):
            continue
        if axis.has_title and axis.axis_title.has_text_frame:
            judge_text(t, axis.axis_title.text_frame.text, f"{where} 축 제목", ref)
    for plot in chart.plots:
        for c in plot.categories:
            judge_text(t, str(c), f"{where} 차트 항목", ref)
        for s in plot.series:
            judge_text(t, s.name, f"{where} 계열 이름", ref)
            for i, v in enumerate(s.values):
                judge_value(t, v, f"{where} {s.name}[{i}]", ref)
    sizes = [el.get("sz") for el in chart._chartSpace.iter() if el.tag in (qn("a:defRPr"), qn("a:rPr")) and el.get("sz")]
    if not sizes:
        t.unk("레이아웃", f"{where} 차트 글자 크기 미지정 — 18pt 하한 미판정")
    for sz in sizes:
        if int(sz) / 100 < MIN_PT:
            t.bad("레이아웃", f"{where} 차트 글자 {int(sz) / 100:g}pt < {MIN_PT}pt")
        else:
            t.ok("레이아웃")


def judge_pptx(path, cfg, ref, t):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(str(path))
    SW, SH = prs.slide_width, prs.slide_height
    slides = list(prs.slides)
    n = len(slides)
    if "slides_exact" in cfg:
        (t.ok("구성") if n == cfg["slides_exact"] else t.bad("구성", f"슬라이드 {n}장 ≠ {cfg['slides_exact']}장"))
    if "slides_max" in cfg:
        (t.ok("구성") if 1 <= n <= cfg["slides_max"] else t.bad("구성", f"슬라이드 {n}장 — {cfg['slides_max']}장 이내가 아님"))
    (t.ok("레이아웃") if abs(SW / SH - 16 / 9) < 0.01 else t.bad("레이아웃", f"화면비 {SW}×{SH}가 16:9가 아님"))
    heads, charts = [], 0
    for si, slide in enumerate(slides, 1):
        shapes = list(slide.shapes)
        boxes = []
        text_shapes = [s for s in shapes if s.has_text_frame and s.text_frame.text.strip() and s.top is not None]
        heads.append(min(text_shapes, key=lambda s: (s.top, s.left)).text_frame.text if text_shapes else "")
        for shp in shapes:
            where = f"슬라이드 {si} · {shp.name}"
            if shp.shape_type == MSO_SHAPE_TYPE.GROUP:
                t.unk("레이아웃", f"{where} 그룹 도형 — 좌표를 풀지 않아 넘침·겹침·숫자 미판정")
                continue
            if None in (shp.left, shp.top, shp.width, shp.height):
                t.unk("레이아웃", f"{where} 위치 상속(자리표시자) — 범위·겹침 미판정")
            else:
                inside = shp.left >= 0 and shp.top >= 0 and shp.left + shp.width <= SW and shp.top + shp.height <= SH
                (t.ok("레이아웃") if inside else t.bad("레이아웃", f"{where} 슬라이드 밖으로 나감"))
                if shp.name.startswith("장식"):
                    t.c["장식제외"] += 1
                else:
                    boxes.append((shp.name, shp.left, shp.top, shp.left + shp.width, shp.top + shp.height))
            if shp.has_text_frame and shp.width is not None:
                judge_frame(t, shp.text_frame, shp.width, shp.height, where, ref, t.detail)
            if getattr(shp, "has_table", False) and shp.has_table:
                tbl = shp.table
                for ri, row in enumerate(tbl.rows):
                    for ci, cell in enumerate(row.cells):
                        judge_frame(t, cell.text_frame, tbl.columns[ci].width, row.height,
                                    f"{where} 표[{ri},{ci}]", ref, t.detail)
            if getattr(shp, "has_chart", False) and shp.has_chart:
                charts += 1
                judge_chart(t, shp.chart, where, ref)
            if shp.shape_type == MSO_SHAPE_TYPE.PICTURE:
                t.unk("숫자", f"{where} 그림 — 픽셀 속 숫자 미판정")
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                a, b = boxes[i], boxes[j]
                if min(a[3], b[3]) - max(a[1], b[1]) > 0 and min(a[4], b[4]) - max(a[2], b[2]) > 0:
                    t.bad("레이아웃", f"슬라이드 {si} 겹침: {a[0]} ↔ {b[0]}")
                else:
                    t.ok("레이아웃")
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide and slide.notes_slide.notes_text_frame else ""
        k = len(re.sub(r"\s", "", notes))
        if k == 0:
            t.bad("노트", f"슬라이드 {si} 발표자 노트 없음")
        elif not NOTE_MIN <= k <= NOTE_MAX:
            t.bad("노트", f"슬라이드 {si} 노트 {k}자 — 1분 분량 추정 {NOTE_MIN}~{NOTE_MAX}자 밖")
        else:
            t.ok("노트")
        judge_text(t, notes, f"슬라이드 {si} 노트", ref)
    judge_order(t, {lab: next((i for i, h in enumerate(heads) if lab in h), None) for lab in cfg["sections"]},
                cfg["sections"], "장")
    (t.ok("구성") if charts >= cfg.get("charts_min", 0) else t.bad("구성", f"차트 {charts}개 < {cfg['charts_min']}개"))


# ── DOCX ─────────────────────────────────────────────────────────────
W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
WP_NS = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
DOCX_XML = ("document.main+xml", "header+xml", "footer+xml", "footnotes+xml", "endnotes+xml", "comments+xml")


def para_text(p):
    out = []
    for el in p.iter(W_NS + "t", W_NS + "tab", W_NS + "br"):
        anc = el.getparent()
        while anc is not None and anc.tag != W_NS + "p":
            anc = anc.getparent()
        if anc is p:
            out.append((el.text or "") if el.tag == W_NS + "t" else " ")
    return "".join(out)


def judge_docx(path, cfg, ref, t):
    from docx import Document
    from lxml import etree
    doc = Document(str(path))
    parts = list(doc.part.package.iter_parts())
    blips = 0
    for part in parts:
        if not part.content_type.endswith(DOCX_XML):
            continue
        root = etree.fromstring(part.blob)
        for i, p in enumerate(root.iter(W_NS + "p"), 1):
            judge_text(t, para_text(p), f"{part.partname} 문단 {i}", ref)
        for dp in root.iter(WP_NS + "docPr"):
            judge_text(t, dp.get("descr") or "", f"{part.partname} 그림 대체텍스트", ref)
        if part.content_type.endswith("document.main+xml"):
            blips = sum(1 for _ in root.iter(A_NS + "blip"))
    body_imgs = {rel.target_part.partname for part in parts if part.content_type.endswith(DOCX_XML)
                 for rel in part.rels.values() if rel.reltype.endswith("/image") and not rel.is_external}
    for part in parts:
        if not part.content_type.startswith("image/"):
            continue
        if part.partname not in body_imgs:   # 문서 속성 썸네일 등 — 본문에 보이지 않는다(범위 밖 줄에 개수 공개)
            t.c["범위밖_속성그림"] += 1
            continue
        labels = None
        try:
            from PIL import Image
            labels = Image.open(io.BytesIO(part.blob)).info.get("labels")
        except Exception:
            labels = None
        if labels is None:
            t.unk("숫자", f"{part.partname} 그림 — 라벨 목록이 없어 그림 속 숫자 미판정")
            continue
        for s in json.loads(labels):
            judge_text(t, s, f"{part.partname} 그림 라벨", ref)
    heads = [p.text.strip() for p in doc.paragraphs if p.style is not None and p.style.name.startswith("Heading")]
    has_title = any(p.style is not None and p.style.name == "Title" and p.text.strip() for p in doc.paragraphs)
    (t.ok("구성") if has_title else t.bad("구성", "제목(Title) 문단 없음"))
    judge_order(t, {lab: (heads.index(lab) if lab in heads else None) for lab in cfg["sections"]}, cfg["sections"], "칸")
    (t.ok("구성") if len(doc.tables) >= cfg["tables_min"] else t.bad("구성", f"표 {len(doc.tables)}개 < {cfg['tables_min']}개"))
    (t.ok("구성") if blips >= cfg["images_min"] else t.bad("구성", f"그림 {blips}개 < {cfg['images_min']}개"))


# ── MD ───────────────────────────────────────────────────────────────
def judge_md(path, cfg, ref, t):
    lines = path.read_text(encoding="utf-8").splitlines()
    for i, line in enumerate(lines, 1):
        judge_text(t, line, f"{i}행", ref)
    heads = [ln[3:].strip() for ln in lines if ln.startswith("## ")]
    judge_order(t, {lab: (heads.index(lab) if lab in heads else None) for lab in cfg["sections"]}, cfg["sections"], "칸")
    (t.ok("구성") if len(lines) <= cfg["max_lines"] else t.bad("구성", f"{len(lines)}줄 > {cfg['max_lines']}줄"))


# ── HTML ─────────────────────────────────────────────────────────────
def diff_json(a, b, path=""):
    """(비교한 수 값 개수, 다른 경로 목록)"""
    if isinstance(b, dict):
        if not isinstance(a, dict):
            return 1, [path or "/"]
        n, bad = 0, []
        for k in set(a) | set(b):
            if k not in a or k not in b:
                n, bad = n + 1, bad + [f"{path}/{k}"]
            else:
                dn, db = diff_json(a[k], b[k], f"{path}/{k}")
                n, bad = n + dn, bad + db
        return n, bad
    if isinstance(b, list):
        if not isinstance(a, list) or len(a) != len(b):
            return 1, [path]
        n, bad = 0, []
        for i, (x, y) in enumerate(zip(a, b)):
            dn, db = diff_json(x, y, f"{path}[{i}]")
            n, bad = n + dn, bad + db
        return n, bad
    return 1, ([] if a == b and type(a) is type(b) else [path])


def judge_html(path, cfg, ref, t):
    s = path.read_text(encoding="utf-8")
    base_m = re.search(r'<script[^>]*id="baseline-data"[^>]*>(.*?)</script>', s, re.S)
    rows_m = re.search(r'<script[^>]*id="rows-data"[^>]*>(.*?)</script>', s, re.S)
    if base_m is None:
        t.bad("대시보드데이터", "baseline-data 블록 없음")
    else:
        base = json.loads(base_m.group(1))
        n, bad = diff_json(base, ref.data)
        for p in bad:
            t.bad("대시보드데이터", f"baseline-data가 지표기준값.json과 다름: {p}")
        t.ok("대시보드데이터", n - len(bad))
        for sec, key, _ in CORE:
            if base.get(sec, {}).get(key) == ref.data[sec][key]:
                t.values.add(dec(ref.data[sec][key]))
    if rows_m is None:
        t.bad("대시보드데이터", "rows-data 블록 없음")
    else:
        pack = json.loads(rows_m.group(1))
        cols, rows = pack.get("cols", []), pack.get("rows", [])
        need = [c for c in ("revenue", "margin", "rating") if c not in cols]
        if need or not rows:
            t.bad("대시보드데이터", f"rows-data 열 없음 {need} 또는 행 0")
        else:
            ci = {c: i for i, c in enumerate(cols)}
            rated = [r[ci["rating"]] for r in rows if r[ci["rating"]] not in (None, "")]
            avg = (sum(dec(x) for x in rated) / len(rated)).quantize(Decimal("0.01"), ROUND_HALF_UP) if rated else None
            checks = [
                ("행 수", Decimal(len(rows)), ref.data["ORD"]["ORD-01_분석대상주문수"]),
                ("매출 합", sum(dec(r[ci["revenue"]]) for r in rows), ref.data["REV"]["REV-01_순매출_krw"]),
                ("마진 합", sum(dec(r[ci["margin"]]) for r in rows), ref.data["MGN"]["MGN-02_마진_krw"]),
                ("평점 분모", Decimal(len(rated)), ref.data["RATE"]["RATE-02_분모행수"]),
                ("평균 평점", avg, ref.data["RATE"]["RATE-01_평균"]),
            ]
            for label, got, want in checks:
                if got is not None and got == dec(want):
                    t.ok("대시보드데이터")
                    t.values.add(dec(want))
                else:
                    t.bad("대시보드데이터", f"rows-data 재집계 {label} {got} ≠ 기준 {want}")
    code = "".join(m.group(1) for m in re.finditer(r"<script(?![^>]*application/json)[^>]*>(.*?)</script>", s, re.S))
    t.c["범위밖_스크립트숫자"] = len(NUM_RE.findall(code))
    body = re.sub(r"<script\b.*?</script>|<style\b.*?</style>|<!--.*?-->", " ", s, flags=re.S)
    text = re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", body)))
    judge_text(t, text, "화면 정적 텍스트", ref)


# ── 실행 ─────────────────────────────────────────────────────────────
JUDGES = dict(pptx=judge_pptx, docx=judge_docx, md=judge_md, html=judge_html)
CATS = ["숫자", "핵심값", "구성", "레이아웃", "노트", "원인단정", "대시보드데이터"]


def run_target(root, cfg, ref):
    t = Tally(cfg["path"], cfg["new"])
    path = root / cfg["path"]
    if not path.is_file():
        t.bad("입력", "파일 없음")
        return t
    if path.stat().st_size == 0:
        t.bad("입력", "파일이 비어 있음")
        return t
    try:
        JUDGES[cfg["kind"]](path, cfg, ref, t)
    except Exception as e:  # 읽지 못한 파일은 통과가 아니다
        t.bad("입력", f"읽기 실패 {type(e).__name__}: {e}")
        return t
    if t.c["숫자_판정"] == 0:
        t.bad("숫자", "숫자 토큰 0 — 대상 계수 0은 통과가 아니다")
    for sec, key, label in CORE:
        v = dec(ref.data[sec][key])
        (t.ok("핵심값") if v in t.values else t.bad("핵심값", f"{label} {ref.data[sec][key]} 없음"))
    if t.new:
        hits = [(i, m) for i, txt in enumerate(t.texts) for m in CAUSAL_RE.finditer(txt)]
        for i, m in hits:
            t.bad("원인단정", f"«{snippet(t.texts[i], m.start(), m.end())}»")
        if not hits:
            t.ok("원인단정")
    return t


def summary(t):
    parts = []
    for cat in ["입력"] + CATS:
        judged, bad = t.c[cat + "_판정"], t.c[cat + "_위반"]
        if not judged:
            continue
        if cat == "숫자":
            parts.append(f"숫자 {judged} (일치 {t.c['숫자_일치']} · 허용 {t.c['숫자_허용']} · 위반 {bad})")
        elif cat == "핵심값":
            parts.append(f"핵심값 {judged - bad}/{judged}")
        else:
            parts.append(f"{cat} {judged} (위반 {bad})")
    return f"[{t.name}] " + " · ".join(parts) + f" · 미판정 {t.c['미판정']}"


def main():
    ap = argparse.ArgumentParser(description="결과물 숫자 = 지표기준값.json 대조 · PPTX 넘침 추정")
    ap.add_argument("--root", default=str(HERE), help="결과물과 지표기준값.json이 있는 폴더")
    ap.add_argument("--detail", action="store_true", help="텍스트 상자별 추정 높이 출력")
    a = ap.parse_args()
    root = Path(a.root).resolve()
    ref_path = root / REF_NAME
    try:
        ref = Ref(json.loads(ref_path.read_text(encoding="utf-8")))
        if not ref.nums or any(k not in ref.data.get(s, {}) for s, k, _ in CORE):
            raise ValueError("기준값이 비었거나 핵심 지표 키가 없다")
    except Exception as e:
        print(f"FAIL 기준 파일 {ref_path} — {type(e).__name__}: {e}")
        print("RESULT FAIL (판정 0 · 미판정 0 · 기준 파일을 읽지 못함)")
        return 1
    print(f"기준 {REF_NAME}: 수 값 {len(ref.nums)}개 · 날짜 {len(ref.dates)}개 (_meta 제외) · 폴더 {root}")
    tallies = [run_target(root, cfg, ref) for cfg in TARGETS]
    judged = unk = bad = 0
    allow, ids, decor, code_nums, prop_imgs = Counter(), 0, 0, 0, 0
    for t in tallies:
        print(summary(t))
        for line in t.fails[:40]:
            print(f"  FAIL {line}")
        if len(t.fails) > 40:
            print(f"  FAIL … 외 {len(t.fails) - 40}건")
        for line in t.unjudged[:20]:
            print(f"  미판정 {line}")
        if a.detail and t.detail:
            print("\n".join(t.detail))
        judged += sum(v for k, v in t.c.items() if k.endswith("_판정"))
        bad += sum(v for k, v in t.c.items() if k.endswith("_위반"))
        unk += t.c["미판정"]
        allow.update(t.allow)
        ids += t.c["식별자"]
        decor += t.c["장식제외"]
        code_nums += t.c["범위밖_스크립트숫자"]
        prop_imgs += t.c["범위밖_속성그림"]
    present = sum(1 for t in tallies if t.c["입력_판정"] == 0)
    print("허용 목록 적용: " + (" · ".join(f"{k} {v}" for k, v in allow.items()) or "없음")
          + f" · 식별자로 뺀 토큰 {ids}")
    print(f"범위 밖(판정하지 않음): dashboard <script> 코드 숫자 리터럴 {code_nums}개와 <style>·태그 속성 · "
          f"JS가 그린 화면 값(브라우저 자기검증 배너 담당) · PPTX 차트 축 눈금(렌더 시 계산) · "
          f"그림 픽셀(OCR 안 함) · 장식 도형 {decor}개(겹침만 제외) · 문서 속성(본문에 없는 썸네일 그림 {prop_imgs}개)")
    verdict = "PASS" if bad == 0 and unk == 0 else ("FAIL" if bad else "UNJUDGED")
    print(f"RESULT {verdict} (판정 {judged} · 위반 {bad} · 미판정 {unk} · 파일 {present}/{len(TARGETS)})")
    return 0 if verdict == "PASS" else 1


if __name__ == "__main__":
    sys.exit(main())
