#!/usr/bin/env python3
"""gen_04.py - 실습 4 「여러 파일 → 주간 보고서」 자료 생성 (Phase D · D5).

정본: plans/FRAME-개편/준비물_사양.md §0(메인 결정 A~E) · §1 · §5.

산출(base = courses/AI_에이전트_실습워크숍_4시간/sessions/1주차)
  base/실습자료/실습자료_FRAME/04_주간보고서/**            작업 폴더(일일 메모 5개 + 매출.xlsx)
  base/실습자료/실습자료_FRAME/이어가기/04_주간보고서/**    첫 요청 결과(주간보고서.md, 수요일 매출이 메모 값인 상태)
  base/실습자료_강사용/04_주간보고서/**                    정답 주간 보고서 + 정답 기준값
  tmp/frame/D245/매출.csv                                 같은 표의 csv(zip에는 넣지 않는다, 사양 §0-3)

재현성: 난수가 없다. XLSX는 작성자 「도담수납」 · zip 항목 시각 · core.xml 수정 시각을 고정해 다시 묶는다.
파일 수정 시각은 2026년 3월로 고정(os.utime). 같은 스크립트를 두 번 돌리면 sha256이 같다.

사용: python plans/FRAME-개편/gen/gen_04.py [--base DIR] [--sha FILE]
"""
from __future__ import annotations

import sys as _sys

if hasattr(_sys.stdout, "reconfigure"):
    _sys.stdout.reconfigure(encoding="utf-8")
    _sys.stderr.reconfigure(encoding="utf-8")

import argparse
import datetime
import hashlib
import io
import os
import re
import sys
import unicodedata
import zipfile
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
except ImportError as e:  # pragma: no cover
    print(f"[gen_04] openpyxl 없음: {e}", file=sys.stderr)
    sys.exit(2)

REPO = Path(__file__).resolve().parents[3]
COURSE = "AI_에이전트_실습워크숍_4시간"
DEFAULT_BASE = REPO / "courses" / COURSE / "sessions" / "1주차"
DEFAULT_SHA = REPO / "tmp" / "frame" / "D245" / "sha256_04.txt"
DEFAULT_CSV_DIR = REPO / "tmp" / "frame" / "D245"
YEAR = 2026
FOLDER = "04_주간보고서"
AUTHOR = "도담수납"

# ------------------------------------------------------------------ 매출 표 (사양 §5-2)
SALES = [  # (일, 구분, 고객, 내용, 금액)
    (23, "계약금", "서씨댁", "붙박이장 84형 계약금", 1350000),
    (23, "잔금", "송씨댁", "붙박이장 잔금", 2160000),
    (23, "유상수리", "노씨댁", "경첩 교체", 180000),
    (24, "계약금", "문씨댁", "신발장 계약금", 1620000),
    (24, "중도금", "노씨댁", "붙박이장 중도금", 1460000),
    (24, "샘플", "현장 판매", "마감재 소품 판매", 90000),
    (25, "계약금", "배씨댁", "샘플하이츠 1세대 계약금", 1860000),
    (25, "잔금", "조씨댁", "팬트리 잔금", 2520000),
    (25, "유상수리", "하씨댁", "서랍 레일 교체", 240000),
    (26, "계약금", "유씨댁", "드레스룸 계약금", 2100000),
    (26, "잔금", "한씨댁", "붙박이장 잔금", 1245000),
    (26, "유상수리", "민씨댁", "문짝 조정", 120000),
    (27, "계약금", "심씨댁", "붙박이장 계약금", 1290000),
    (27, "잔금", "오씨댁", "수납장 잔금", 2310000),
    (27, "샘플", "현장 판매", "마감재 소품 판매", 70000),
]
DAYS = [(23, "월", "mon"), (24, "화", "tue"), (25, "수", "wed"), (26, "목", "thu"), (27, "금", "fri")]
TABLE_SUM = {d: sum(r[4] for r in SALES if r[0] == d) for d, _, _ in DAYS}
MEMO_SUM = dict(TABLE_SUM)
MEMO_SUM[25] = 4260000  # 수요일 메모만 자릿수가 바뀐 값(4,620,000 → 4,260,000)
WEEK_TABLE = sum(TABLE_SUM.values())   # 18,615,000
WEEK_MEMO = sum(MEMO_SUM.values())     # 18,255,000
assert TABLE_SUM == {23: 3690000, 24: 3170000, 25: 4620000, 26: 3465000, 27: 3670000}
assert WEEK_TABLE == 18615000 and WEEK_MEMO == 18255000
assert all(sum(1 for r in SALES if r[0] == d) == 3 for d, _, _ in DAYS)


def won(n: int) -> str:
    return f"{n:,}"


# ------------------------------------------------------------------ 일일 메모 5개 (사양 §5-1 · §5-4)
MEMO = {  # 일 → (날씨 · 현장 한 줄, 한 일 3줄, 이슈, 섹션 제목, 마지막 줄)
    23: ("맑음. 전 현장에서 작업 조건이 좋았다.",
         ["- 서씨댁 붙박이장 84형 계약금 입금을 확인했다.",
          "- 송씨댁 설치를 마쳤다(오지훈 확인). 잔금 입금도 확인했다.",
          "- 노씨댁 경첩 교체 방문을 했다(유상 수리)."],
         "- 이슈: 없음", "[내일]", "- 자재 단가표 회신이 왔는지 확인한다."),
    24: ("흐림. 오후에 비 소식이 있어 외부 작업은 오전에 마쳤다.",
         ["- 자재 업체의 단가표가 도착했다.",
          "- 문씨댁 신발장 계약금과 노씨댁 붙박이장 중도금 입금을 확인했다.",
          "- 전시장에서 마감재 소품을 판매했다."],
         "- 이슈: 합판 단가가 지난달보다 올랐다.", "[내일]", "- 전시장 벽면 샘플 교체를 진행한다."),
    25: ("맑음. 전시장 조명이 밝아 샘플 교체 작업이 수월했다.",
         ["- 전시장 벽면 샘플 교체를 마쳤다.",
          "- 발주서를 제출했다.",
          "- 배씨댁(샘플하이츠 1세대) 계약금, 조씨댁 잔금 입금과 하씨댁 서랍 레일 교체(유상)를 처리했다."],
         "- 이슈: 조씨댁 잔금이 오후 늦게 확인됐다.", "[내일]", "- 한씨댁 설치 준비를 한다."),
    26: ("흐리고 바람이 강했다. 야외 자재 반입은 최소로 했다.",
         ["- 유씨댁 드레스룸 계약금 입금을 확인했다.",
          "- 한씨댁 붙박이장 잔금 입금을 확인했다.",
          "- 민씨댁 문짝 조정 방문을 마쳤다(유상)."],
         "- 이슈: 한씨댁 설치 일정이 하루 늦춰졌다.", "[내일]", "- 샘플하이츠 단체 견적서를 마무리한다."),
    27: ("맑음. 주말 전이라 사무실이 한산했다.",
         ["- 샘플하이츠 단체 견적서를 제출했다(마감일).",
          "- 심씨댁 붙박이장 계약금과 오씨댁 수납장 잔금 입금을 확인했다.",
          "- 전시장에서 마감재 소품을 판매했다."],
         "- 이슈: 없음", "[다음 주]",
         "- 4월 시공 일정표 갱신(박도윤), 샘플하이츠 입주자 대표에게 견적 설명(김지원), 3월 결산 자료 정리(정하윤)"),
}


def memo_lines(day: int, yoil: str) -> list[str]:
    weather, done, issue, nxt_head, nxt = MEMO[day]
    L = [f"일일 메모 — {YEAR}-03-{day:02d} ({yoil})", "", "작성: 정하윤", weather, "",
         "[오늘 한 일]", *done, "", "[매출·이슈]",
         f"- 오늘 매출 합계: {won(MEMO_SUM[day])}원 (입금 확인 기준, 3건)", issue, "", nxt_head, nxt]
    assert len(L) == 16
    return L


# ------------------------------------------------------------------ 주간 보고서 (사양 §5-7 · §5-8)
def report_lines(corrected: bool) -> list[str]:
    sums = TABLE_SUM if corrected else MEMO_SUM
    week = WEEK_TABLE if corrected else WEEK_MEMO
    L = [f"# 주간 보고서 — {YEAR}-03-23 ~ {YEAR}-03-27", "",
         "- 작성: 정하윤 (경영지원팀)",
         "- 기간: 2026년 3월 23일(월) ~ 3월 27일(금)",
         "- 자료: 일일 메모 5건, 매출.xlsx", "",
         "## 1. 이번 주 요약", "",
         f"1. 이번 주 매출 합계는 {won(week)}원이고, 매일 3건씩 모두 15건의 입금을 확인했다.",
         "2. 화요일에 자재 단가표가 도착했고, 수요일에 발주서 제출과 전시장 벽면 샘플 교체를 마쳤다.",
         "3. 금요일에 샘플하이츠 단체 견적서를 제출해 마감일을 지켰다.", "",
         "## 2. 일자별 매출과 주간 합계", "",
         "| 일자 | 매출(원) | 입금 건수 |", "|---|---:|---:|"]
    for d, y, _ in DAYS:
        L.append(f"| 3월 {d}일({y}) | {won(sums[d])} | 3건 |")
    L += [f"| 주간 합계 | {won(week)} | 15건 |", "",
          "## 3. 요일별 주요 일", "",
          "- 월(3/23): 서씨댁 계약금과 송씨댁 잔금이 입금됐다. 송씨댁 설치를 마쳤다(오지훈 확인).",
          "- 화(3/24): 자재 단가표가 도착했다. 합판 단가가 지난달보다 올랐다.",
          "- 수(3/25): 전시장 벽면 샘플 교체를 마치고 발주서를 제출했다. 조씨댁 잔금은 오후 늦게 확인됐다.",
          "- 목(3/26): 유씨댁 계약금과 한씨댁 잔금이 들어왔다. 한씨댁 설치 일정이 하루 늦춰졌다.",
          "- 금(3/27): 샘플하이츠 단체 견적서를 제출했다(마감일). 심씨댁 계약금과 오씨댁 잔금이 들어왔다.", "",
          "## 4. 다음 주 할 일", "",
          "- 4월 시공 일정표 갱신 — 박도윤",
          "- 샘플하이츠 입주자 대표에게 견적 설명 — 김지원",
          "- 3월 결산 자료 정리 — 정하윤"]
    if corrected:
        L += ["", f"> 참고: 3월 25일(수) 메모의 합계({won(MEMO_SUM[25])}원)와 매출.xlsx의 합계"
                  f"({won(TABLE_SUM[25])}원)가 달라 매출.xlsx 기준으로 적었다."]
    return L


def answer_doc() -> list[str]:
    L = ["# 04_주간보고서 정답 기준값 (강사용)", ""]
    L += ["정답_주간보고서.md가 고친 결과의 예시다.",
          "이어가기 폴더의 주간보고서.md는 수요일 매출이 메모 값(4,260,000원)인 상태다.", ""]
    L += ["## 1. 메모 합계와 표 합계", "",
          "| 요일 | 메모 12줄의 합계 | 매출.xlsx 그날 세 행의 합계 | 일치 |", "|---|---|---|---|"]
    for d, y, _ in DAYS:
        same = "일치" if MEMO_SUM[d] == TABLE_SUM[d] else "다름"
        L.append(f"| {y} 3/{d} | {won(MEMO_SUM[d])}원 | {won(TABLE_SUM[d])}원 | {same} |")
    L += [f"| 주간 | {won(WEEK_MEMO)}원 | {won(WEEK_TABLE)}원 | 다름 |", "",
          "- 다른 곳은 수요일 한 군데뿐이다(자릿수가 바뀜, 차이 360,000원).",
          "- 메모 다섯 개에는 주간 합계를 적지 않았다. 메모끼리 더한 값이 18,255,000원이다.",
          "- 메모의 「3건」은 표의 그날 행 수와 같다.", ""]
    L += ["## 2. 정답 기준값", "",
          "| 항목 | 값 |", "|---|---|",
          "| 일자별 매출(표 기준) | 월 3,690,000 · 화 3,170,000 · 수 4,620,000 · 목 3,465,000 · 금 3,670,000 |",
          "| 주간 합계(표 기준) | 18,615,000원 |",
          "| 확인 거리에 걸린 결과 | 수요일 4,260,000원 또는 주간 합계 18,255,000원 |",
          "| 요일별 주요 일 | 월: 서씨댁 계약 또는 송씨댁 잔금 · 화: 자재 단가표 도착 · 수: 전시장 샘플 교체 완료 |",
          "| (이어서) | 목: 한씨댁 설치 일정 하루 순연 · 금: 샘플하이츠 단체 견적서 제출 |",
          "| 다음 주 할 일 | 3건(4월 시공 일정표 갱신 · 견적 설명 · 3월 결산 자료 정리) |",
          "| 고친 결과의 판정 | 수요일 4,620,000 · 주간 합계 18,615,000, 메모와 표가 다르다는 한 줄이 맨 아래에 있다 |",
          "", "각 요일의 주요 일은 1개 이상 있어야 한다. 고친 결과에서 다른 요일의 값은 바뀌지 않는다.", ""]
    L += ["## 3. 확인 거리", "",
          "- 위치: 메모_0325_수.txt 12줄(4,260,000원)과 매출.xlsx 8~10행(합 4,620,000원).",
          "- 보고서의 일자별 매출 다섯 개를 각각 매출.xlsx의 그날 행 세 개를 더한 값과 맞춘다.",
          "- 다른 요일(수요일)을 찾고, 주간 합계를 표의 열 전체를 더한 값과 맞춘다.",
          "- 매출.xlsx에는 합계 행과 수식이 없다. 합계는 행을 직접 더해서 구한다."]
    return L


# ------------------------------------------------------------------ XLSX · CSV
def xlsx_bytes() -> bytes:
    """openpyxl 출력의 zip 항목 시각 · core.xml 수정 시각을 고정해 다시 묶는다(작성자 도담수납)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "3월4주"
    ws.append(["일자", "구분", "고객", "내용", "금액(원)"])
    for d, kind, who, what, amt in SALES:
        ws.append([datetime.date(YEAR, 3, d), kind, who, what, amt])
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = PatternFill("solid", start_color="DDEBF7", end_color="DDEBF7")
    for c, w in zip("ABCDE", (14, 12, 14, 30, 14)):
        ws.column_dimensions[c].width = w
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=5):
        for cell in row:
            cell.border = border
            if cell.row > 1 and cell.column == 1:
                cell.number_format = "yyyy-mm-dd"
            if cell.row > 1 and cell.column == 5:
                cell.number_format = "#,##0"
    iso = f"{YEAR}-03-27T10:00:00Z"
    wb.properties.title = "3월 4주 매출"
    wb.properties.creator = AUTHOR
    wb.properties.lastModifiedBy = AUTHOR
    wb.properties.created = datetime.datetime(YEAR, 3, 27, 10, 0, 0)
    buf = io.BytesIO()
    wb.save(buf)
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for name in src.namelist():
            data = src.read(name)
            if name == "docProps/core.xml":
                txt = data.decode("utf-8")
                txt = re.sub(r"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                             lambda m: m.group(1) + iso + m.group(2), txt)
                data = txt.encode("utf-8")
            zi = zipfile.ZipInfo(name, date_time=(YEAR, 3, 27, 10, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.create_system = 0
            zi.external_attr = 0
            zo.writestr(zi, data)
    return out.getvalue()


def csv_bytes() -> bytes:
    L = ["일자,구분,고객,내용,금액(원)"]
    for d, kind, who, what, amt in SALES:
        L.append(f"{YEAR}-03-{d:02d},{kind},{who},{what},{amt}")
    return ("\n".join(L) + "\n").encode("utf-8")


# ------------------------------------------------------------------ 쓰기
def _ts(day: int, hh=10, mm=0) -> float:
    return datetime.datetime(YEAR, 3, day, hh, mm, 0).timestamp()


def text_bytes(lines: list[str], name: str) -> bytes:
    for i, s in enumerate(lines, 1):
        if len(s) > 100:
            raise ValueError(f"{name} {i}줄이 {len(s)}자다(100자 이내): {s[:30]}")
        if unicodedata.normalize("NFC", s) != s or "\r" in s or "\n" in s:
            raise ValueError(f"{name} {i}줄: NFC 아님 또는 줄바꿈 포함")
    return ("\n".join(lines) + "\n").encode("utf-8")


class Writer:
    def __init__(self):
        self.files: dict[Path, tuple[bytes, int]] = {}

    def put(self, path: Path, data: bytes, day: int):
        self.files[path] = (data, day)

    def flush(self, tops: list[Path], shared: list[Path]):
        for path, (data, day) in sorted(self.files.items(), key=lambda kv: str(kv[0])):
            path.parent.mkdir(parents=True, exist_ok=True)
            with open(path, "wb") as f:
                f.write(data)
            os.utime(path, (_ts(day), _ts(day)))
        for top in tops:  # 아래에서 위로
            dirs = [p for p in [top, *top.rglob("*")] if p.is_dir()]
            for p in sorted(dirs, key=lambda q: len(q.parts), reverse=True):
                os.utime(p, (_ts(30), _ts(30)))
        for p in shared:  # 여러 생성 스크립트가 함께 쓰는 상위 폴더(수정 시각만 3월로 맞춘다)
            if p.exists():
                os.utime(p, (_ts(30), _ts(30)))


def generate(base: Path | None = None, sha_out: Path | None = None, csv_dir: Path | None = None) -> dict:
    base = Path(base) if base else DEFAULT_BASE
    csv_dir = Path(csv_dir) if csv_dir else (DEFAULT_CSV_DIR if base == DEFAULT_BASE else base / "_csv_04")
    work = base / "실습자료" / "실습자료_FRAME" / FOLDER
    cont = base / "실습자료" / "실습자료_FRAME" / "이어가기" / FOLDER
    ans = base / "실습자료_강사용" / FOLDER
    w = Writer()
    for d, y, key in DAYS:
        name = f"메모_03{d}_{y}.txt"
        w.put(work / name, text_bytes(memo_lines(d, y), name), d)
    w.put(work / "매출.xlsx", xlsx_bytes(), 27)
    w.put(cont / "주간보고서.md", text_bytes(report_lines(False), "이어가기/주간보고서.md"), 30)
    w.put(ans / "정답_주간보고서.md", text_bytes(report_lines(True), "정답_주간보고서.md"), 30)
    w.put(ans / "정답_기준값.md", text_bytes(answer_doc(), "정답_기준값.md"), 30)
    frame = base / "실습자료" / "실습자료_FRAME"
    w.flush([work, cont, ans], [frame, frame.parent, frame / "이어가기", base / "실습자료_강사용"])

    # 매출.csv는 zip 밖(tmp)에만 둔다
    csv_dir.mkdir(parents=True, exist_ok=True)
    csv_path = csv_dir / "매출.csv"
    csv_path.write_bytes(csv_bytes())
    os.utime(csv_path, (_ts(27), _ts(27)))

    shas: dict[str, str] = {}
    for top in (work, cont, ans):
        for p in sorted(top.rglob("*")):
            if p.is_file():
                shas[p.relative_to(base).as_posix()] = hashlib.sha256(p.read_bytes()).hexdigest()
    csv_sha = hashlib.sha256(csv_path.read_bytes()).hexdigest()
    if sha_out:
        Path(sha_out).parent.mkdir(parents=True, exist_ok=True)
        with open(sha_out, "w", encoding="utf-8", newline="\n") as f:
            for rel, h in sorted(shas.items()):
                f.write(f"{h}  {rel}\n")
            f.write(f"{csv_sha}  tmp/frame/D245/매출.csv\n")
    return {"sha": shas, "csv_sha": csv_sha, "csv_path": csv_path, "paths": {"work": work, "cont": cont, "ans": ans}}


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--base", default=None, help="출력 기준 폴더(기본: 1주차 폴더)")
    ap.add_argument("--sha", default=None, help="sha256 목록 파일(기본: tmp/frame/D245/sha256_04.txt)")
    a = ap.parse_args(argv)
    base = Path(a.base) if a.base else DEFAULT_BASE
    sha = Path(a.sha) if a.sha else (DEFAULT_SHA if a.base is None else base / "sha256_04.txt")
    r = generate(base, sha)
    print(f"[gen_04] 파일 {len(r['sha'])}개 작성 · csv {r['csv_path']} · sha256 목록 {sha}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
