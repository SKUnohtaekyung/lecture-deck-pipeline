#!/usr/bin/env python3
"""check_02.py - 실습 2 「폴더·파일 정리」 자료 검사 (Phase D · D3).

정본 사양: plans/FRAME-개편/준비물_사양.md §0 · §1-4 · §1-5 · §3. 기대값은 이 파일에
사양 표에서 옮겨 적은 리터럴로 두고(gen_02.py와 독립), gen_02.py는 ⑦ 재생성에만 쓴다.

검사 7항
 ① 전 파일 열림(PIL · openpyxl · pypdf · 텍스트)
 ② 「거의 같은 이름」 쌍: 쌍 A · B는 sha256이 다르고, 쌍 C는 사양대로 바이트 동일
 ③ AGENTS.md 본문 일치 · CLAUDE.md = `@AGENTS.md`
 ④ 자료 수 · 형식 계수 · 폴더별 위치 · 이동기록이 사양 표와 일치
 ⑤ 전화 · 이메일 형식 · D31 금지어 0
 ⑥ 파일 이름에 Windows 금지 문자 0 · NFC · 공백/괄호는 쌍 파일만
 ⑦ 두 번 생성 sha256 동일 · 실제 파일과 일치 · 수정 시각 2026-03

종료코드 0 = 7항 전부 PASS, 1 = FAIL 있음 또는 판정 0건인 항목 있음(눈먼 0).
「판정 N건 · 미판정 M건」을 항목마다 출력한다. 미판정은 FAIL/PASS 계수에 넣지 않는다.
"""
from __future__ import annotations

import datetime
import hashlib
import io
import os
import re
import sys
import unicodedata
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
REPO = HERE.parents[2]
TMP = REPO / "tmp" / "frame" / "D3"

try:
    from PIL import Image
    import openpyxl
except ImportError as e:  # pragma: no cover
    print(f"[check_02] PIL/openpyxl 없음: {e}")
    sys.exit(1)
try:
    import pypdf
except ImportError:
    pypdf = None

sys.dont_write_bytecode = True  # gen/ 안에 __pycache__를 남기지 않는다
import gen_02  # ⑦ 재생성 전용

BASE = gen_02.DEFAULT_BASE
RECV = BASE / "실습자료" / "실습자료_FRAME" / "02_폴더정리" / "받은자료"
CONT = BASE / "실습자료" / "실습자료_FRAME" / "이어가기" / "02_폴더정리" / "정리후"
ANS = BASE / "실습자료_강사용" / "02_폴더정리"
ANS_TREE = ANS / "정답_정리후"
ROOTS = {"받은자료": RECV, "정리후": CONT, "정답": ANS_TREE}
TOP_DIRS = [RECV.parent, CONT.parent, ANS]

# ---------------------------------------------------------------- 사양 §3-2 리터럴
LOCK = {  # 삭제금지/ 4개
    "계약서_원본스캔_송씨댁.pdf": "PDF",
    "계약서_원본스캔_노씨댁.pdf": "PDF",
    "사업자등록증_사본.png": "PNG",
    "보험가입증서.pdf": "PDF",
}
SPEC = [  # (번호, 파일, 형식, 정답 폴더)
    (1, "견적서_송씨댁_0303.pdf", "PDF", "01_견적서"),
    (2, "견적서_노씨댁_0304.pdf", "PDF", "01_견적서"),
    (3, "견적서_서씨댁_0305.pdf", "PDF", "01_견적서"),
    (4, "견적서_최종.pdf", "PDF", "01_견적서"),
    (5, "견적서_최종(1).pdf", "PDF", "01_견적서"),
    (6, "견적서_문씨댁_0309.pdf", "PDF", "01_견적서"),
    (7, "견적서_양식.xlsx", "XLSX", "01_견적서"),
    (8, "계약서_송씨댁_0310.pdf", "PDF", "02_계약서"),
    (9, "계약서_노씨댁_0311.pdf", "PDF", "02_계약서"),
    (10, "계약서_샘플하이츠_초안.pdf", "PDF", "02_계약서"),
    (11, "현장A_전경.jpg", "JPG", "03_현장사진"),
    (12, "현장A_전경 - 복사본.jpg", "JPG", "03_현장사진"),
    (13, "현장A_설치전.jpg", "JPG", "03_현장사진"),
    (14, "현장A_설치후.jpg", "JPG", "03_현장사진"),
    (15, "현장B_실측.jpg", "JPG", "03_현장사진"),
    (16, "현장B_자재입고.png", "PNG", "03_현장사진"),
    (17, "현장C_현관.jpg", "JPG", "03_현장사진"),
    (18, "현장C_수납장.jpg", "JPG", "03_현장사진"),
    (19, "샘플하이츠_동배치.png", "PNG", "03_현장사진"),
    (20, "영수증_자재_0302.pdf", "PDF", "04_영수증"),
    (21, "영수증_자재_0309.pdf", "PDF", "04_영수증"),
    (22, "영수증_운송_0306.jpg", "JPG", "04_영수증"),
    (23, "영수증_공구_0311.jpg", "JPG", "04_영수증"),
    (24, "영수증_식대_0310.png", "PNG", "04_영수증"),
    (25, "도면_84타입.png", "PNG", "05_도면"),
    (26, "도면_59타입.png", "PNG", "05_도면"),
    (27, "도면_A현장.pdf", "PDF", "05_도면"),
    (28, "도면_C현장.pdf", "PDF", "05_도면"),
    (29, "재고현황.xlsx", "XLSX", "06_기타"),
    (30, "재고현황(1).xlsx", "XLSX", "06_기타"),
    (31, "자재단가표.xlsx", "XLSX", "06_기타"),
    (32, "메모_연락처.txt", "TXT", "06_기타"),
    (33, "메모_할일.txt", "TXT", "06_기타"),
    (34, "협력업체_목록.txt", "TXT", "06_기타"),
    (35, "공지_0310.txt", "TXT", "06_기타"),
    (36, "임시.txt", "TXT", "06_기타"),
]
NAME_OF = {n: name for n, name, _, _ in SPEC}
FMT_TOTAL = {"PDF": 16, "JPG": 9, "PNG": 6, "XLSX": 4, "TXT": 5}
FOLDERS = ["01_견적서", "02_계약서", "03_현장사진", "04_영수증", "05_도면", "06_기타"]
# 사양 §3-7 · §3-6
CONT_COUNTS = {"01_견적서": 6, "02_계약서": 3, "03_현장사진": 8, "04_영수증": 5, "05_도면": 4,
               "06_기타": 6, "보류": 4, "삭제금지": 4}
CONT_HOLD = {5, 12, 30, 36}
ANS_COUNTS = {"01_견적서": 7, "02_계약서": 3, "03_현장사진": 8, "04_영수증": 5, "05_도면": 4,
              "06_기타": 7, "보류": 2, "삭제금지": 4}
ANS_HOLD = {12, 36}
ALLOWED_ALT = {12: {"03_현장사진", "보류"}, 36: {"06_기타", "보류"}}
TXT_LINES = {"메모_연락처.txt": 8, "메모_할일.txt": 10, "협력업체_목록.txt": 12, "공지_0310.txt": 9, "임시.txt": 2}
XLSX_DIMS = {"견적서_양식.xlsx": (11, 6), "재고현황.xlsx": (13, 4), "재고현황(1).xlsx": (13, 4),
             "자재단가표.xlsx": (16, 5)}
PAIR_NAMES_WITH_SPACE_OR_PAREN = {"견적서_최종(1).pdf", "재고현황(1).xlsx", "현장A_전경 - 복사본.jpg"}

AGENTS_MD = """# 받은자료 폴더 정리 규칙

이 폴더의 파일을 정리할 때 아래 규칙을 따른다.

1. `삭제금지` 폴더 안의 파일은 옮기지 않고, 이름을 바꾸지 않고, 지우지 않는다.
2. 파일을 지우지 않는다. 남길 필요가 없어 보이는 파일은 `보류` 폴더로 옮긴다.
3. 파일 이름은 바꾸지 않는다.
4. 분류 폴더는 아래 여섯 개만 쓴다.
   - 01_견적서
   - 02_계약서
   - 03_현장사진
   - 04_영수증
   - 05_도면
   - 06_기타
5. 파일을 옮길 때마다 `이동기록.md`에 「원래 위치 → 새 위치 | 이유」를 한 줄씩 적는다.
6. 같은 파일이 둘 이상이면 하나만 분류 폴더에 두고 나머지는 `보류`로 옮긴다.
"""

D31_WORDS = ["중간 유실", "환각", "아첨", "컨텍스트", "하네스", "인젝션", "자동화 편향", "외주화",
             "토큰", "지식 마감일", "다음 말 예측"]
WIN_BAD = set('\\/:*?"<>|')
WIN_RESERVED = {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}
PHONE_RE = re.compile(r"(?<![\d-])\d{2,4}-\d{3,4}-\d{4}(?![\d-])")
PHONE_OK = re.compile(r"^010-0000-00\d\d$")
PHONE_FLAT = re.compile(r"(?<!\d)01\d{8,9}(?!\d)")
MAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+")
MOVE_RE = re.compile(r"^(?P<src>.+?) → (?P<dst>.+?) \| (?P<why>.+)$")


# ---------------------------------------------------------------- 공용
def targets() -> list[tuple[str, Path, Path]]:
    """(라벨, 기준 폴더, 파일) - 세 트리의 모든 파일과 강사용 정답_트리.md."""
    out = []
    for label, root in ROOTS.items():
        out += [(label, root, p) for p in files_under(root)]
    doc = ANS / "정답_트리.md"
    if doc.exists():
        out.append(("정답(문서)", ANS, doc))
    return out


class Item:
    def __init__(self, no: int, title: str):
        self.no, self.title = no, title
        self.judged = 0
        self.unjudged = 0
        self.fails: list[str] = []
        self.notes: list[str] = []

    def ok(self, n=1):
        self.judged += n

    def fail(self, msg: str):
        self.judged += 1
        self.fails.append(msg)

    def skip(self, msg: str):
        self.unjudged += 1
        self.notes.append("미판정: " + msg)

    @property
    def status(self) -> str:
        if self.fails:
            return "FAIL"
        if self.judged == 0:
            return "미판정(판정 0건)"
        return "PASS"


def sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


def files_under(root: Path) -> list[Path]:
    return sorted(p for p in root.rglob("*") if p.is_file()) if root.exists() else []


def rel(p: Path, root: Path) -> str:
    return p.relative_to(root).as_posix()


def fmt_of(name: str) -> str:
    ext = name.rsplit(".", 1)[-1].lower()
    return {"pdf": "PDF", "jpg": "JPG", "png": "PNG", "xlsx": "XLSX", "txt": "TXT", "md": "MD"}.get(ext, ext.upper())


def read_text(p: Path) -> str | None:
    b = p.read_bytes()
    if b.startswith(b"\xef\xbb\xbf"):
        return None
    try:
        return b.decode("utf-8")
    except UnicodeDecodeError:
        return None


def locate(root: Path, name: str) -> Path | None:
    hits = [p for p in files_under(root) if p.name == name]
    return hits[0] if len(hits) == 1 else None


# ---------------------------------------------------------------- ① 전 파일 열림
def check1(it: Item, regen_texts: dict):
    for label, root in ROOTS.items():
        if not files_under(root):
            it.fail(f"{label}: 파일 없음({root})")
    for label, root, p in targets():
        r = f"{label}/{rel(p, root)}"
        f = fmt_of(p.name)
        try:
            if f in ("PNG", "JPG"):
                with Image.open(p) as im:
                    im.load()
                    if im.size != (800, 600):
                        it.fail(f"{r}: 크기 {im.size} != 800x600")
                        continue
                    if im.getexif():
                        it.fail(f"{r}: EXIF가 있다")
                        continue
                    extra = {k for k in im.info if k not in ("dpi", "jfif", "jfif_version", "jfif_unit", "jfif_density", "progressive", "progression")}
                    if extra:
                        it.fail(f"{r}: 이미지 메타데이터 {sorted(extra)}")
                        continue
                if p.stat().st_size > 200_000:
                    it.fail(f"{r}: 200KB 초과 {p.stat().st_size}")
                    continue
                it.ok()
            elif f == "PDF":
                data = p.read_bytes()
                if not data.startswith(b"%PDF-") or b"%%EOF" not in data[-32:]:
                    it.fail(f"{r}: PDF 머리·꼬리 이상")
                    continue
                if len(data) > 100_000:
                    it.fail(f"{r}: 100KB 초과 {len(data)}")
                    continue
                if pypdf is None:
                    it.ok()
                    it.skip(f"{r}: pypdf 없음 - 쪽 수·이미지 디코딩 미확인")
                    continue
                rd = pypdf.PdfReader(io.BytesIO(data))
                if len(rd.pages) != 1:
                    it.fail(f"{r}: 쪽 수 {len(rd.pages)}")
                    continue
                imgs = rd.pages[0].images
                if len(imgs) != 1:
                    it.fail(f"{r}: 이미지 {len(imgs)}개")
                    continue
                imgs[0].image.load()
                md = rd.metadata
                if not md or md.author != "도담수납":
                    it.fail(f"{r}: PDF 작성자 {getattr(md, 'author', None)!r}")
                    continue
                it.ok()
            elif f == "XLSX":
                wb = openpyxl.load_workbook(p)
                if len(wb.sheetnames) != 1:
                    it.fail(f"{r}: 시트 {len(wb.sheetnames)}개")
                    continue
                if wb.properties.creator != "도담수납" or wb.properties.lastModifiedBy != "도담수납":
                    it.fail(f"{r}: 작성자 {wb.properties.creator!r}/{wb.properties.lastModifiedBy!r}")
                    continue
                ws = wb.active
                want = XLSX_DIMS.get(p.name)
                if want and (ws.max_row, ws.max_column) != want:
                    it.fail(f"{r}: 크기 {(ws.max_row, ws.max_column)} != {want}")
                    continue
                it.ok()
            elif f in ("TXT", "MD"):
                txt = read_text(p)
                if txt is None:
                    it.fail(f"{r}: UTF-8(BOM 없음) 아님")
                    continue
                if "\r" in txt:
                    it.fail(f"{r}: CR 있음(LF만 허용)")
                    continue
                want = TXT_LINES.get(p.name)
                if want is not None and txt.count("\n") != want:
                    it.fail(f"{r}: 줄 수 {txt.count(chr(10))} != {want}")
                    continue
                too_long = [i for i, ln in enumerate(txt.split("\n"), 1) if len(ln) > 100]
                if too_long:
                    it.fail(f"{r}: 100자 넘는 줄 {too_long[:3]}")
                    continue
                it.ok()
            else:
                it.fail(f"{r}: 알 수 없는 형식")
        except Exception as e:  # noqa: BLE001
            it.fail(f"{r}: 열기 실패 {type(e).__name__}: {e}")


# ---------------------------------------------------------------- ② 쌍
def pair_shas(name_a, name_b, where: dict[str, Path]):
    out = {}
    for label, root in where.items():
        a, b = locate(root, name_a), locate(root, name_b)
        out[label] = (sha(a) if a else None, sha(b) if b else None, a, b)
    return out


def check2(it: Item, regen_texts: dict):
    pairs = [("A", "견적서_최종.pdf", "견적서_최종(1).pdf", False),
             ("B", "재고현황.xlsx", "재고현황(1).xlsx", False),
             ("C", "현장A_전경.jpg", "현장A_전경 - 복사본.jpg", True)]
    for tag, na, nb, same in pairs:
        for label, (ha, hb, a, b) in pair_shas(na, nb, ROOTS).items():
            if ha is None or hb is None:
                it.fail(f"쌍 {tag}/{label}: 파일을 찾지 못함")
            elif same and ha != hb:
                it.fail(f"쌍 {tag}/{label}: 사양은 바이트 동일인데 다르다")
            elif not same and ha == hb:
                it.fail(f"쌍 {tag}/{label}: sha256이 같다(서로 달라야 한다)")
            else:
                it.ok()
    # 내용 차이: 쌍 A(고객·합계·발행일), 쌍 B(기준 날짜 · 수량 8행)
    ta, tb = regen_texts.get("견적서_최종.pdf"), regen_texts.get("견적서_최종(1).pdf")
    if not ta or not tb:
        it.skip("쌍 A 그림 글자 목록 없음 - 고객·합계·발행일 차이 미확인")
    else:
        ja, jb = " ".join(ta), " ".join(tb)
        for label, xa, xb in (("고객", "유씨댁", "한씨댁"), ("합계", "7,000,000", "4,150,000"), ("발행일", "3월 6일", "3월 10일")):
            if xa in ja and xb in jb and xb not in ja and xa not in jb:
                it.ok()
            else:
                it.fail(f"쌍 A 내용 차이 없음: {label}")
    a29, a30 = locate(RECV, "재고현황.xlsx"), locate(RECV, "재고현황(1).xlsx")
    if a29 and a30:
        w1, w2 = openpyxl.load_workbook(a29).active, openpyxl.load_workbook(a30).active
        h1 = [c.value for c in w1[1]]
        h2 = [c.value for c in w2[1]]
        if h1 == h2 == ["품목", "규격", "수량", "단위"]:
            it.ok()
        else:
            it.fail(f"쌍 B 열 구성: {h1} / {h2}")
        diff = sum(1 for r in range(2, 14) if w1.cell(r, 3).value != w2.cell(r, 3).value)
        same_other = all(w1.cell(r, c).value == w2.cell(r, c).value for r in range(2, 14) for c in (1, 2, 4))
        if diff == 8 and same_other:
            it.ok()
        else:
            it.fail(f"쌍 B 수량이 다른 행 {diff} != 8 (또는 품목·규격·단위가 다름)")
        if w1.title != w2.title and "3월5일" in w1.title and "3월12일" in w2.title:
            it.ok()
        else:
            it.fail(f"쌍 B 기준 날짜 표시: {w1.title!r} / {w2.title!r}")


# ---------------------------------------------------------------- ③ 규칙 파일
def check3(it: Item, regen_texts: dict):
    lines = AGENTS_MD.count("\n")
    if lines != 16:
        it.fail(f"기대 AGENTS.md가 16줄이 아니다({lines})")
    for label, root in ROOTS.items():
        ag, cl = root / "AGENTS.md", root / "CLAUDE.md"
        if not ag.exists() or not cl.exists():
            it.fail(f"{label}: AGENTS.md/CLAUDE.md 없음")
            continue
        t = read_text(ag)
        if t == AGENTS_MD:
            it.ok()
        else:
            it.fail(f"{label}: AGENTS.md가 사양 §3-3과 다르다")
        if t and t.count("\n") == 16:
            it.ok()
        else:
            it.fail(f"{label}: AGENTS.md 16줄 아님")
        c = read_text(cl)
        if c is not None and c.strip() == "@AGENTS.md" and c.count("\n") <= 1:
            it.ok()
        else:
            it.fail(f"{label}: CLAUDE.md가 `@AGENTS.md` 한 줄이 아니다: {c!r}")
        # 두 파일의 규칙 내용: CLAUDE.md는 AGENTS.md를 가져오므로 가져온 본문 = AGENTS.md
        imported = (root / c.strip().lstrip("@")).exists() if c else False
        if imported:
            it.ok()
        else:
            it.fail(f"{label}: CLAUDE.md가 가리키는 파일이 같은 폴더에 없다")
    shas = {sha(root / "AGENTS.md") for root in ROOTS.values() if (root / "AGENTS.md").exists()}
    if len(shas) == 1:
        it.ok()
    else:
        it.fail("세 위치의 AGENTS.md가 서로 다르다")


# ---------------------------------------------------------------- ④ 계수 · 위치 · 이동기록
def parse_log(path: Path):
    t = read_text(path)
    if t is None:
        return None
    rows = []
    for ln in t.split("\n")[:-1] if t.endswith("\n") else t.split("\n"):
        m = MOVE_RE.match(ln)
        rows.append(m.groupdict() if m else None)
    return rows


def check_tree_counts(it: Item, label, root: Path, counts: dict, hold: set, hold_reason: dict, answer_tree: bool):
    top = sorted(p.name for p in root.iterdir())
    want_top = sorted(["AGENTS.md", "CLAUDE.md", "이동기록.md", *FOLDERS, "보류", "삭제금지"])
    # 비어 있는 분류 폴더는 없어야 한다(모두 채워짐)
    if top == want_top:
        it.ok()
    else:
        it.fail(f"{label}: 최상위 {top} != {want_top}")
    for folder, n in counts.items():
        d = root / folder
        got = sorted(p.name for p in d.iterdir()) if d.is_dir() else []
        if len(got) == n:
            it.ok()
        else:
            it.fail(f"{label}/{folder}: {len(got)}개 != {n}")
    # 위치 대조: 정답 트리는 허용 범위 안이어야 하고, 정리후는 사양 §3-7의 위치(보류 4개)여야 한다
    mismatch = []
    for n, name, fmt, folder in SPEC:
        want = {folder} | ALLOWED_ALT.get(n, set())
        expect = "보류" if n in hold else folder
        p = root / expect / name
        if p.is_file():
            it.ok()
        else:
            it.fail(f"{label}: {n}번 {name}이 {expect}/에 없다")
        if expect not in want:
            mismatch.append(n)
    if answer_tree:
        if mismatch:
            it.fail(f"{label}: 정답 폴더와 다른 파일 {mismatch}")
        else:
            it.ok()
    else:
        if sorted(mismatch) == [5, 30]:
            it.ok()
            it.notes.append(f"{label}: 정답 폴더와 다른 파일 {len(mismatch)}개({mismatch}) - 사양 §3-7의 확인 거리에 걸린 상태")
        else:
            it.fail(f"{label}: 정답 폴더와 다른 파일 {mismatch} != [5, 30]")
    for name in LOCK:
        if (root / "삭제금지" / name).is_file():
            it.ok()
        else:
            it.fail(f"{label}: 삭제금지/{name} 없음")
    # 이동기록
    lg = root / "이동기록.md"
    rows = parse_log(lg) if lg.exists() else None
    if rows is None:
        it.fail(f"{label}: 이동기록.md 없음/읽기 실패")
        return
    if len(rows) == 36 and all(rows):
        it.ok()
    else:
        it.fail(f"{label}: 이동기록 줄 수/형식 {len(rows)} (형식 불일치 {sum(1 for r in rows if not r)}줄)")
        return
    srcs = [r["src"] for r in rows]
    if sorted(srcs) == sorted(NAME_OF.values()):
        it.ok()
    else:
        it.fail(f"{label}: 이동기록 원래 위치가 36개 자료와 일치하지 않는다")
    bad = 0
    for r in rows:
        if not (root / r["dst"]).is_file():
            bad += 1
    if bad == 0:
        it.ok()
    else:
        it.fail(f"{label}: 이동기록의 새 위치에 파일이 없는 줄 {bad}개")
    dup = [r for r in rows if r["why"] == "중복"]
    if len(dup) == sum(1 for v in hold_reason.values() if v == "중복"):
        it.ok()
    else:
        it.fail(f"{label}: 이유 「중복」 줄 {len(dup)} != {sum(1 for v in hold_reason.values() if v == '중복')}")
    empty = [r for r in rows if r["why"] == "내용 없음"]
    if len(empty) == 1 and empty[0]["src"] == "임시.txt":
        it.ok()
    else:
        it.fail(f"{label}: 「내용 없음」 줄이 임시.txt 1줄이 아니다")
    for n, why in hold_reason.items():
        r = next((x for x in rows if x["src"] == NAME_OF[n]), None)
        if r and r["dst"] == f"보류/{NAME_OF[n]}" and r["why"] == why:
            it.ok()
        else:
            it.fail(f"{label}: {n}번 이동기록 줄이 보류·{why}가 아니다")


def check4(it: Item, regen_texts: dict):
    # (a) 원본 40개의 형식 계수
    fmts: dict[str, int] = {}
    for name, f in LOCK.items():
        fmts[f] = fmts.get(f, 0) + 1
    for _, name, f, _ in SPEC:
        fmts[f] = fmts.get(f, 0) + 1
    if fmts == FMT_TOTAL and sum(fmts.values()) == 40:
        it.ok()
    else:
        it.fail(f"사양 표 자체의 형식 합계 {fmts} != {FMT_TOTAL}")
    # (b) 받은자료 실제 상태
    top_files = sorted(p.name for p in RECV.iterdir() if p.is_file()) if RECV.exists() else []
    want_top = sorted(["AGENTS.md", "CLAUDE.md", *NAME_OF.values()])
    if top_files == want_top:
        it.ok()
    else:
        miss = sorted(set(want_top) - set(top_files))
        extra = sorted(set(top_files) - set(want_top))
        it.fail(f"받은자료 바로 아래 파일 불일치: 없음 {miss[:4]} 여분 {extra[:4]}")
    top_dirs = sorted(p.name for p in RECV.iterdir() if p.is_dir()) if RECV.exists() else []
    if top_dirs == ["삭제금지"]:
        it.ok()
    else:
        it.fail(f"받은자료 바로 아래 폴더 {top_dirs} != ['삭제금지']")
    lock_files = sorted(p.name for p in (RECV / "삭제금지").iterdir()) if (RECV / "삭제금지").is_dir() else []
    if lock_files == sorted(LOCK):
        it.ok()
    else:
        it.fail(f"삭제금지/ 파일 {lock_files} != {sorted(LOCK)}")
    got_fmts: dict[str, int] = {}
    for p in files_under(RECV):
        if p.name in ("AGENTS.md", "CLAUDE.md"):
            continue
        got_fmts[fmt_of(p.name)] = got_fmts.get(fmt_of(p.name), 0) + 1
    if got_fmts == FMT_TOTAL:
        it.ok()
    else:
        it.fail(f"받은자료 실제 형식 계수 {got_fmts} != {FMT_TOTAL}")
    total = len(files_under(RECV))
    if total == 42:
        it.ok()
    else:
        it.fail(f"받은자료 파일 수 {total} != 42(자료 40 + 규칙 2)")
    # (c) 정리후 · 정답
    cont_files = len(files_under(CONT))
    if cont_files == 43:
        it.ok()
    else:
        it.fail(f"정리후 파일 수 {cont_files} != 43")
    ans_files = len(files_under(ANS_TREE))
    if ans_files == 43:
        it.ok()
    else:
        it.fail(f"정답 트리 파일 수 {ans_files} != 43")
    if CONT.exists():
        check_tree_counts(it, "정리후", CONT, CONT_COUNTS, CONT_HOLD,
                          {5: "중복", 12: "중복", 30: "중복", 36: "내용 없음"}, False)
    else:
        it.fail("정리후 폴더 없음")
    if ANS_TREE.exists():
        check_tree_counts(it, "정답", ANS_TREE, ANS_COUNTS, ANS_HOLD, {12: "중복", 36: "내용 없음"}, True)
    else:
        it.fail("정답 트리 폴더 없음")
    ans_top = sorted(p.name for p in ANS.iterdir()) if ANS.exists() else []
    if ans_top == ["정답_정리후", "정답_트리.md"]:
        it.ok()
    else:
        it.fail(f"강사용 02_폴더정리 최상위 {ans_top}")
    # (d) 같은 이름 파일은 세 위치에서 바이트 동일(내용이 바뀌지 않았다)
    for n, name, *_ in SPEC:
        hs = {label: sha(locate(root, name)) for label, root in ROOTS.items() if locate(root, name)}
        if len(hs) == 3 and len(set(hs.values())) == 1:
            it.ok()
        else:
            it.fail(f"{name}: 세 위치의 sha256이 다르거나 없음")
    for name in LOCK:
        hs = {label: sha(root / "삭제금지" / name) for label, root in ROOTS.items() if (root / "삭제금지" / name).exists()}
        if len(hs) == 3 and len(set(hs.values())) == 1:
            it.ok()
        else:
            it.fail(f"삭제금지/{name}: 세 위치의 sha256이 다르거나 없음")


# ---------------------------------------------------------------- ⑤ 개인정보 · 금지어
def scan_text(it: Item, where: str, text: str, counters: dict):
    norm = text.replace(" ", "")
    for w in D31_WORDS:
        if w in text or w.replace(" ", "") in norm:
            it.fail(f"{where}: D31 금지어 「{w}」")
            return
    for m in PHONE_RE.finditer(text):
        counters["phone"] += 1
        if not PHONE_OK.match(m.group(0)):
            it.fail(f"{where}: 전화 형식 위반 {m.group(0)}")
            return
    if PHONE_FLAT.search(text):
        it.fail(f"{where}: 붙여 쓴 전화번호 {PHONE_FLAT.search(text).group(0)}")
        return
    for m in MAIL_RE.finditer(text):
        counters["mail"] += 1
        if not m.group(0).endswith("@example.com"):
            it.fail(f"{where}: 이메일 형식 위반 {m.group(0)}")
            return
    it.ok()


def leak_keys() -> list[str]:
    """제작 PC의 사용자 이름 · PC 이름이 파일 바이트에 남았는지 찾는 키(짧은 토막은 우연 일치가 생겨 뺀다)."""
    import socket
    raw = [os.environ.get("USERNAME", ""), os.environ.get("USER", ""), socket.gethostname(),
           "AppData", "C:" + chr(92) + "Users", "C:/Users"]
    raw += [w for n in (os.environ.get("USERNAME", ""), os.environ.get("USER", "")) for w in n.split()]
    return sorted({k for k in raw if len(k) >= 7})


def scan_bytes(it: Item, where: str, p: Path, keys: list[str]):
    blobs = [p.read_bytes()]
    if p.suffix.lower() == ".xlsx":
        import zipfile
        z = zipfile.ZipFile(p)
        blobs += [z.read(n) for n in z.namelist()]
    for b in blobs:
        low = b.lower()
        for k in keys:
            for enc in ("utf-8", "utf-16-be", "utf-16-le"):
                if k.lower().encode(enc) in low:
                    it.fail(f"{where}: 제작 PC 흔적 「{k}」({enc})")
                    return
    it.ok()


def check5(it: Item, regen_texts: dict):
    counters = {"phone": 0, "mail": 0}
    seen_img: set[str] = set()
    keys = leak_keys()
    if not keys:
        it.skip("제작 PC 사용자·PC 이름을 알 수 없어 바이트 검색 키가 없다")
    for label, root, p in targets():
        scan_bytes(it, f"{label}/{rel(p, root)}", p, keys)
    for label, root, p in targets():
        r = f"{label}/{rel(p, root)}"
        f = fmt_of(p.name)
        scan_text(it, f"{r}(파일 이름)", "/".join(p.relative_to(root).parts), counters)
        if f in ("TXT", "MD"):
            t = read_text(p)
            if t is None:
                it.skip(f"{r}: 읽지 못함")
            else:
                scan_text(it, r, t, counters)
        elif f == "XLSX":
            wb = openpyxl.load_workbook(p)
            cells = []
            for ws in wb.worksheets:
                cells.append(ws.title)
                cells += [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
            pr = wb.properties
            cells += [str(x) for x in (pr.title, pr.creator, pr.lastModifiedBy, pr.subject, pr.description) if x]
            scan_text(it, r, "\n".join(cells), counters)
        elif f == "PDF" and pypdf is not None:
            md = pypdf.PdfReader(io.BytesIO(p.read_bytes())).metadata
            scan_text(it, r + "(PDF 메타데이터)", " ".join(str(v) for v in (md or {}).values()), counters)
            if p.name in regen_texts and regen_texts[p.name]:
                if p.name not in seen_img:
                    scan_text(it, r + "(그림 글자)", "\n".join(regen_texts[p.name]), counters)
                    seen_img.add(p.name)
            else:
                it.skip(f"{r}: 그림 글자 목록 없음")
        elif f in ("PNG", "JPG"):
            if p.name in regen_texts and regen_texts[p.name]:
                if p.name not in seen_img:
                    scan_text(it, r + "(그림 글자)", "\n".join(regen_texts[p.name]), counters)
                    seen_img.add(p.name)
            else:
                it.skip(f"{r}: 그림 글자 목록 없음")
        elif f == "PDF":
            it.skip(f"{r}: pypdf 없음")
    it.notes.append(f"전화 {counters['phone']}건 · 이메일 {counters['mail']}건 발견, 전부 허용 형식")
    if counters["phone"] == 0 or counters["mail"] == 0:
        it.fail("전화 또는 이메일이 0건 - 검사가 본문을 못 본 것일 수 있다")


# ---------------------------------------------------------------- ⑥ 파일 이름
def check6(it: Item, regen_texts: dict):
    seen: dict[str, int] = {}
    for label, root in ROOTS.items():
        for p in [*files_under(root), *[d for d in root.rglob("*") if d.is_dir()]]:
            for part in p.relative_to(root).parts:
                it.ok()
                bad = [c for c in part if c in WIN_BAD or ord(c) < 32]
                if bad:
                    it.fail(f"{label}/{part}: 금지 문자 {bad}")
                    continue
                if not unicodedata.is_normalized("NFC", part):
                    it.fail(f"{label}/{part}: NFC 아님")
                if part != part.rstrip(" .") or part != part.lstrip(" "):
                    it.fail(f"{label}/{part}: 이름 끝/앞 공백·점")
                if part.split(".")[0].upper() in WIN_RESERVED:
                    it.fail(f"{label}/{part}: Windows 예약 이름")
                if len(part) > 100:
                    it.fail(f"{label}/{part}: 이름 길이 {len(part)}")
                if (" " in part or "(" in part or ")" in part) and part not in PAIR_NAMES_WITH_SPACE_OR_PAREN:
                    it.fail(f"{label}/{part}: 공백·괄호는 쌍 파일에만 허용")
                seen[part] = seen.get(part, 0) + 1
    it.notes.append(f"고유 이름 {len(seen)}개")


# ---------------------------------------------------------------- ⑦ 재현성
def list_sha(root: Path, base: Path) -> dict[str, str]:
    out = {}
    for top in (root / "실습자료" / "실습자료_FRAME" / "02_폴더정리",
                root / "실습자료" / "실습자료_FRAME" / "이어가기" / "02_폴더정리",
                root / "실습자료_강사용" / "02_폴더정리"):
        for p in sorted(top.rglob("*")) if top.exists() else []:
            if p.is_file():
                out[p.relative_to(root).as_posix()] = sha(p)
    return out


def check7(it: Item, regen_texts: dict, run_a: dict, run_b: dict):
    real = list_sha(BASE, BASE)
    a, b = run_a["sha"], run_b["sha"]
    if a == b and len(a) > 0:
        it.ok(len(a))
    else:
        diff = sorted(k for k in set(a) | set(b) if a.get(k) != b.get(k))
        it.fail(f"두 번 생성한 sha256이 다르다: {len(diff)}개 {diff[:3]}")
    if real == a:
        it.ok(len(real))
    else:
        diff = sorted(k for k in set(a) | set(real) if a.get(k) != real.get(k))
        it.fail(f"실제 파일이 재생성 결과와 다르다(gen_02.py를 다시 돌려야 한다): {len(diff)}개 {diff[:3]}")
    # 수정 시각 2026-03
    for top in TOP_DIRS:
        for p in [top, *(top.rglob("*") if top.exists() else [])]:
            m = datetime.datetime.fromtimestamp(p.stat().st_mtime)
            if (m.year, m.month) == (2026, 3):
                it.ok()
            else:
                it.fail(f"{p.relative_to(BASE).as_posix()}: 수정 시각 {m:%Y-%m-%d}")
    # 작성 시각: XLSX core.xml · PDF 정보
    for label, root in ROOTS.items():
        for p in files_under(root):
            f = fmt_of(p.name)
            if f == "XLSX":
                pr = openpyxl.load_workbook(p).properties
                ok = all(d and (d.year, d.month) == (2026, 3) for d in (pr.created, pr.modified))
                (it.ok if ok else (lambda: it.fail(f"{label}/{p.name}: XLSX 작성 시각 {pr.created}/{pr.modified}")))()
            elif f == "PDF" and pypdf is not None:
                md = pypdf.PdfReader(io.BytesIO(p.read_bytes())).metadata
                d1, d2 = (md.creation_date if md else None), (md.modification_date if md else None)
                ok = all(d and (d.year, d.month) == (2026, 3) for d in (d1, d2))
                (it.ok if ok else (lambda: it.fail(f"{label}/{p.name}: PDF 작성 시각 {d1}/{d2}")))()
    it.notes.append("PDF·XLSX 작성 시각은 고정되어 두 번 생성해도 같다(위 결과로 확인)")


# ---------------------------------------------------------------- 실행
def main() -> int:
    for need in (gen_02.MALGUN,):
        if not os.path.exists(need):
            print(f"[check_02] 글꼴 없음: {need}")
            return 1
    if not (RECV.exists() and CONT.exists() and ANS_TREE.exists()):
        print("[check_02] 산출 폴더가 없다. 먼저 python plans/FRAME-개편/gen/gen_02.py 를 실행한다.")
        return 1
    va, vb = TMP / "verify_A", TMP / "verify_B"
    run_a = gen_02.generate(va, va / "sha256.txt")
    run_b = gen_02.generate(vb, vb / "sha256.txt")
    texts = run_a["texts"]

    items = [Item(1, "전 파일 열림"), Item(2, "거의 같은 이름 쌍"), Item(3, "AGENTS.md · CLAUDE.md"),
             Item(4, "자료 수 · 형식 · 위치 · 이동기록"), Item(5, "전화 · 이메일 · D31 금지어"),
             Item(6, "파일 이름 Windows 금지 문자"), Item(7, "재현성 · 수정 시각")]
    check1(items[0], texts)
    check2(items[1], texts)
    check3(items[2], texts)
    check4(items[3], texts)
    check5(items[4], texts)
    check6(items[5], texts)
    check7(items[6], texts, run_a, run_b)

    tj = tu = 0
    nfail = nblind = 0
    for it in items:
        tj += it.judged
        tu += it.unjudged
        print(f"[{it.no}] {it.title:<24} {it.status:<14} 판정 {it.judged}건 · 미판정 {it.unjudged}건")
        for n in it.notes[:4]:
            print(f"      - {n}")
        for f in it.fails[:6]:
            print(f"      ! {f}")
        if len(it.fails) > 6:
            print(f"      ! ... 외 {len(it.fails) - 6}건")
        nfail += 1 if it.fails else 0
        nblind += 1 if (not it.fails and it.judged == 0) else 0
    npass = sum(1 for it in items if it.status == "PASS")
    print(f"RESULT: {'PASS' if npass == 7 else 'FAIL'} {npass}/7 · 판정 {tj}건 · 미판정 {tu}건 · FAIL {nfail}항목 · 판정 0건 {nblind}항목")
    return 0 if npass == 7 else 1


if __name__ == "__main__":
    sys.exit(main())
