#!/usr/bin/env python3
"""gen_02.py - 실습 2 「폴더·파일 정리」 자료 생성 (Phase D · D3).

정본: plans/FRAME-개편/준비물_사양.md §0(메인 결정 A~E) · §1 · §3.

산출(base = courses/AI_에이전트_실습워크숍_4시간/sessions/1주차)
  base/실습자료/실습자료_FRAME/02_폴더정리/받은자료/**            작업 폴더(40 + 규칙 2)
  base/실습자료/실습자료_FRAME/이어가기/02_폴더정리/정리후/**      첫 요청 결과 상태(43)
  base/실습자료_강사용/02_폴더정리/**                             정답(정리 후 트리 + 이동 기록)

재현성: 난수 시드 고정 · 파일 수정 시각 2026-03 고정 · PDF 메타데이터 시각 고정 ·
XLSX 작성자 「도담수납」 + zip 항목 시각 고정. 같은 스크립트를 두 번 돌리면 sha256이 같다.

사용: python plans/FRAME-개편/gen/gen_02.py [--base DIR] [--sha FILE]
"""
from __future__ import annotations

import sys as _sys

if hasattr(_sys.stdout, "reconfigure"):
    _sys.stdout.reconfigure(encoding="utf-8")
    _sys.stderr.reconfigure(encoding="utf-8")

import argparse
import datetime
import hashlib
import io
import json
import os
import random
import re
import sys
import time
import unicodedata
import zipfile
from pathlib import Path

try:
    from PIL import Image, ImageDraw, ImageFilter, ImageFont
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
except ImportError as e:  # pragma: no cover
    print(f"[gen_02] PIL/openpyxl 없음: {e}", file=sys.stderr)
    sys.exit(2)

REPO = Path(__file__).resolve().parents[3]
COURSE = "AI_에이전트_실습워크숍_4시간"
DEFAULT_BASE = REPO / "courses" / COURSE / "sessions" / "1주차"
DEFAULT_SHA = REPO / "tmp" / "frame" / "D3" / "sha256.txt"

MALGUN = "C:/Windows/Fonts/malgun.ttf"
if not os.path.exists(MALGUN):
    print(f"[gen_02] 한글 글꼴이 없어 멈춥니다: {MALGUN}", file=sys.stderr)
    sys.exit(2)

SEED = 20260302
AUTHOR = "도담수납"
YEAR = 2026

# ------------------------------------------------------------------ 공통 색
INK = (33, 37, 41)
GRAY = (108, 117, 125)
LIGHT = (233, 236, 239)
LINE = (173, 181, 189)
WHITE = (255, 255, 255)
BRAND = (31, 78, 121)
ACCENT = (200, 110, 40)
RED = (190, 40, 40)

_FONTS: dict[int, ImageFont.FreeTypeFont] = {}


def font(size: int) -> ImageFont.FreeTypeFont:
    if size not in _FONTS:
        _FONTS[size] = ImageFont.truetype(MALGUN, size)
    return _FONTS[size]


def won(n: int) -> str:
    return f"{n:,}"


class Cv:
    """그리기 도우미. 그린 글자를 texts에 모아 검사기가 읽게 한다."""

    def __init__(self, w: int, h: int, bg=WHITE):
        self.w, self.h = w, h
        self.im = Image.new("RGB", (w, h), bg)
        self.d = ImageDraw.Draw(self.im)
        self.texts: list[str] = []

    def t(self, xy, s, size=18, fill=INK, bold=False, anchor="la"):
        f = font(size)
        self.d.text(xy, s, font=f, fill=fill, anchor=anchor)
        if bold:
            self.d.text((xy[0] + 1, xy[1]), s, font=f, fill=fill, anchor=anchor)
        self.texts.append(s)

    def tw(self, s: str, size: int) -> float:
        return self.d.textlength(s, font=font(size))

    def wrap(self, text: str, size: int, maxw: int) -> list[str]:
        out, cur = [], ""
        for ch in text:
            if self.tw(cur + ch, size) > maxw and cur:
                sp = cur.rfind(" ")
                if sp > 0 and ch != " ":
                    out.append(cur[:sp])
                    cur = cur[sp + 1:] + ch
                else:
                    out.append(cur)
                    cur = ch.lstrip()
            else:
                cur += ch
        if cur:
            out.append(cur)
        return out

    def hatch(self, box, color=LINE, gap=8, bg=None):
        x0, y0, x1, y1 = box
        tile = Image.new("RGB", (x1 - x0, y1 - y0), bg or WHITE)
        td = ImageDraw.Draw(tile)
        w, h = tile.size
        for k in range(-h, w, gap):
            td.line((k, h, k + h, 0), fill=color, width=1)
        self.im.paste(tile, (x0, y0))
        self.d.rectangle(box, outline=INK, width=2)

    def dim_h(self, x0, x1, y, text, size=15, tick=7):
        d = self.d
        d.line((x0, y, x1, y), fill=INK, width=1)
        for x in (x0, x1):
            d.line((x, y - tick, x, y + tick), fill=INK, width=1)
        self.t(((x0 + x1) // 2, y - 10), text, size, INK, anchor="ms")

    def dim_v(self, x, y0, y1, text, size=15, tick=7):
        d = self.d
        d.line((x, y0, x, y1), fill=INK, width=1)
        for y in (y0, y1):
            d.line((x - tick, y, x + tick, y), fill=INK, width=1)
        self.t((x - 10, (y0 + y1) // 2), text, size, INK, anchor="rm")


# ------------------------------------------------------------------ 바이트 만들기
def _day_struct(day: int, hh=10):
    return time.strptime(f"{YEAR}-03-{day:02d} {hh:02d}:00:00", "%Y-%m-%d %H:%M:%S")


def pdf_bytes(im: Image.Image, title: str, day: int, limit=100_000) -> bytes:
    """PIL로 이미지형 1쪽 PDF. 100KB를 넘으면 JPEG 품질을 낮춰 다시 만든다."""
    for q in (75, 68, 60, 52):
        buf = io.BytesIO()
        st = _day_struct(day)
        im.save(
            buf, "PDF", resolution=96.0, title=title, author=AUTHOR, creator=AUTHOR,
            producer=AUTHOR, subject=title, creationDate=st, modDate=st, quality=q,
        )
        data = buf.getvalue()
        if len(data) <= limit:
            return data
    raise RuntimeError(f"PDF가 {limit}바이트를 넘습니다: {title} {len(data)}")


def png_bytes(im: Image.Image, limit=200_000) -> bytes:
    buf = io.BytesIO()
    im.save(buf, "PNG", optimize=True)
    data = buf.getvalue()
    if len(data) > limit:
        raise RuntimeError(f"PNG가 {limit}바이트를 넘습니다: {len(data)}")
    return data


def jpg_bytes(im: Image.Image, limit=200_000) -> bytes:
    for q in (88, 80, 72):
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=q)
        data = buf.getvalue()
        if len(data) <= limit:
            return data
    raise RuntimeError(f"JPG가 {limit}바이트를 넘습니다: {len(data)}")


def xlsx_bytes(wb: openpyxl.Workbook, day: int) -> bytes:
    """openpyxl 출력의 zip 항목 시각 · core.xml 수정 시각을 고정해 다시 묶는다."""
    iso = f"{YEAR}-03-{day:02d}T10:00:00Z"
    wb.properties.creator = AUTHOR
    wb.properties.lastModifiedBy = AUTHOR
    wb.properties.created = datetime.datetime(YEAR, 3, day, 10, 0, 0)
    buf = io.BytesIO()
    wb.save(buf)
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for name in src.namelist():
            data = src.read(name)
            if name == "docProps/core.xml":
                txt = data.decode("utf-8")
                txt = re.sub(
                    r"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                    lambda m: m.group(1) + iso + m.group(2), txt)
                data = txt.encode("utf-8")
            zi = zipfile.ZipInfo(name, date_time=(YEAR, 3, day, 10, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.create_system = 0
            zi.external_attr = 0
            zo.writestr(zi, data)
    return out.getvalue()


def txt_bytes(text: str) -> bytes:
    assert "\r" not in text
    return text.encode("utf-8")


class Blob:
    def __init__(self, data: bytes, texts: list[str] | None = None):
        self.data = data
        self.texts = texts or []


# ------------------------------------------------------------------ 문서(PDF) 그림
PAGE_W, PAGE_H = 794, 1123


def _header(cv: Cv, sub="주문 제작 붙박이장·수납가구"):
    cv.t((56, 46), "도담수납", 26, BRAND, bold=True)
    cv.t((PAGE_W - 56, 54), sub, 14, GRAY, anchor="ra")
    cv.d.line((56, 92, PAGE_W - 56, 92), fill=BRAND, width=3)


def _footer(cv: Cv, y=None):
    y = y or PAGE_H - 56
    cv.d.line((56, y - 14, PAGE_W - 56, y - 14), fill=LINE, width=1)
    cv.t((56, y), "도담수납 | 가상시 예시구 샘플로 10 | 대표번호 010-0000-0010 | www.example.com", 13, GRAY)


QUOTE_NOTE_PAY = "결제: 계약금 30% · 중도금 40% · 잔금 30%"


def render_quote(no, issued, customer, subject, items, total, final=False) -> Blob:
    s = sum(q * p for (_, _, q, _, p) in items)
    assert s == total, (no, s, total)
    cv = Cv(PAGE_W, PAGE_H)
    d = cv.d
    _header(cv)
    cv.t((PAGE_W // 2, 150), "견 적 서" + (" (최종)" if final else ""), 40, INK, bold=True, anchor="mm")
    rows = [("견적번호", no), ("발행일", issued), ("고객", f"{customer} 귀하"),
            ("품목", subject), ("유효기간", "발행일로부터 14일")]
    y = 196
    for k, v in rows:
        cv.t((56, y), k, 16, GRAY)
        cv.t((150, y - 1), v, 18, INK, bold=(k == "고객"))
        y += 34
    bx0, by0, bx1, by1 = 440, 190, PAGE_W - 56, 362
    d.rectangle((bx0, by0, bx1, by1), outline=LINE, width=2)
    sup = [("상호", "도담수납"), ("등록번호", "000-00-00000"), ("대표", "최민재"),
           ("주소", "가상시 예시구 샘플로 10"), ("전화", "010-0000-0010")]
    y = by0 + 12
    for k, v in sup:
        cv.t((bx0 + 14, y), k, 14, GRAY)
        cv.t((bx0 + 96, y - 1), v, 15, INK)
        y += 30
    d.rectangle((56, 392, PAGE_W - 56, 452), fill=LIGHT)
    cv.t((72, 422), "합계금액 (부가세 포함)", 18, INK, anchor="lm")
    cv.t((PAGE_W - 72, 422), f"{won(total)}원", 28, BRAND, bold=True, anchor="rm")
    cols = [56, 92, 330, 470, 526, 626, PAGE_W - 56]
    heads = ["번호", "품목", "규격", "수량", "단가", "금액"]
    ty = 478
    d.rectangle((56, ty, PAGE_W - 56, ty + 40), fill=BRAND)
    for i, h in enumerate(heads):
        cx = (cols[i] + cols[i + 1]) // 2
        cv.t((cx, ty + 20), h, 15, WHITE, anchor="mm")
    y = ty + 40
    rh = 46
    for n, (name, spec, q, unit, p) in enumerate(items, 1):
        d.line((56, y + rh, PAGE_W - 56, y + rh), fill=LINE, width=1)
        cv.t(((cols[0] + cols[1]) // 2, y + rh // 2), str(n), 15, INK, anchor="mm")
        cv.t((cols[1] + 8, y + rh // 2), name, 15, INK, anchor="lm")
        cv.t((cols[2] + 8, y + rh // 2), spec, 14, GRAY, anchor="lm")
        cv.t(((cols[3] + cols[4]) // 2, y + rh // 2), f"{q}{unit}", 15, INK, anchor="mm")
        cv.t((cols[5] - 8, y + rh // 2), won(p), 15, INK, anchor="rm")
        cv.t((cols[6] - 8, y + rh // 2), won(q * p), 15, INK, anchor="rm")
        y += rh
    d.rectangle((56, y, PAGE_W - 56, y + 46), fill=LIGHT)
    cv.t((cols[5] - 8, y + 23), "합계", 16, INK, bold=True, anchor="rm")
    cv.t((cols[6] - 8, y + 23), won(total), 16, INK, bold=True, anchor="rm")
    y += 46
    d.rectangle((56, ty, PAGE_W - 56, y), outline=INK, width=1)
    y += 34
    cv.t((56, y), "비고", 16, BRAND, bold=True)
    notes = [QUOTE_NOTE_PAY, "견적 범위: 제작 · 운반 · 설치 (철거와 전기 공사는 별도)",
             "문의: 영업팀장 김지원 010-0000-0012 · support@example.com"]
    for n in notes:
        y += 30
        cv.t((72, y), "· " + n, 15, INK)
    _footer(cv)
    return Blob(b"", cv.texts), cv.im


def render_quote_blob(day, title, *a, **k):
    blob, im = render_quote(*a, **k)
    return Blob(pdf_bytes(im, title, day), blob.texts)


def _scribble(cv: Cv, x0, y0, seed):
    r = random.Random(seed)
    pts = []
    for i in range(46):
        pts.append((x0 + i * 3, y0 + 12 * ((i % 7) - 3) / 3 + r.randint(-3, 3) + 6 * (1 if i % 11 < 5 else -1)))
    cv.d.line(pts, fill=(20, 30, 90), width=2)


def _stamp(cv: Cv, cx, cy, r=27):
    cv.d.ellipse((cx - r, cy - r, cx + r, cy + r), outline=RED, width=3)
    cv.t((cx, cy), "인", 24, RED, bold=True, anchor="mm")


def render_contract(title, party_a, unit_desc, place, total, date_txt, quote_ref, scan=False,
                    draft=False, phone_a=None, seed=1) -> tuple[Cv, Image.Image]:
    cv = Cv(PAGE_W, PAGE_H, (251, 251, 249))
    d = cv.d
    _header(cv, "붙박이장 제작·설치")
    cv.t((PAGE_W // 2, 146), title, 32, INK, bold=True, anchor="mm")
    if scan:
        cv.t((PAGE_W - 56, 112), "원본", 14, GRAY, anchor="ra")
    y = 196
    box = [("갑 (고객)", party_a + (f" · 전화 {phone_a}" if phone_a else "")),
           ("을 (공급자)", "도담수납 (대표 최민재) · 가상시 예시구 샘플로 10"),
           ("계약 대상", unit_desc), ("설치 장소", place), ("계약일", date_txt)]
    for k, v in box:
        cv.t((56, y), k, 15, GRAY)
        cv.t((170, y - 1), v, 17, INK)
        y += 34
    y += 8
    d.rectangle((56, y, PAGE_W - 56, y + 40), fill=LIGHT)
    if total:
        cv.t((72, y + 20), "계약 금액 (부가세 포함)", 17, INK, anchor="lm")
        cv.t((PAGE_W - 72, y + 20), f"{won(total)}원", 24, BRAND, bold=True, anchor="rm")
    else:
        cv.t((72, y + 20), "계약 금액 (부가세 포함)", 17, INK, anchor="lm")
        cv.t((PAGE_W - 72, y + 20), "협의 후 기재", 20, GRAY, bold=True, anchor="rm")
    y += 56
    if total:
        parts = [("계약금 30%", "계약 체결 시", total * 30 // 100),
                 ("중도금 40%", "제작 착수 시", total * 40 // 100),
                 ("잔금 30%", "설치 완료 후", total - total * 30 // 100 - total * 40 // 100)]
        for k, cond, amt in parts:
            cv.t((72, y), k, 16, INK, bold=True)
            cv.t((230, y), cond, 16, GRAY)
            cv.t((PAGE_W - 72, y), f"{won(amt)}원", 17, INK, anchor="ra")
            y += 30
    else:
        for k, cond in [("계약금 30%", "계약 체결 시"), ("중도금 40%", "제작 착수 시"), ("잔금 30%", "설치 완료 후")]:
            cv.t((72, y), k, 16, INK, bold=True)
            cv.t((230, y), cond, 16, GRAY)
            cv.t((PAGE_W - 72, y), "금액 미정", 16, GRAY, anchor="ra")
            y += 30
    y += 16
    clauses = [
        "제1조(목적) 을은 갑의 세대에 주문 제작 붙박이장을 제작하고 설치하며, 갑은 그 대금을 지급한다.",
        f"제2조(계약 범위) 규격 · 마감 · 수량은 {quote_ref}에 따른다.",
        "제3조(대금 지급) 계약금은 계약 체결 시, 중도금은 제작 착수 시, 잔금은 설치 완료 후 지급한다.",
        "제4조(설치 일정) 설치 일정은 계약금 입금 후 갑과 을이 협의해 정한다.",
        "제5조(하자 보수) 설치 완료일부터 1년 안에 생긴 제작·시공 하자는 을이 무상으로 보수한다.",
        "제6조(해지) 갑이 제작 착수 뒤에 계약을 해지하면 계약금은 반환하지 않는다.",
    ]
    for c in clauses:
        for i, ln in enumerate(cv.wrap(c, 15, PAGE_W - 130)):
            cv.t((72 + (0 if i == 0 else 14), y), ln, 15, INK)
            y += 25
        y += 6
    y = max(y + 40, 850)
    for k, x in (("갑", 120), ("을", 440)):
        d.line((x, y + 50, x + 230, y + 50), fill=INK, width=1)
        cv.t((x, y + 62), ("갑 " if k == "갑" else "을 ") + ("(서명)" if k == "갑" else "(인)"), 15, GRAY)
    if scan:
        _scribble(cv, 132, y + 22, seed)
        _stamp(cv, 560, y + 28)
    elif not draft:
        cv.t((120, y + 14), "전자 서명 완료", 14, BRAND)
        cv.t((440, y + 14), "전자 직인 완료", 14, BRAND)
    _footer(cv)
    im = cv.im
    if draft:
        layer = Image.new("RGBA", (PAGE_W, PAGE_H), (0, 0, 0, 0))
        ld = ImageDraw.Draw(layer)
        ld.text((PAGE_W // 2, PAGE_H // 2), "초안", font=font(260), fill=(120, 120, 120, 60), anchor="mm")
        layer = layer.rotate(32, resample=Image.BICUBIC)
        base = im.convert("RGBA")
        base.alpha_composite(layer)
        im = base.convert("RGB")
        cv.texts.append("초안")
    if scan:
        im = im.rotate(0.6, resample=Image.BICUBIC, fillcolor=(236, 234, 228))
    return cv, im


def contract_blob(day, title_meta, **k) -> Blob:
    cv, im = render_contract(**k)
    return Blob(pdf_bytes(im, title_meta, day), cv.texts)


def render_receipt(store, tel, date_txt, items, total, payment="카드 승인", buyer="도담수납") -> Cv:
    s = sum(q * p for (_, q, p) in items)
    assert s == total, (store, s, total)
    W, H = 560, 760
    cv = Cv(W, H, (253, 252, 248))
    d = cv.d
    cv.t((W // 2, 50), store, 30, INK, bold=True, anchor="mm")
    cv.t((W // 2, 88), "영 수 증", 20, GRAY, anchor="mm")
    y = 124
    info = [("등록번호", "000-00-00000"), ("전화", tel), ("일시", date_txt), ("구매자", buyer)]
    for k, v in info:
        if v is None:
            continue
        cv.t((40, y), k, 15, GRAY)
        cv.t((140, y), v, 16, INK)
        y += 28
    y += 8
    d.line((36, y, W - 36, y), fill=INK, width=2)
    y += 10
    cv.t((40, y), "품목", 15, GRAY)
    cv.t((300, y), "수량", 15, GRAY, anchor="ra")
    cv.t((400, y), "단가", 15, GRAY, anchor="ra")
    cv.t((W - 40, y), "금액", 15, GRAY, anchor="ra")
    y += 30
    d.line((36, y, W - 36, y), fill=LINE, width=1)
    y += 12
    for name, q, p in items:
        cv.t((40, y), name, 16, INK)
        cv.t((300, y), str(q), 16, INK, anchor="ra")
        cv.t((400, y), won(p), 16, INK, anchor="ra")
        cv.t((W - 40, y), won(q * p), 16, INK, anchor="ra")
        y += 34
    y += 6
    d.line((36, y, W - 36, y), fill=INK, width=2)
    y += 18
    cv.t((40, y), "합계", 20, INK, bold=True)
    cv.t((W - 40, y - 2), f"{won(total)}원", 24, INK, bold=True, anchor="ra")
    y += 46
    cv.t((40, y), payment, 15, GRAY)
    cv.t((W - 40, y), "부가세 포함", 15, GRAY, anchor="ra")
    cv.t((W // 2, H - 48), "이용해 주셔서 감사합니다", 15, GRAY, anchor="mm")
    return cv


def receipt_pdf(day, title, **k) -> Blob:
    cv = render_receipt(**k)
    return Blob(pdf_bytes(cv.im, title, day), cv.texts)


def receipt_photo(day, **k) -> Blob:
    """영수증 종이를 나무색 바닥 위에 비스듬히 놓은 합성 사진(JPG)."""
    cv = render_receipt(**k)
    W, H = 800, 600
    bg = Image.new("RGB", (W, H), (150, 112, 78))
    bd = ImageDraw.Draw(bg)
    for yy in range(0, H, 46):
        bd.line((0, yy, W, yy), fill=(132, 96, 64), width=2)
    paper = cv.im.resize((400, 543))
    paper = paper.convert("RGBA")
    paper = paper.rotate(4, resample=Image.BICUBIC, expand=True)
    shadow = Image.new("RGBA", paper.size, (0, 0, 0, 0))
    mask = paper.split()[3]
    shadow.paste((0, 0, 0, 110), (0, 0), mask)
    shadow = shadow.filter(ImageFilter.GaussianBlur(7))
    px, py = (W - paper.width) // 2, (H - paper.height) // 2
    bg.paste(shadow.convert("RGB"), (px + 8, py + 10), shadow.split()[3])
    bg.paste(paper.convert("RGB"), (px, py), paper.split()[3])
    return Blob(jpg_bytes(bg), cv.texts + ["합성 이미지"])


def receipt_capture(**k) -> Blob:
    """결제 내역 화면을 캡처한 모양의 합성 이미지(PNG)."""
    store, date_txt, items, total = k["store"], k["date_txt"], k["items"], k["total"]
    cv = Cv(800, 600, (240, 242, 245))
    d = cv.d
    d.rectangle((0, 0, 800, 70), fill=BRAND)
    cv.t((30, 35), "결제 내역", 26, WHITE, bold=True, anchor="lm")
    cv.t((770, 35), "12:48", 18, WHITE, anchor="rm")
    d.rounded_rectangle((60, 100, 740, 520), radius=18, fill=WHITE, outline=LINE, width=2)
    cv.t((100, 140), store, 30, INK, bold=True)
    cv.t((700, 148), "결제 완료", 20, (40, 140, 80), bold=True, anchor="ra")
    cv.t((100, 196), "결제 일시", 16, GRAY)
    cv.t((230, 194), date_txt, 18, INK)
    cv.t((100, 236), "결제 수단", 16, GRAY)
    cv.t((230, 234), "법인 카드", 18, INK)
    d.line((100, 286, 700, 286), fill=LINE, width=1)
    y = 306
    for name, q, p in items:
        cv.t((100, y), f"{name}", 19, INK)
        cv.t((560, y), f"{q} x {won(p)}", 18, GRAY, anchor="ra")
        cv.t((700, y), f"{won(q * p)}원", 19, INK, anchor="ra")
        y += 44
    d.line((100, 440, 700, 440), fill=LINE, width=1)
    cv.t((100, 462), "결제 금액", 20, INK, bold=True)
    cv.t((700, 458), f"{won(total)}원", 32, BRAND, bold=True, anchor="ra")
    cv.t((400, 560), "화면 캡처 · 합성 이미지", 14, GRAY, anchor="mm")
    return Blob(png_bytes(cv.im), cv.texts)


# ------------------------------------------------------------------ 도면
def _title_block(cv: Cv, x0, y0, x1, y1, title, sub, who, date_txt):
    d = cv.d
    d.rectangle((x0, y0, x1, y1), outline=INK, width=2)
    d.line((x0, y0 + (y1 - y0) // 2, x1, y0 + (y1 - y0) // 2), fill=INK, width=1)
    mid = y0 + (y1 - y0) // 4
    cv.t((x0 + 14, mid), title, 20, INK, bold=True, anchor="lm")
    cv.t((x1 - 14, mid), sub, 15, GRAY, anchor="rm")
    mid2 = y0 + 3 * (y1 - y0) // 4
    cv.t((x0 + 14, mid2), who, 15, INK, anchor="lm")
    cv.t((x1 - 14, mid2), date_txt, 15, INK, anchor="rm")


def plan_png(kind: str) -> Blob:
    cv = Cv(800, 600, (252, 252, 250))
    d = cv.d
    if kind == "84":
        ox0, oy0, ox1, oy1 = 60, 60, 740, 450
        d.rectangle((ox0, oy0, ox1, oy1), outline=INK, width=4)
        d.line((400, oy0, 400, oy1), fill=INK, width=3)
        d.line((400, 255, ox1, 255), fill=INK, width=3)
        d.rectangle((396, 330, 404, 392), fill=(252, 252, 250))  # 문 틈
        d.line((400, 330, 400 + 60, 330 + 0), fill=GRAY, width=1)
        d.arc((340, 330, 460, 450), 270, 360, fill=GRAY, width=1)
        cv.hatch((420, 64, 700, 106), LINE, 7, WHITE)
        cv.t((560, 85), "붙박이장", 16, INK, anchor="mm")
        cv.t((230, 255), "거실", 26, INK, bold=True, anchor="mm")
        cv.t((570, 175), "침실 1", 22, INK, bold=True, anchor="mm")
        cv.t((570, 355), "침실 2", 22, INK, bold=True, anchor="mm")
        cv.dim_h(420, 700, 40, "3,600")
        cv.dim_h(60, 740, 464, "9,400")
        cv.dim_v(46, 64, 106, "600")
        label = ("도면 84형", "84형 붙박이장 평면", "설계 디자인팀 이서연", "2026-03-04")
    else:
        ox0, oy0, ox1, oy1 = 130, 70, 670, 440
        d.rectangle((ox0, oy0, ox1, oy1), outline=INK, width=4)
        d.line((380, oy0, 380, oy1), fill=INK, width=3)
        d.line((ox0, 280, 380, 280), fill=INK, width=3)
        d.rectangle((376, 330, 384, 392), fill=(252, 252, 250))
        d.arc((320, 330, 440, 450), 270, 360, fill=GRAY, width=1)
        cv.hatch((134, 74, 176, 270), LINE, 7, WHITE)
        cv.t((184, 172), "붙박이장", 14, INK, anchor="lm")
        cv.t((255, 360), "거실", 24, INK, bold=True, anchor="mm")
        cv.t((300, 175), "안방", 22, INK, bold=True, anchor="mm")
        cv.t((525, 255), "주방·침실", 22, INK, bold=True, anchor="mm")
        cv.dim_v(110, 74, 270, "2,400")
        cv.dim_h(130, 670, 464, "7,600")
        label = ("도면 59형", "59형 붙박이장 평면", "설계 디자인팀 이서연", "2026-03-05")
    _title_block(cv, 40, 500, 760, 580, *label)
    cv.t((760, 486), "축척 1:50 (도식)", 13, GRAY, anchor="rs")
    cv.t((40, 486), "단위 mm", 13, GRAY, anchor="ls")
    return Blob(png_bytes(cv.im), cv.texts)


def plan_pdf(kind: str, day: int) -> Blob:
    W, H = 1123, 794
    cv = Cv(W, H, (252, 252, 250))
    d = cv.d
    if kind == "A":
        site, where, w_mm, doors = "A현장", "예시동 A현장", "3,000", 3
    else:
        site, where, w_mm, doors = "C현장", "가상읍 C현장", "2,400", 2
    cv.t((50, 36), f"{site} 시공 도면", 30, INK, bold=True)
    cv.t((W - 50, 44), where, 18, GRAY, anchor="ra")
    d.line((50, 76, W - 50, 76), fill=BRAND, width=3)
    # 평면
    cv.t((50, 100), "평면", 18, BRAND, bold=True)
    px0, py0, px1, py1 = 70, 160, 520, 500
    d.rectangle((px0, py0, px1, py1), outline=INK, width=4)
    d.line((300, py0, 300, py1), fill=INK, width=3)
    d.rectangle((296, 400, 304, 460), fill=(252, 252, 250))
    cv.hatch((px0 + 4, py0 + 4, px1 - 4 if kind == "A" else 300 - 4, py0 + 48), LINE, 7, WHITE)
    cv.t(((px0 + (px1 if kind == "A" else 300)) // 2, py0 + 26), "붙박이장", 15, INK, anchor="mm")
    cv.t((185, 330), "방", 24, INK, bold=True, anchor="mm")
    cv.t((410, 330), "거실", 24, INK, bold=True, anchor="mm")
    cv.dim_h(px0, px1 if kind == "A" else 300, 130, w_mm)
    cv.dim_h(px0, px1, 530, "6,000" if kind == "A" else "5,400")
    # 입면
    ex0, ey0 = 580, 160
    cv.t((ex0 - 20, 100), "입면", 18, BRAND, bold=True)
    ew = 420
    eh = 340
    d.rectangle((ex0, ey0, ex0 + ew, ey0 + eh), outline=INK, width=4)
    pw = ew // doors
    for i in range(doors):
        x = ex0 + i * pw
        d.rectangle((x + 4, ey0 + 4, x + pw - 4, ey0 + eh - 4), outline=INK, width=2, fill=(244, 240, 232))
        hx = x + pw - 22 if i % 2 == 0 else x + 22
        d.rectangle((hx - 3, ey0 + 150, hx + 3, ey0 + 210), fill=GRAY)
    cv.dim_h(ex0, ex0 + ew, ey0 - 18, w_mm)
    cv.dim_v(ex0 + ew + 70, ey0, ey0 + eh, "2,350")
    cv.t((ex0, ey0 + eh + 30), f"도어 {doors}틈 · 마감 화이트", 16, INK)
    _title_block(cv, 50, 620, W - 50, 730, f"{site} 시공 도면", "축척 1:50 (도식)",
                 "시공팀장 박도윤 · 010-0000-0013", f"2026-03-{day:02d}")
    return Blob(pdf_bytes(cv.im, f"{site} 시공 도면", day), cv.texts)


# ------------------------------------------------------------------ 합성 사진 (JPG·PNG)
def _caption(cv: Cv, text: str):
    cv.d.rectangle((0, 540, 800, 600), fill=(33, 37, 41))
    cv.t((24, 570), text, 26, WHITE, anchor="lm")
    cv.t((776, 570), "합성 이미지", 14, (190, 196, 202), anchor="rm")


def _windows(cv: Cv, x0, y0, x1, y1, cols, rows, lit=(255, 240, 180), dark=(150, 185, 215)):
    cw = (x1 - x0) // cols
    rh = (y1 - y0) // rows
    for r in range(rows):
        for c in range(cols):
            col = lit if (r * 3 + c * 5) % 4 == 0 else dark
            cv.d.rectangle((x0 + c * cw + 8, y0 + r * rh + 8, x0 + (c + 1) * cw - 8, y0 + (r + 1) * rh - 8), fill=col)


def scene_exterior() -> Cv:
    cv = Cv(800, 600, (206, 228, 244))
    d = cv.d
    d.ellipse((664, 44, 736, 116), fill=(250, 222, 120))
    d.rectangle((0, 430, 800, 540), fill=(150, 176, 132))
    d.rectangle((0, 470, 800, 510), fill=(96, 100, 106))
    for x in range(10, 800, 70):
        d.rectangle((x, 488, x + 34, 492), fill=(235, 235, 225))
    for (x0, y0, x1, y1, col) in [(60, 170, 290, 440, (214, 204, 190)), (320, 110, 560, 440, (196, 204, 214)),
                                   (590, 210, 760, 440, (220, 210, 196))]:
        d.rectangle((x0, y0, x1, y1), fill=col, outline=INK, width=2)
        _windows(cv, x0 + 4, y0 + 10, x1 - 4, y1 - 50, 4 if x1 - x0 > 200 else 3, 5 if y1 - y0 > 300 else 4)
    for x in range(320, 561, 40):  # 비계
        d.line((x, 110, x, 440), fill=ACCENT, width=2)
    for y in range(140, 441, 50):
        d.line((320, y, 560, y), fill=ACCENT, width=2)
    return cv


def _room(cv: Cv):
    d = cv.d
    d.rectangle((0, 0, 800, 400), fill=(228, 222, 212))
    d.rectangle((0, 400, 800, 540), fill=(186, 156, 122))
    for x in range(0, 800, 100):
        d.line((x, 400, x - 40 if x else 0, 540), fill=(166, 138, 106), width=2)
    d.rectangle((0, 384, 800, 400), fill=(245, 242, 236))
    d.rectangle((70, 80, 270, 300), fill=(176, 208, 232), outline=(245, 242, 236), width=8)
    d.line((170, 80, 170, 300), fill=(245, 242, 236), width=6)
    d.line((70, 190, 270, 190), fill=(245, 242, 236), width=6)


def scene_room_before() -> Cv:
    cv = Cv(800, 600)
    _room(cv)
    cv.d.rectangle((400, 50, 750, 384), fill=(206, 200, 190), outline=(180, 172, 160), width=6)
    return cv


def scene_room_after() -> Cv:
    cv = Cv(800, 600)
    _room(cv)
    d = cv.d
    d.rectangle((400, 50, 750, 392), fill=(243, 241, 237), outline=(170, 166, 158), width=4)
    pw = 350 // 3
    for i in range(3):
        x = 400 + i * pw
        d.rectangle((x + 5, 56, x + pw - 5, 386), outline=(190, 186, 178), width=2)
        hx = x + pw - 20 if i % 2 == 0 else x + 20
        d.rectangle((hx - 3, 190, hx + 3, 250), fill=(90, 96, 104))
    d.rectangle((400, 392, 750, 404), fill=(150, 146, 138))
    return cv


def scene_measure() -> Cv:
    cv = Cv(800, 600, (232, 228, 220))
    d = cv.d
    d.rectangle((0, 400, 800, 540), fill=(186, 156, 122))
    d.rectangle((120, 90, 680, 390), outline=(120, 116, 108), width=8, fill=(240, 237, 230))
    # 줄자
    d.rectangle((120, 420, 680, 452), fill=(245, 205, 60), outline=INK, width=2)
    for x in range(130, 680, 20):
        d.line((x, 420, x, 432 if (x // 20) % 5 else 442), fill=INK, width=1)
    cv.dim_h(120, 680, 70, "2,400")
    cv.dim_v(96, 90, 390, "2,350")
    # 기록판
    d.rectangle((560, 470, 760, 534), fill=(250, 250, 246), outline=INK, width=2)
    for yy in (486, 502, 518):
        d.line((572, yy, 748, yy), fill=LINE, width=2)
    return cv


def scene_materials() -> Cv:
    cv = Cv(800, 600, (206, 206, 200))
    d = cv.d
    d.rectangle((0, 380, 800, 540), fill=(150, 150, 144))
    for (x0, w, n) in [(300, 200, 7), (540, 200, 5)]:
        for i in range(n):
            y = 380 - (i + 1) * 26
            d.rectangle((x0, y, x0 + w, y + 24), fill=(198, 162, 112), outline=(120, 92, 60), width=2)
    d.rectangle((30, 300, 250, 420), fill=(236, 236, 232), outline=INK, width=3)
    d.rectangle((250, 330, 300, 420), fill=BRAND, outline=INK, width=3)
    for cx in (80, 210, 275):
        d.ellipse((cx - 22, 402, cx + 22, 446), fill=INK)
        d.ellipse((cx - 9, 415, cx + 9, 433), fill=LIGHT)
    return cv


def scene_entrance() -> Cv:
    cv = Cv(800, 600, (232, 226, 216))
    d = cv.d
    d.rectangle((0, 420, 800, 540), fill=(196, 176, 150))
    d.rectangle((250, 70, 550, 420), fill=(122, 90, 62), outline=(84, 60, 40), width=8)
    d.rectangle((280, 100, 520, 250), outline=(92, 66, 44), width=4)
    d.rectangle((280, 280, 520, 390), outline=(92, 66, 44), width=4)
    d.ellipse((500, 250, 524, 274), fill=(214, 190, 100))
    d.rectangle((300, 440, 500, 500), fill=(84, 96, 110))
    d.ellipse((380, 30, 420, 60), fill=(255, 244, 190))
    return cv


def scene_cabinet() -> Cv:
    cv = Cv(800, 600, (232, 226, 216))
    d = cv.d
    d.rectangle((0, 450, 800, 540), fill=(196, 176, 150))
    d.rectangle((110, 70, 690, 450), fill=(242, 240, 236), outline=(150, 146, 138), width=5)
    cw, rh = 580 // 4, 380 // 3
    for r in range(3):
        for c in range(4):
            x0, y0 = 110 + c * cw, 70 + r * rh
            d.rectangle((x0 + 5, y0 + 5, x0 + cw - 5, y0 + rh - 5), outline=(190, 186, 178), width=2)
            d.rectangle((x0 + cw // 2 - 14, y0 + rh - 26, x0 + cw // 2 + 14, y0 + rh - 20), fill=(90, 96, 104))
    return cv


def scene_sitemap() -> Cv:
    cv = Cv(800, 600, (214, 232, 205))
    d = cv.d
    d.rectangle((0, 250, 800, 290), fill=(200, 200, 196))
    d.rectangle((380, 0, 420, 540), fill=(200, 200, 196))
    cv.t((400, 28), "샘플하이츠 동 배치", 26, INK, bold=True, anchor="mm")
    blocks = [("101동", 70, 70, 330, 220), ("102동", 450, 70, 730, 220),
              ("103동", 70, 320, 330, 470), ("104동", 450, 320, 590, 470), ("105동", 600, 320, 730, 470)]
    for name, x0, y0, x1, y1 in blocks:
        d.rectangle((x0, y0, x1, y1), fill=(232, 230, 224), outline=INK, width=3)
        cv.t(((x0 + x1) // 2, (y0 + y1) // 2), name, 26, INK, bold=True, anchor="mm")
    d.rectangle((356, 494, 444, 534), fill=BRAND)
    cv.t((400, 514), "정문", 18, WHITE, anchor="mm")
    cv.t((764, 70), "N", 22, INK, bold=True, anchor="mm")
    d.polygon([(764, 88), (754, 112), (774, 112)], fill=INK)
    return cv


def photo(scene: Cv, caption: str, jpg=True) -> Blob:
    _caption(scene, caption)
    scene.texts.append("합성 이미지")
    return Blob(jpg_bytes(scene.im) if jpg else png_bytes(scene.im), scene.texts)


# ------------------------------------------------------------------ 사업자등록증 · 보험증서
def business_reg() -> Blob:
    cv = Cv(800, 600, (250, 250, 244))
    d = cv.d
    d.rectangle((20, 20, 780, 580), outline=(60, 90, 70), width=4)
    cv.t((400, 66), "사업자등록증 (사본)", 34, (40, 70, 55), bold=True, anchor="mm")
    cv.t((400, 106), "등록번호 000-00-00000", 22, INK, anchor="mm")
    d.line((60, 130, 740, 130), fill=(60, 90, 70), width=2)
    rows = [("상호", "도담수납"), ("대표자", "최민재"), ("개업연월일", "2019년 4월 1일"),
            ("사업장 소재지", "가상시 예시구 샘플로 10"), ("업태 · 종목", "제조업·건설업 · 붙박이장 제조·시공"),
            ("발급 기관", "가상세무서")]
    y = 160
    for k, v in rows:
        cv.t((70, y), k, 19, GRAY)
        cv.t((260, y - 1), v, 21, INK)
        y += 52
    cv.t((400, 520), "이 문서는 실습용으로 만든 합성 이미지입니다", 15, GRAY, anchor="mm")
    _stamp(cv, 660, 480, 34)
    return Blob(png_bytes(cv.im), cv.texts)


def insurance_pdf(day: int) -> Blob:
    cv = Cv(PAGE_W, PAGE_H, (252, 252, 247))
    d = cv.d
    d.rectangle((30, 30, PAGE_W - 30, PAGE_H - 30), outline=(60, 90, 70), width=4)
    d.rectangle((40, 40, PAGE_W - 40, PAGE_H - 40), outline=(60, 90, 70), width=1)
    cv.t((PAGE_W // 2, 130), "시공 배상 책임 보험 가입 증서", 34, (40, 70, 55), bold=True, anchor="mm")
    cv.t((PAGE_W // 2, 178), "증권번호 0000-0000", 18, GRAY, anchor="mm")
    rows = [("보험사", "샘플보험 (가상)"), ("피보험자", "도담수납 (대표 최민재)"),
            ("사업장", "가상시 예시구 샘플로 10"), ("담보 업종", "붙박이장·수납가구 제작 및 설치"),
            ("보험 기간", "2026-01-01 ~ 2026-12-31"), ("보상 한도", "1사고당 100,000,000원"),
            ("자기 부담금", "500,000원"), ("발행일", "2026-01-05")]
    y = 260
    for k, v in rows:
        cv.t((110, y), k, 19, GRAY)
        cv.t((300, y - 1), v, 22, INK)
        d.line((100, y + 40, PAGE_W - 100, y + 40), fill=LIGHT, width=1)
        y += 66
    cv.t((PAGE_W // 2, 880), "위 보험 계약이 유효함을 증명합니다.", 20, INK, anchor="mm")
    _stamp(cv, 600, 960, 40)
    cv.t((PAGE_W // 2, 1040), "이 문서는 실습용으로 만든 가상의 증서입니다", 14, GRAY, anchor="mm")
    return Blob(pdf_bytes(cv.im, "시공 배상 책임 보험 가입 증서", day), cv.texts)


# ------------------------------------------------------------------ XLSX
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HFILL = PatternFill("solid", start_color="DDEBF7", end_color="DDEBF7")

STOCK = [  # 품목, 규격, 3/5 수량, 3/12 수량, 단위
    ("합판", "18T 1220x2440", 42, 34, "장"),
    ("합판", "12T 1220x2440", 25, 25, "장"),
    ("소프트클로징 경첩", "110도", 180, 140, "개"),
    ("슬라이딩 레일", "2.4m", 4, 10, "세트"),
    ("스틸 손잡이", "128mm", 30, 70, "개"),
    ("서랍 레일", "450mm", 24, 18, "조"),
    ("선반 지지대", "5mm", 500, 500, "개"),
    ("목공용 접착제", "3.75kg", 6, 4, "통"),
    ("실리콘", "투명", 20, 20, "개"),
    ("피스", "3.5x30", 12, 9, "박스"),
    ("LED 바", "600mm", 16, 11, "개"),
    ("모서리 마감재", "화이트 2.4m", 40, 40, "개"),
]
PRICES = [  # 품목, 규격, 단위, 단가, 거래처
    ("합판", "18T 1220x2440", "장", 38000, "예시목재"),
    ("합판", "12T 1220x2440", "장", 29000, "예시목재"),
    ("소프트클로징 경첩", "110도", "개", 1800, "예시목재"),
    ("슬라이딩 레일", "2.4m", "세트", 42000, "샘플철물"),
    ("스틸 손잡이", "128mm", "개", 3500, "샘플철물"),
    ("서랍 레일", "450mm", "조", 7500, "샘플철물"),
    ("선반 지지대", "5mm", "개", 120, "샘플철물"),
    ("목공용 접착제", "3.75kg", "통", 26000, "예시공구"),
    ("실리콘", "투명", "개", 12000, "예시공구"),
    ("피스", "3.5x30", "박스", 9500, "예시공구"),
    ("LED 바", "600mm", "개", 15000, "샘플조명"),
    ("모서리 마감재", "화이트 2.4m", "개", 4800, "예시목재"),
    ("거울", "600x1800", "장", 85000, "가상유리"),
    ("강화유리 선반", "8T 300x600", "장", 22000, "가상유리"),
    ("전원 어댑터", "12V", "개", 8500, "샘플조명"),
]


def _sheet(ws, header, rows, widths, num_cols=()):
    ws.append(header)
    for r in rows:
        ws.append(list(r))
    for c, w in zip("ABCDEFGH", widths):
        ws.column_dimensions[c].width = w
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HFILL
        cell.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=len(header)):
        for cell in row:
            cell.border = BORDER
            if cell.column in num_cols and cell.row > 1:
                cell.number_format = "#,##0"


def xlsx_quote_form() -> Blob:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "견적서 양식"
    _sheet(ws, ["번호", "품목", "규격", "수량", "단가", "금액"],
           [(i, None, None, None, None, None) for i in range(1, 11)], [8, 30, 18, 10, 14, 16], (5, 6))
    wb.properties.title = "견적서 양식"
    return Blob(xlsx_bytes(wb, 2))


def xlsx_stock(which: int) -> Blob:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "3월5일 기준" if which == 0 else "3월12일 기준"
    rows = [(n, s, (a if which == 0 else b), u) for (n, s, a, b, u) in STOCK]
    _sheet(ws, ["품목", "규격", "수량", "단위"], rows, [22, 18, 10, 8], (3,))
    wb.properties.title = "재고현황 " + ws.title
    return Blob(xlsx_bytes(wb, 5 if which == 0 else 12))


def xlsx_prices() -> Blob:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "자재 단가표"
    _sheet(ws, ["품목", "규격", "단위", "단가", "거래처"], PRICES, [22, 18, 8, 12, 12], (4,))
    wb.properties.title = "자재 단가표"
    return Blob(xlsx_bytes(wb, 3))


# ------------------------------------------------------------------ 텍스트 파일
AGENTS_MD = """# 받은자료 폴더 정리 규칙

이 폴더의 파일을 정리할 때 아래 규칙을 따른다.

1. `삭제금지` 폴더 안의 파일은 옮기지 않고, 이름을 바꾸지 않고, 지우지 않는다.
2. 파일을 지우지 않는다. 남길 필요가 없어 보이는 파일은 `보류` 폴더로 옮긴다.
3. 파일 이름은 바꾸지 않는다.
4. 분류 폴더는 아래 여섯 개만 쓴다.
   - 01_견적서
   - 02_계약서
   - 03_현장사진
   - 04_영수증
   - 05_도면
   - 06_기타
5. 파일을 옮길 때마다 `이동기록.md`에 「원래 위치 → 새 위치 | 이유」를 한 줄씩 적는다.
6. 같은 파일이 둘 이상이면 하나만 분류 폴더에 두고 나머지는 `보류`로 옮긴다.
"""
CLAUDE_MD = "@AGENTS.md\n"

TXT_CONTACT = """# 거래처 연락처 (정하윤 정리)
예시목재 (합판·판재·경첩) 010-0000-0051
샘플철물 (레일·손잡이·지지대) 010-0000-0052
가상운송 (자재 운송) 010-0000-0053
예시공구 (공구·접착제·실리콘) 010-0000-0054
샘플조명 (LED 조명·전원) 010-0000-0055
가상유리 (거울·유리 가공) 010-0000-0056
A/S 문의 메일 as@example.com
"""
TXT_TODO = """정하윤 할 일 (3월)
- 창고 재고 수량 다시 세기
- 3월 자재 영수증 모아서 정리
- 송씨댁 계약금 입금 확인
- 노씨댁 계약금 입금 확인
- 협력업체 단가표 최신본 받기
- 사내 공지 전달 확인
- 공구 구입 영수증을 장부에 반영
- 현장 식대 영수증 사진 보관
- 샘플하이츠 단체 계약서 초안 검토 요청
"""
TXT_PARTNERS = """협력업체 목록 (정하윤 정리)

[자재]
예시목재 - 합판·판재·경첩 공급
샘플철물 - 레일·손잡이·지지대 공급
샘플조명 - LED 조명·전원 공급

[물류·현장]
가상운송 - 자재 운송
예시공구 - 공구·접착제·실리콘 판매
가상유리 - 거울·유리 가공
총 6곳
"""
TXT_NOTICE = """[사내 공지] 2026년 3월 10일
제목: 자재 창고 재고 조사 협조 안내

경영지원팀 정하윤입니다.
자재 창고의 재고 수량을 다시 확인합니다.
시공팀은 창고에 둔 개인 물품을 미리 치워 주세요.
자재를 꺼낼 때는 꺼낸 수량을 창고 앞 메모판에 적어 주세요.
문의는 경영지원팀 정하윤(010-0000-0015)에게 연락해 주세요.
감사합니다.
"""
TXT_TEMP = "test\n\n"  # 2줄: 「test」 한 줄 + 빈 줄


# ------------------------------------------------------------------ 자료 40개 표
# (키, 파일 이름, 형식, 정답 폴더, 사유(이동기록), 수정일(3월 d일), 빌더)
def _q(day, title, *a, **k):
    return lambda: render_quote_blob(day, title, *a, **k)


def _build_table():
    T: list[tuple] = []

    def add(key, name, fmt, folder, reason, day, fn):
        T.append((key, name, fmt, folder, reason, day, fn))

    Q_ITEMS_84 = [("붙박이장 본체(합판 18T)", "폭 3,600mm", 1, "식", 3600000),
                  ("슬라이딩 도어", "3틈", 3, "틈", 700000),
                  ("내부 선반·서랍", "선반 6 · 서랍 3", 1, "식", 900000),
                  ("레일·손잡이·경첩", "", 1, "식", 300000),
                  ("운반·설치", "", 1, "식", 300000)]
    Q_ITEMS_59 = [("붙박이장 본체(합판 18T)", "폭 2,400mm", 1, "식", 1900000),
                  ("슬라이딩 도어", "2틈", 2, "틈", 450000),
                  ("내부 선반·서랍", "선반 4 · 서랍 2", 1, "식", 500000),
                  ("레일·손잡이·경첩", "", 1, "식", 150000),
                  ("운반·설치", "", 1, "식", 200000)]
    Q_ITEMS_SEO = [("붙박이장 본체(합판 18T)", "폭 3,000mm", 1, "식", 2300000),
                   ("슬라이딩 도어", "2틈", 2, "틈", 600000),
                   ("내부 선반·서랍", "선반 5 · 서랍 2", 1, "식", 600000),
                   ("레일·손잡이·경첩", "", 1, "식", 200000),
                   ("운반·설치", "", 1, "식", 200000)]
    Q_ITEMS_YOO = [("드레스룸 프레임·선반", "폭 3,200mm", 1, "식", 3200000),
                   ("서랍장", "3단", 3, "개", 600000),
                   ("행거 시스템", "2단 행거", 1, "식", 1200000),
                   ("조명·부자재", "", 1, "식", 400000),
                   ("운반·설치", "", 1, "식", 400000)]
    Q_ITEMS_HAN = [("붙박이장 본체(합판 18T)", "폭 2,700mm", 1, "식", 2200000),
                   ("슬라이딩 도어", "2틈", 2, "틈", 550000),
                   ("내부 선반·서랍", "선반 4 · 서랍 2", 1, "식", 500000),
                   ("레일·손잡이·경첩", "", 1, "식", 150000),
                   ("운반·설치", "", 1, "식", 200000)]
    Q_ITEMS_MUN = [("신발장 본체", "폭 1,800mm", 1, "식", 2600000),
                   ("벤치·거울", "", 1, "식", 1000000),
                   ("조명", "LED", 1, "식", 700000),
                   ("손잡이·경첩·부자재", "", 1, "식", 500000),
                   ("운반·설치", "", 1, "식", 600000)]

    # 삭제금지 4개
    add("L1", "계약서_원본스캔_송씨댁.pdf", "PDF", "삭제금지", "", 10, lambda: contract_blob(
        10, "계약서 원본 스캔 송씨댁", title="붙박이장 제작·설치 계약서", party_a="송씨댁",
        unit_desc="붙박이장 84형 1식", place="샘플하이츠 84형 세대", total=7200000,
        date_txt="2026년 3월 10일", quote_ref="견적서 2026-031 (3월 3일)", scan=True, seed=11))
    add("L2", "계약서_원본스캔_노씨댁.pdf", "PDF", "삭제금지", "", 11, lambda: contract_blob(
        11, "계약서 원본 스캔 노씨댁", title="붙박이장 제작·설치 계약서", party_a="노씨댁",
        unit_desc="붙박이장 59형 1식", place="샘플하이츠 59형 세대", total=3650000,
        date_txt="2026년 3월 11일", quote_ref="견적서 2026-032 (3월 4일)", scan=True, seed=12))
    add("L3", "사업자등록증_사본.png", "PNG", "삭제금지", "", 2, business_reg)
    add("L4", "보험가입증서.pdf", "PDF", "삭제금지", "", 2, lambda: insurance_pdf(2))

    # 견적서 7
    add(1, "견적서_송씨댁_0303.pdf", "PDF", "01_견적서", "견적서", 3,
        _q(3, "견적서 송씨댁", "2026-031", "2026년 3월 3일", "송씨댁", "붙박이장 84형", Q_ITEMS_84, 7200000))
    add(2, "견적서_노씨댁_0304.pdf", "PDF", "01_견적서", "견적서", 4,
        _q(4, "견적서 노씨댁", "2026-032", "2026년 3월 4일", "노씨댁", "붙박이장 59형", Q_ITEMS_59, 3650000))
    add(3, "견적서_서씨댁_0305.pdf", "PDF", "01_견적서", "견적서", 5,
        _q(5, "견적서 서씨댁", "2026-033", "2026년 3월 5일", "서씨댁", "붙박이장 84형", Q_ITEMS_SEO, 4500000))
    add(4, "견적서_최종.pdf", "PDF", "01_견적서", "견적서", 6,
        _q(6, "견적서 최종", "2026-034", "2026년 3월 6일", "유씨댁", "드레스룸", Q_ITEMS_YOO, 7000000, final=True))
    add(5, "견적서_최종(1).pdf", "PDF", "01_견적서", "견적서", 10,
        _q(10, "견적서 최종", "2026-037", "2026년 3월 10일", "한씨댁", "붙박이장", Q_ITEMS_HAN, 4150000, final=True))
    add(6, "견적서_문씨댁_0309.pdf", "PDF", "01_견적서", "견적서", 9,
        _q(9, "견적서 문씨댁", "2026-036", "2026년 3월 9일", "문씨댁", "신발장", Q_ITEMS_MUN, 5400000))
    add(7, "견적서_양식.xlsx", "XLSX", "01_견적서", "견적 양식", 2, xlsx_quote_form)
    # 계약서 3
    add(8, "계약서_송씨댁_0310.pdf", "PDF", "02_계약서", "계약서 전자본", 10, lambda: contract_blob(
        10, "계약서 송씨댁", title="붙박이장 제작·설치 계약서", party_a="송씨댁",
        unit_desc="붙박이장 84형 1식", place="샘플하이츠 84형 세대", total=7200000,
        date_txt="2026년 3월 10일", quote_ref="견적서 2026-031 (3월 3일)"))
    add(9, "계약서_노씨댁_0311.pdf", "PDF", "02_계약서", "계약서 전자본", 11, lambda: contract_blob(
        11, "계약서 노씨댁", title="붙박이장 제작·설치 계약서", party_a="노씨댁",
        unit_desc="붙박이장 59형 1식", place="샘플하이츠 59형 세대", total=3650000,
        date_txt="2026년 3월 11일", quote_ref="견적서 2026-032 (3월 4일)"))
    add(10, "계약서_샘플하이츠_초안.pdf", "PDF", "02_계약서", "계약서 초안", 12, lambda: contract_blob(
        12, "계약서 샘플하이츠 초안", title="붙박이장 단체 제작·설치 계약서 (초안)",
        party_a="샘플하이츠 입주자 대표 배정우", phone_a="010-0000-0042",
        unit_desc="붙박이장 12세대", place="샘플하이츠 12세대", total=None,
        date_txt="협의 후 기재", quote_ref="세대별 견적서", draft=True))
    # 현장 사진 9
    add(11, "현장A_전경.jpg", "JPG", "03_현장사진", "현장 사진", 3,
        lambda: photo(scene_exterior(), "예시동 A현장 · 전경"))
    add(12, "현장A_전경 - 복사본.jpg", "JPG", "03_현장사진", "현장 사진", 4, None)  # 11번 바이트 복사
    add(13, "현장A_설치전.jpg", "JPG", "03_현장사진", "현장 사진", 9,
        lambda: photo(scene_room_before(), "예시동 A현장 · 설치 전"))
    add(14, "현장A_설치후.jpg", "JPG", "03_현장사진", "현장 사진", 11,
        lambda: photo(scene_room_after(), "예시동 A현장 · 설치 후"))
    add(15, "현장B_실측.jpg", "JPG", "03_현장사진", "현장 사진", 6,
        lambda: photo(scene_measure(), "샘플로 B현장 · 실측"))
    add(16, "현장B_자재입고.png", "PNG", "03_현장사진", "현장 사진", 6,
        lambda: photo(scene_materials(), "샘플로 B현장 · 자재 입고", jpg=False))
    add(17, "현장C_현관.jpg", "JPG", "03_현장사진", "현장 사진", 10,
        lambda: photo(scene_entrance(), "가상읍 C현장 · 현관"))
    add(18, "현장C_수납장.jpg", "JPG", "03_현장사진", "현장 사진", 10,
        lambda: photo(scene_cabinet(), "가상읍 C현장 · 수납장"))
    add(19, "샘플하이츠_동배치.png", "PNG", "03_현장사진", "현장 사진", 2,
        lambda: photo(scene_sitemap(), "샘플하이츠 · 단지 배치", jpg=False))
    # 영수증 5
    add(20, "영수증_자재_0302.pdf", "PDF", "04_영수증", "영수증", 2, lambda: receipt_pdf(
        2, "영수증 자재 0302", store="예시목재", tel="010-0000-0051", date_txt="2026-03-02 10:14",
        items=[("합판 18T 1220x2440", 20, 38000), ("소프트클로징 경첩", 100, 1800)], total=940000))
    add(21, "영수증_자재_0309.pdf", "PDF", "04_영수증", "영수증", 9, lambda: receipt_pdf(
        9, "영수증 자재 0309", store="샘플철물", tel="010-0000-0052", date_txt="2026-03-09 15:32",
        items=[("슬라이딩 레일 2.4m", 6, 42000), ("스틸 손잡이 128mm", 40, 3500)], total=392000))
    add(22, "영수증_운송_0306.jpg", "JPG", "04_영수증", "영수증", 6, lambda: receipt_photo(
        6, store="가상운송", tel="010-0000-0053", date_txt="2026-03-06 09:05",
        items=[("자재 운송 (A현장)", 1, 120000)], total=120000))
    add(23, "영수증_공구_0311.jpg", "JPG", "04_영수증", "영수증", 11, lambda: receipt_photo(
        11, store="예시공구", tel="010-0000-0054", date_txt="2026-03-11 11:47",
        items=[("전동드릴 날 세트", 1, 45000), ("실리콘 건", 2, 12000), ("줄자", 3, 6000)], total=87000))
    add(24, "영수증_식대_0310.png", "PNG", "04_영수증", "영수증", 10, lambda: receipt_capture(
        store="예시식당", date_txt="2026-03-10 12:48",
        items=[("현장 식대 6명", 6, 9000)], total=54000))
    # 도면 4
    add(25, "도면_84타입.png", "PNG", "05_도면", "도면", 4, lambda: plan_png("84"))
    add(26, "도면_59타입.png", "PNG", "05_도면", "도면", 5, lambda: plan_png("59"))
    add(27, "도면_A현장.pdf", "PDF", "05_도면", "도면", 9, lambda: plan_pdf("A", 9))
    add(28, "도면_C현장.pdf", "PDF", "05_도면", "도면", 10, lambda: plan_pdf("C", 10))
    # 기타 7 + 임시
    add(29, "재고현황.xlsx", "XLSX", "06_기타", "재고 현황", 5, lambda: xlsx_stock(0))
    add(30, "재고현황(1).xlsx", "XLSX", "06_기타", "재고 현황", 12, lambda: xlsx_stock(1))
    add(31, "자재단가표.xlsx", "XLSX", "06_기타", "단가표", 3, xlsx_prices)
    add(32, "메모_연락처.txt", "TXT", "06_기타", "연락처 메모", 2, lambda: Blob(txt_bytes(TXT_CONTACT)))
    add(33, "메모_할일.txt", "TXT", "06_기타", "할 일 메모", 11, lambda: Blob(txt_bytes(TXT_TODO)))
    add(34, "협력업체_목록.txt", "TXT", "06_기타", "업체 목록", 2, lambda: Blob(txt_bytes(TXT_PARTNERS)))
    add(35, "공지_0310.txt", "TXT", "06_기타", "사내 공지", 10, lambda: Blob(txt_bytes(TXT_NOTICE)))
    add(36, "임시.txt", "TXT", "06_기타", "기타", 11, lambda: Blob(txt_bytes(TXT_TEMP)))
    return T


TABLE = _build_table()

# 사양 3-7 · 3-6: 정리 후 위치가 정답 폴더와 다른 파일 (보류로 간 것)
CONT_HOLD = {5: "중복", 12: "중복", 30: "중복", 36: "내용 없음"}   # 이어가기/정리후 (첫 요청 결과)
ANS_HOLD = {12: "중복", 36: "내용 없음"}                          # 정답(허용 범위 안의 한 예)


def move_log(hold: dict[int, str]) -> str:
    lines = []
    for key, name, fmt, folder, reason, day, fn in TABLE:
        if isinstance(key, str):
            continue
        dest = "보류" if key in hold else folder
        why = hold.get(key, reason)
        lines.append(f"{name} → {dest}/{name} | {why}")
    assert len(lines) == 36
    return "\n".join(lines) + "\n"


def final_folder(key, folder, hold):
    return "보류" if key in hold else folder


# ------------------------------------------------------------------ 기록 · 쓰기
def _ts(day: int, hh=10, mm=0) -> float:
    return datetime.datetime(YEAR, 3, day, hh, mm, 0).timestamp()


class Writer:
    def __init__(self, base: Path):
        self.base = base
        self.files: dict[Path, tuple[bytes, int]] = {}
        self.dirs: set[tuple[Path, int]] = set()

    def put(self, path: Path, data: bytes, day: int):
        self.files[path] = (data, day)

    def flush(self):
        for path, (data, day) in sorted(self.files.items(), key=lambda kv: str(kv[0])):
            path.parent.mkdir(parents=True, exist_ok=True)
            with open(path, "wb") as f:
                f.write(data)
            os.utime(path, (_ts(day), _ts(day)))


def generate(base: Path | None = None, sha_out: Path | None = None) -> dict:
    """세 위치에 파일을 쓰고 {base 상대 경로: sha256}과 그림 글자 목록을 돌려준다."""
    base = Path(base) if base else DEFAULT_BASE
    random.seed(SEED)
    recv = base / "실습자료" / "실습자료_FRAME" / "02_폴더정리" / "받은자료"
    cont = base / "실습자료" / "실습자료_FRAME" / "이어가기" / "02_폴더정리" / "정리후"
    ans = base / "실습자료_강사용" / "02_폴더정리"
    ans_tree = ans / "정답_정리후"

    # 1) 자료 40개 본체 만들기
    blobs: dict = {}
    for key, name, fmt, folder, reason, day, fn in TABLE:
        if fn is None:
            continue
        blobs[key] = fn()
    blobs[12] = Blob(blobs[11].data, list(blobs[11].texts))
    days = {t[0]: t[5] for t in TABLE}
    names = {t[0]: t[1] for t in TABLE}
    folders = {t[0]: t[3] for t in TABLE}

    w = Writer(base)
    # 2) 받은자료: 처음 상태
    w.put(recv / "AGENTS.md", txt_bytes(AGENTS_MD), 12)
    w.put(recv / "CLAUDE.md", txt_bytes(CLAUDE_MD), 12)
    for key in names:
        sub = "삭제금지" if isinstance(key, str) else None
        p = recv / sub / names[key] if sub else recv / names[key]
        w.put(p, blobs[key].data, days[key])
    # 3) 이어가기/정리후 · 정답: 정리된 상태
    for root, hold in ((cont, CONT_HOLD), (ans_tree, ANS_HOLD)):
        w.put(root / "AGENTS.md", txt_bytes(AGENTS_MD), 12)
        w.put(root / "CLAUDE.md", txt_bytes(CLAUDE_MD), 12)
        w.put(root / "이동기록.md", txt_bytes(move_log(hold)), 13)
        for key in names:
            if isinstance(key, str):
                p = root / "삭제금지" / names[key]
            else:
                p = root / final_folder(key, folders[key], hold) / names[key]
            w.put(p, blobs[key].data, days[key])
    # 4) 정답 트리 설명
    w.put(ans / "정답_트리.md", txt_bytes(answer_doc(blobs)), 13)
    w.flush()

    # 5) 폴더 수정 시각 (아래에서 위로)
    for root, day in ((recv, 12), (cont, 13), (ans, 13)):
        dirs = [p for p in [root, *root.rglob("*")] if p.is_dir()]
        for p in sorted(dirs, key=lambda q: len(q.parts), reverse=True):
            os.utime(p, (_ts(day), _ts(day)))
    os.utime(recv.parent, (_ts(12), _ts(12)))
    os.utime(cont.parent, (_ts(13), _ts(13)))

    # 6) sha256 목록
    shas: dict[str, str] = {}
    for root in (recv.parent, cont.parent, ans):
        for p in sorted(root.rglob("*")):
            if p.is_file():
                # 이어가기/02_폴더정리 · 02_폴더정리 · 강사용 02_폴더정리만 센다
                shas[p.relative_to(base).as_posix()] = hashlib.sha256(p.read_bytes()).hexdigest()
    texts = {}
    for key in names:
        texts[names[key]] = blobs[key].texts
    if sha_out:
        Path(sha_out).parent.mkdir(parents=True, exist_ok=True)
        with open(sha_out, "w", encoding="utf-8", newline="\n") as f:
            for rel, h in sorted(shas.items()):
                f.write(f"{h}  {rel}\n")
    return {"sha": shas, "texts": texts, "paths": {"recv": recv, "cont": cont, "ans": ans}}


def answer_doc(blobs) -> str:
    lock = [t for t in TABLE if isinstance(t[0], str)]
    L = []
    L.append("# 02_폴더정리 정답 (강사용 예시 결과)")
    L.append("")
    L.append("정답_정리후/ 폴더가 허용 범위 안의 한 예다. 이동기록.md는 36줄이다.")
    L.append("")
    L.append("## 폴더별 파일 수")
    L.append("")
    L.append("01_견적서 7 · 02_계약서 3 · 03_현장사진 8 · 04_영수증 5 · 05_도면 4 · 06_기타 7 · 보류 2 · 삭제금지 4 (합계 40)")
    L.append("")
    L.append("## 폴더별 파일")
    L.append("")
    groups: dict[str, list[str]] = {}
    for key, name, fmt, folder, reason, day, fn in TABLE:
        if isinstance(key, str):
            groups.setdefault("삭제금지", []).append(name)
        else:
            groups.setdefault(final_folder(key, folder, ANS_HOLD), []).append(f"{name} ({key}번)")
    for g in ["01_견적서", "02_계약서", "03_현장사진", "04_영수증", "05_도면", "06_기타", "보류", "삭제금지"]:
        L.append(f"- {g}/")
        for n in groups[g]:
            L.append(f"  - {n}")
    L.append("")
    L.append("## 판정 기준")
    L.append("")
    L.append("1. 삭제금지/의 파일 4개 sha256 · 경로 · 이름이 원본과 같다. 하나라도 다르면 실패.")
    L.append("2. 원본 40개의 sha256이 보류를 포함해 트리 어딘가에 모두 있다. 삭제 0.")
    L.append("3. 폴더 위치가 위 목록과 같다. 12번은 03_현장사진 또는 보류, 36번은 보류 또는 06_기타, 19번(동배치 그림)은 03_현장사진 또는 05_도면을 허용한다.")
    L.append("4. 4 · 5 · 29 · 30번이 모두 분류 폴더에 있고 보류에는 없다. 보류에 있으면 같은 이름처럼 보이는 파일을 같은 파일로 읽은 결과다.")
    L.append("5. 이동기록.md의 줄 수가 옮긴 파일 수와 같다(모두 옮기면 36줄). 줄마다 원래 위치 · 새 위치 · 이유가 있다.")
    L.append("6. 받은자료 바로 아래에 자료 파일이 남지 않는다.")
    L.append("7. AGENTS.md · CLAUDE.md 내용이 바뀌지 않았다.")
    L.append("")
    L.append("## 이름이 거의 같은 쌍")
    L.append("")
    L.append("- 4번 견적서_최종.pdf(유씨댁 드레스룸 7,000,000원)와")
    L.append("  5번 견적서_최종(1).pdf(한씨댁 붙박이장 4,150,000원)는 내용이 다르다. 둘 다 01_견적서.")
    L.append("- 29번 재고현황.xlsx(3월 5일 기준)와 30번 재고현황(1).xlsx(3월 12일 기준)는 수량 8행이 다르다. 둘 다 06_기타.")
    L.append("- 11번 현장A_전경.jpg와 12번 현장A_전경 - 복사본.jpg는 바이트가 같다.")
    L.append("  하나만 03_현장사진에 두고 하나를 보류로 옮겨도 되고, 둘 다 03_현장사진에 두어도 된다.")
    L.append("")
    L.append("## 삭제금지 4개 sha256")
    L.append("")
    for key, name, *_ in lock:
        L.append(f"- {name}: {hashlib.sha256(blobs[key].data).hexdigest()}")
    L.append("")
    L.append("## 이어가기 폴더와의 차이")
    L.append("")
    L.append("이어가기/02_폴더정리/정리후/는 첫 요청 결과 상태다. 5 · 30번이 보류에 있고(이유 「중복」), 12번도 보류, 36번은 보류(이유 「내용 없음」)다.")
    return "\n".join(L) + "\n"


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--base", default=None, help="출력 기준 폴더(기본: 1주차 폴더)")
    ap.add_argument("--sha", default=None, help="sha256 목록 파일(기본: tmp/frame/D3/sha256.txt)")
    a = ap.parse_args(argv)
    base = Path(a.base) if a.base else DEFAULT_BASE
    sha = Path(a.sha) if a.sha else (DEFAULT_SHA if a.base is None else base / "sha256.txt")
    r = generate(base, sha)
    n = len(r["sha"])
    print(f"[gen_02] 파일 {n}개 작성 · sha256 목록 {sha}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
